#!/usr/bin/env python3
"""Import monthly payroll records from a user-supplied CSV or XLSX file."""

import argparse
import csv
import sys
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from retro.modules.accountant.roster import RosterStore


FIELDS = ('external_key', 'name', 'role', 'salary', 'schedule', 'card', 'cash',
          'advances', 'remaining')


def read_rows(path: Path) -> list[dict]:
    if path.suffix.casefold() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as source:
            return [{key: (row.get(key) or '') for key in FIELDS}
                    for row in csv.DictReader(source)]
    if path.suffix.casefold() not in ('.xlsx', '.xlsm'):
        raise ValueError('Поддерживаются только CSV, XLSX и XLSM.')
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else '' for value in next(values)]
        return [{key: ('' if dict(zip(headers, row)).get(key) is None else dict(zip(headers, row))[key]) for key in FIELDS}
                for row in values]
    finally:
        workbook.close()


def import_rows(store: RosterStore, rows: list[dict]) -> tuple[int, int]:
    cleaned = []
    seen_keys = set()
    seen_names = set()
    for number, row in enumerate(rows, 2):
        key = str(row['external_key']).strip()
        name = str(row['name']).strip()
        if not name:
            continue
        if key and key in seen_keys:
            raise ValueError(f'Строка {number}: внешний ключ повторяется в файле.')
        if not key and name.casefold() in seen_names:
            raise ValueError(f'Строка {number}: имя без внешнего ключа неоднозначно.')
        seen_keys.add(key) if key else seen_names.add(name.casefold())
        values = dict(name=name, role=str(row['role']), salary=str(row['salary']),
                      schedule=str(row['schedule']), card=str(row['card']), cash=str(row['cash']),
                      advances=str(row['advances']), remaining=str(row['remaining']))
        store._monthly_values(**values)
        cleaned.append((key or None, values))

    created = updated = 0
    with closing(store._open()) as connection, connection:
        connection.execute('BEGIN IMMEDIATE')
        for key, values in cleaned:
            if key:
                matches = connection.execute(
                    'SELECT id FROM accountant_monthly_employees WHERE external_key = ?', (key,)).fetchall()
            else:
                matches = connection.execute(
                    'SELECT id FROM accountant_monthly_employees WHERE name = ? COLLATE NOCASE',
                    (values['name'],)).fetchall()
            if len(matches) > 1:
                raise ValueError('В базе найдено несколько сотрудников с одинаковым именем; добавьте внешний ключ.')
            name, role, schedule, money = store._monthly_values(**values)
            record = (name, role, str(money['salary']), schedule, str(money['card']),
                      str(money['cash']), str(money['advances']), str(money['remaining']))
            if matches:
                connection.execute(
                    'UPDATE accountant_monthly_employees SET name=?,role=?,salary=?,schedule=?,'
                    'card=?,cash=?,advances=?,remaining=? WHERE id=?', (*record, matches[0][0]))
                updated += 1
            else:
                connection.execute(
                    'INSERT INTO accountant_monthly_employees '
                    '(name,role,salary,schedule,card,cash,advances,remaining,external_key) '
                    'VALUES (?,?,?,?,?,?,?,?,?)', (*record, key))
                created += 1
    return created, updated


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--database', type=Path, required=True)
    args = parser.parse_args()
    try:
        created, updated = import_rows(RosterStore(args.database), read_rows(args.source))
    except (OSError, ValueError) as error:
        print(f'Ошибка: {error}')
        return 1
    print(f'Импорт завершён: добавлено {created}, обновлено {updated}.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
