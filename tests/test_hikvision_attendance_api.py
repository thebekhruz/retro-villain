from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from retro.app import create_app
from retro.config import HikvisionConfig, Settings
from retro.integrations.hikvision import HikvisionEvent, HikvisionPerson
from retro.integrations.hikvision_poller import PollResult


TZ = ZoneInfo('Asia/Tashkent')
DAY = date(2026, 9, 20)
CONFIG = HikvisionConfig(
    base_url='https://203.0.113.10', username='reader', password='secret',
    source='entry', poll_seconds=30, timeout_seconds=8, verify_tls=True)


class NoopPoller:
    def start(self):
        pass

    async def stop(self):
        pass


class SyncPoller(NoopPoller):
    async def sync_all_people(self):
        return {'people': 35, 'created': 35, 'linked': 0,
                'already_linked': 0, 'ambiguous': 0}

    async def run_once(self):
        return PollResult(True, events=14)


def live_client(tmp_path, *, complete=True):
    app = create_app(
        Settings(hikvision=CONFIG, data_dir=tmp_path, manual_handover_only=True),
        accountant_db_path=tmp_path / 'accountant.sqlite3',
        expense_db_path=tmp_path / 'cashier.sqlite3',
        director_db_path=tmp_path / 'director.sqlite3',
        hikvision_poller=NoopPoller())
    arrived = app.state.accountant_roster.add(
        name='Азиза Каримова', role='официант', rate='250000',
        group_name='Обслуживание зала')
    missing = app.state.accountant_roster.add(
        name='Бахром Алиев', role='бармен', rate='260000', group_name='Бар')
    app.state.accountant_roster.link_hikvision_people((
        HikvisionPerson('100', 'азиза каримова'),
        HikvisionPerson('200', 'бахром алиев'),
    ))
    app.state.attendance_store.ingest(
        HikvisionEvent('entry', 'serial-1', '100',
                        datetime(2026, 9, 20, 9, 12, tzinfo=TZ)), arrived.id)
    if complete:
        app.state.attendance_store.record_success(
            'entry', at=datetime(2026, 9, 21, 0, 5, tzinfo=TZ),
            cursor_at=datetime(2026, 9, 21, 0, 5, tzinfo=TZ),
            covered_from=datetime(2026, 9, 20, 0, 0, tzinfo=TZ),
            covered_through=datetime(2026, 9, 21, 0, 0, tzinfo=TZ))
    return TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)), arrived, missing


def test_accountant_and_director_share_real_first_entries_without_pay_leak(tmp_path):
    client, arrived, missing = live_client(tmp_path)
    with client:
        accountant = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        director = client.get('/api/director/attendance', params={'date': DAY.isoformat()}).json()

    assert accountant['demo'] is False
    assert accountant['attendance']['complete'] is True
    assert accountant['attendance']['status'] in ('ok', 'stale')
    by_id = {row['employee_id']: row for row in accountant['employees']}
    assert by_id[arrived.id]['status'] == 'on_time'
    assert by_id[arrived.id]['first_entry'].startswith('2026-09-20T09:12:00')
    assert by_id[missing.id]['status'] == 'missing'
    assert accountant['payroll']['missing_count'] == 1
    assert accountant['payroll']['unavailable_count'] == 0
    assert director['demo'] is False
    assert director['attendance'] == accountant['attendance']
    assert {row['employee_id']: row['status'] for row in director['employees']} == {
        arrived.id: 'on_time', missing.id: 'missing'}
    assert all('rate' not in row and 'payable' not in row for row in director['employees'])


def test_authenticated_accountant_can_trigger_full_hikvision_people_sync(tmp_path):
    app = create_app(
        Settings(hikvision=CONFIG, data_dir=tmp_path, manual_handover_only=True),
        accountant_db_path=tmp_path / 'accountant.sqlite3',
        expense_db_path=tmp_path / 'cashier.sqlite3',
        director_db_path=tmp_path / 'director.sqlite3',
        hikvision_poller=SyncPoller())

    with TestClient(app, base_url='http://127.0.0.1',
                    client=('127.0.0.1', 50000)) as client:
        response = client.post('/api/accountant/hikvision/sync-people')

    assert response.status_code == 200
    assert response.json() == {
        'source': 'Hikvision ISAPI', 'people': 35, 'created': 35, 'linked': 0,
        'already_linked': 0, 'ambiguous': 0, 'events': 14}


def test_incomplete_source_is_unavailable_and_payroll_confirmation_is_blocked(tmp_path):
    client, arrived, missing = live_client(tmp_path, complete=False)
    with client:
        day = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        confirmation = client.post('/api/accountant/payroll/confirm', json={
            'date': DAY.isoformat(), 'approver': 'Бухгалтер'})

    statuses = {row['employee_id']: row['status'] for row in day['employees']}
    assert statuses[arrived.id] == 'on_time'
    assert statuses[missing.id] == 'unavailable'
    assert day['payroll']['missing_count'] == 0
    assert day['payroll']['unavailable_count'] == 1
    assert confirmation.status_code == 409
    assert 'Hikvision' in confirmation.json()['detail']


def test_real_employee_and_entrance_exports_use_same_rows_without_demo_claim(tmp_path):
    client, arrived, missing = live_client(tmp_path)
    with client:
        employees = client.get('/api/accountant/employees/export', params={
            'date': DAY.isoformat(), 'scope': 'all'})
        entrances = client.get('/api/accountant/entrances/export', params={
            'date': DAY.isoformat()})

    assert employees.status_code == entrances.status_code == 200
    assert 'DEMO' not in employees.headers['content-disposition']
    employee_sheet = load_workbook(BytesIO(employees.content), data_only=True).active
    entrance_sheet = load_workbook(BytesIO(entrances.content), data_only=True).active
    assert 'ДЕМО' not in employee_sheet['A1'].value
    assert employee_sheet['D7'].value == 'Вовремя'
    assert employee_sheet['D8'].value == 'Не пришёл'
    assert entrance_sheet['B4'].value == 1
    assert entrance_sheet['B7'].value == 'Азиза Каримова'
