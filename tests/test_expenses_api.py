from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
from fastapi.testclient import TestClient

from retro.app import create_app
from retro.config import Settings
from retro.modules.cashier.expenses import seed_cashier_expense
from retro.modules.cashier.service import build_snapshot


DAY = date(2026, 9, 12)


def row(*values):
    return {f'field{i}': {'value': value} for i, value in enumerate(values)}


def test_manual_expenses_are_saved_by_day_and_can_be_removed(tmp_path):
    path = tmp_path / 'cashier.sqlite3'
    with TestClient(create_app(Settings(), expense_db_path=path), client=('127.0.0.1', 50000)) as client:
        created = client.post('/api/cashier/expenses', json={
            'date': '2026-09-12', 'description': 'Зарплата', 'amount': '350000',
        })
        assert created.status_code == 201
        expense_id = created.json()['id']
        assert client.get('/api/cashier/expenses?date=2026-09-12').json() == {
            'date': '2026-09-12',
            'expenses': [{'id': expense_id, 'description': 'Зарплата', 'amount': '350000'}],
            'total': '350000',
            'expense_policy_configured': False,
        }
        assert client.get('/api/cashier/expenses?date=2026-09-11').json()['total'] == '0'
        assert client.delete(f'/api/cashier/expenses/{expense_id}?date=2026-09-11').status_code == 404
        assert client.delete(f'/api/cashier/expenses/{expense_id}?date=2026-09-12').status_code == 204
    with TestClient(create_app(Settings(), expense_db_path=path), client=('127.0.0.1', 50000)) as client:
        assert client.get('/api/cashier/expenses?date=2026-09-12').json() == {
            'date': '2026-09-12',
            'expenses': [],
            'total': '0',
            'expense_policy_configured': False,
        }


def test_manual_expense_is_not_duplicated_by_runtime_policy(tmp_path):
    with TestClient(create_app(Settings(), expense_db_path=tmp_path / 'cashier.sqlite3'),
                    client=('127.0.0.1', 50000)) as client:
        for day, description in [('2026-04-25', 'Выплата за смену'), ('2026-09-15', 'Зарплата')]:
            saved = client.post('/api/cashier/expenses', json={
                'date': day, 'description': description, 'amount': '350000',
            }).json()
            result = client.get('/api/cashier/expenses', params={'date': day}).json()
            assert result['total'] == '350000'
            assert result['expenses'] == [saved]


def test_daily_salary_is_added_to_other_expenses_and_export(tmp_path):
    path = tmp_path / 'cashier.sqlite3'
    seed_cashier_expense(path, DAY, DAY, 'Зарплата', Decimal('350000'))
    app = create_app(Settings(), expense_db_path=path)
    snapshot = build_snapshot(DAY, [row('2026-09-12', 2, 450000)],
                              [row('Демо', 100000), row('UzCard', 350000)])
    app.state.cache.put(snapshot)
    with TestClient(app, client=('127.0.0.1', 50000)) as client:
        client.post('/api/cashier/expenses', json={
            'date': DAY.isoformat(), 'description': 'Такси', 'amount': '20000',
        })
        result = client.get('/api/cashier/expenses', params={'date': DAY.isoformat()}).json()
        assert result['total'] == '370000'
        assert [(item['description'], item['amount']) for item in result['expenses']] == [
            ('Зарплата', '350000'), ('Такси', '20000')]
        response = client.get('/api/cashier/export', params={
            'date': DAY.isoformat(), 'snapshot_id': snapshot.id,
        })
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        assert workbook['отчет']['B28'].value == 370000
        assert workbook['отчет']['B29'].value == -270000
        assert workbook['Расходы']['A8'].value == 'Зарплата'
        assert workbook['Расходы']['B8'].value == 350000


def test_invalid_expense_does_not_change_saved_data(tmp_path):
    with TestClient(create_app(Settings(), expense_db_path=tmp_path / 'cashier.sqlite3'),
                    client=('127.0.0.1', 50000)) as client:
        for body in [dict(date='2026-09-12', description='', amount='100'),
                     dict(date='2026-09-12', description='Такси', amount='-1'),
                     dict(date='2099-01-01', description='Такси', amount='100')]:
            assert client.post('/api/cashier/expenses', json=body).status_code == 422
        assert client.get('/api/cashier/expenses?date=2026-09-12').json()['total'] == '0'


def test_export_includes_manual_expenses_and_demo_less_expenses(tmp_path):
    path = tmp_path / 'cashier.sqlite3'
    seed_cashier_expense(path, DAY, DAY, 'Зарплата', Decimal('350000'))
    app = create_app(Settings(), expense_db_path=path)
    snapshot = build_snapshot(DAY, [row('2026-09-12', 2, 450000)],
                              [row('Демо', 100000), row('UzCard', 350000)])
    app.state.cache.put(snapshot)
    with TestClient(app, client=('127.0.0.1', 50000)) as client:
        client.post('/api/cashier/expenses', json={
            'date': '2026-09-12', 'description': 'Такси', 'amount': '20000',
        })
        response = client.get('/api/cashier/export', params={
            'date': '2026-09-12', 'snapshot_id': snapshot.id,
        })
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        assert workbook['отчет']['B2'].value == 450000
        assert workbook['отчет']['B28'].value == 370000
        assert workbook['отчет']['B29'].value == -270000
        assert workbook['Расходы']['A8'].value == 'Зарплата'
        assert workbook['Расходы']['B8'].value == 350000
        assert workbook['Расходы']['A9'].value == 'Такси'
        assert workbook['Расходы']['B9'].value == 20000


def test_other_cash_receipts_affect_handover_and_export(tmp_path):
    from dataclasses import replace
    path = tmp_path / 'cashier.sqlite3'
    seed_cashier_expense(path, DAY, DAY, 'Зарплата', Decimal('350000'))
    app = create_app(Settings(), expense_db_path=path)
    snapshot = replace(build_snapshot(DAY, [row('2026-09-12', 2, 450000)],
                                      [row('Демо', 100000), row('UzCard', 350000)]),
                       cash_prepayment=Decimal('600000'), new_prepayment=Decimal('600000'))
    app.state.cache.put(snapshot)
    with TestClient(app, client=('127.0.0.1', 50000)) as client:
        receipt = client.post('/api/cashier/receipts', json={
            'date': DAY.isoformat(), 'description': 'Вернули долг', 'amount': '854000',
        })
        assert receipt.status_code == 201
        assert client.get('/api/cashier/receipts', params={'date': DAY.isoformat()}).json()['total'] == '854000'
        response = client.get('/api/cashier/export', params={'date': DAY.isoformat(),
                                                            'snapshot_id': snapshot.id})
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        assert workbook['отчет']['B20'].value == 600000
        assert workbook['отчет']['B21'].value == 1904000
        assert workbook['отчет']['B29'].value == 1204000
        assert workbook['Расходы']['B6'].value == 1204000
        assert client.delete(f'/api/cashier/receipts/{receipt.json()["id"]}',
                             params={'date': DAY.isoformat()}).status_code == 204
