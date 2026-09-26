"""Кабинет учредителя (T-377): неделя по дням, дивиденды, прогноз, Счёт Шефа."""

from dataclasses import replace
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from retro.app import create_app
from retro.config import Settings
from retro.integrations.iiko import chef_bills_from_olap, daily_orders_from_olap
from retro.modules.cashier.service import (
    DataError, Payment, RETRO_REGISTER, SCHOOL_REGISTER, RevenueBreakdown, demo_snapshot,
)
from retro.modules.founder import overview
from retro.modules.founder.dividends import DividendTargetError, DividendTargetStore

TODAY = date(2026, 9, 24)  # четверг
MONDAY = date(2026, 9, 21)


def row(*values, children=None):
    result = {f'field{index}': {'value': value} for index, value in enumerate(values) if value is not None}
    if children:
        result['children'] = children
    return result


def snapshot(day, *, retro='30000000', school='4000000', demo='9000000'):
    return replace(demo_snapshot(day), demo=False,
                   payments=(Payment('Демо', Decimal(demo)), Payment('UzCard', Decimal('5000000'))),
                   cash_prepayment=Decimal(0),
                   revenue_breakdown=RevenueBreakdown(Decimal(retro), Decimal(school), Decimal(0)))


class FakeIiko:
    def __init__(self, *, fail_days=(), orders=None, chef=None, error=None):
        self.fail_days, self.orders, self.chef, self.error = set(fail_days), orders, chef, error

    async def load(self, day):
        if day in self.fail_days:
            raise DataError('iiko временно недоступен')
        return snapshot(day)

    async def load_daily_orders(self, start, end):
        if self.error:
            raise self.error
        if self.orders is not None:
            return self.orders
        rows = []
        day = start
        while day <= end:
            rows.append(dict(day=day.isoformat(), direction='retro', orders=100 + day.weekday(),
                             revenue=str(30000000 + day.weekday() * 1000000)))
            rows.append(dict(day=day.isoformat(), direction='school', orders=50, revenue='4000000'))
            day += timedelta(days=1)
        return rows

    async def load_chef_bills(self, start, end):
        if self.error:
            raise self.error
        return self.chef or []


def client(tmp_path, iiko=None, **settings):
    app = create_app(Settings(manual_handover_only=True, **settings),
                     expense_db_path=tmp_path / 'cashier.sqlite3',
                     accountant_db_path=tmp_path / 'accountant.sqlite3',
                     founder_db_path=tmp_path / 'founder.sqlite3')
    app.state.iiko = iiko or FakeIiko()
    return TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000))


@pytest.fixture
def today():
    with patch('retro.modules.founder.cabinet.today_tashkent', return_value=TODAY), \
            patch('retro.modules.founder.routes.today_tashkent', return_value=TODAY), \
            patch('retro.modules.accountant.routes.today_tashkent', return_value=TODAY):
        yield TODAY


# ── Чистые расчёты ─────────────────────────────────────────────────────────

def test_week_runs_monday_to_sunday_with_iso_label():
    assert overview.week_of(date(2026, 9, 27)) == (MONDAY, date(2026, 9, 27), '2026-W39')
    assert overview.parse_week('2026-W39') == MONDAY
    for broken in ('2026-39', '2026-W60', 'week', '2026-W3'):
        with pytest.raises(ValueError):
            overview.parse_week(broken)


def test_dividends_are_only_the_safe_transfer_and_direct_payout():
    def kind(type_, code=None):
        return overview.flow_kind(dict(type=type_, item_code=code, amount='1'))
    assert kind('reserve_transfer') == 'dividends'
    assert kind('other_expense', 'distribution_dividends') == 'dividends'
    assert kind('other_expense', 'distribution_oxbridge') == 'other'  # перевод Oxbridge — не дивиденды
    assert kind('other_expense', 'salary_monthly') == 'salary'
    assert kind('salary_payment') == 'salary'
    assert kind('other_expense', 'proc_meat') == 'procurement'
    assert kind('procurement_advance') == 'procurement'
    assert kind('handover') == 'handover'
    assert kind('other_receipt', 'income_other') == 'receipt'


def test_dividend_plan_grows_by_day_and_flags_falling_behind():
    flows = overview.daily_flows([
        dict(day='2026-09-21', type='reserve_transfer', amount='1000000'),
        dict(day='2026-09-22', type='reserve_transfer', amount='1000000'),
    ])
    result = overview.dividend_week(TODAY, Decimal('10000000'), flows)
    assert result['week'] == '2026-W39'
    assert result['collected'] == '2000000.00'
    assert result['pace'] == '5714285.71'  # четыре седьмых недели
    assert result['behind'] is True
    # Осталось 8 млн на четыре дня с сегодняшним — по 2 млн, шаг 50 000.
    assert result['suggest_today'] == '2000000.00'
    assert result['days_left'] == 4
    assert [day['future'] for day in result['days']] == [False] * 4 + [True] * 3


def test_dividend_week_without_target_does_not_invent_a_plan():
    result = overview.dividend_week(TODAY, None, {})
    assert result['target'] is None and result['pace'] is None
    assert result['behind'] is False and result['suggest_today'] is None


def test_free_cash_skips_days_without_a_recorded_handover():
    flows = overview.daily_flows([
        dict(day='2026-09-22', type='handover', amount='10000000'),
        dict(day='2026-09-22', type='salary_payment', amount='2000000'),
        dict(day='2026-09-22', type='procurement_advance', amount='3000000'),
        dict(day='2026-09-23', type='salary_payment', amount='9000000'),  # передачи нет
    ])
    free = overview.free_cash_per_week(flows, [date(2026, 9, 22), date(2026, 9, 23)])
    assert free == Decimal('35000000')  # (10 − 2 − 3) млн × 7
    assert overview.free_cash_per_week({}, [date(2026, 9, 22)]) is None


def test_handover_check_tolerates_rounding_only():
    assert overview.handover_check('1000000', '1000000.40')['status'] == 'ok'
    short = overview.handover_check('700000', '1000000')
    assert short == {'status': 'mismatch', 'difference': '-300000.00'}
    assert overview.handover_check(None, '1000000')['status'] == 'missing'
    assert overview.handover_check('1000000', None)['status'] == 'unknown'


def test_expense_categories_leave_out_income_and_split_dividends():
    result = overview.expense_categories([
        dict(day='2026-09-21', type='handover', amount='50000000'),
        dict(day='2026-09-21', type='other_receipt', item_code='income_other', amount='100'),
        dict(day='2026-09-21', type='salary_payment', amount='2000000'),
        dict(day='2026-09-21', type='other_expense', item_code='salary_monthly', amount='3000000'),
        dict(day='2026-09-21', type='other_expense', item_code='proc_shoh', amount='1000000'),
        dict(day='2026-09-21', type='reserve_transfer', amount='4000000'),
        dict(day='2026-09-21', type='other_expense', item_code='utilities', amount='500000'),
    ])
    labels = {item['label']: item['amount'] for item in result['categories']}
    assert result['total'] == '10500000.00'
    assert labels == {'Дивиденды': '4000000.00', 'Зарплаты · оклады': '3000000.00',
                      'Зарплаты · сменные': '2000000.00', 'Закуп · наличные Шоху': '1000000.00',
                      'Коммунальные и охрана': '500000.00'}
    assert result['categories'][0]['label'] == 'Дивиденды'


def test_weekday_forecast_averages_history_and_ignores_the_open_day():
    history = overview.register_days([
        dict(day='2026-09-17', direction='retro', orders=100, revenue='30000000'),  # чт
        dict(day='2026-09-10', direction='retro', orders=120, revenue='40000000'),  # чт
        dict(day='2026-09-17', direction='school', orders=60, revenue='4000000'),
        dict(day='2026-09-24', direction='retro', orders=5, revenue='900000'),  # сегодня
    ])
    forecast = overview.weekday_forecast(history, TODAY)
    thursday = forecast[3]
    assert thursday['label'] == 'Чт' and thursday['samples'] == 2
    assert thursday['retro'] == '35000000.00'
    assert thursday['school'] == '4000000.00'  # одно значение — само и есть среднее
    assert thursday['orders'] == 170
    assert forecast[0]['revenue'] is None  # по понедельникам истории нет


def test_month_forecast_is_fact_through_yesterday_plus_weekday_averages():
    history = overview.register_days([
        dict(day='2026-09-01', direction='retro', orders=1, revenue='10'),
        dict(day='2026-09-23', direction='school', orders=1, revenue='5'),
        dict(day='2026-09-24', direction='retro', orders=1, revenue='1000'),  # сегодня — не факт
    ])
    weekdays = [dict(revenue='1') for _ in range(7)]
    result = overview.month_forecast(history, weekdays, TODAY)
    assert result['fact'] == '15.00'
    assert result['forecast'] == '7.00'  # 24–30 сентября
    assert result['fact_through'] == '2026-09-23'


def test_chef_bills_flag_the_limit_and_split_week_from_month():
    rows = [dict(day='2026-09-22', order_id='a', table=3, waiters=['Алина'], amount='610000', cost='1'),
            dict(day='2026-09-22', order_id='b', table=4, waiters=[], amount='90000', cost='1'),
            dict(day='2026-09-02', order_id='c', table=5, waiters=[], amount='400000', cost='1')]
    result = overview.chef_bills(rows, week_start=MONDAY, week_end=date(2026, 9, 27),
                                 month_start=date(2026, 9, 1))
    assert [bill['order_id'] for bill in result['week']] == ['a', 'b']
    assert result['week_over'] == 1 and result['month_over'] == 1  # ровно 400 000 — не выше
    assert result['week_total'] == '700000.00' and result['month_total'] == '1100000.00'


def test_shokh_month_separates_advance_from_direct_procurement():
    flows = [dict(day='2026-09-21', type='procurement_advance', description='Шох: базар', amount='3000000'),
             dict(day='2026-09-21', type='other_expense', item_code='proc_bread', amount='200000'),
             dict(day='2026-09-22', type='other_expense', item_code='proc_shoh', amount='1000000')]
    purchases = [dict(id=1, day='2026-09-21', item='Лук', total='100000', has_photo=False,
                      price_above_usual=False, accepted_at=None),
                 dict(id=2, day='2026-09-21', item='Мясо', total='900000', has_photo=True,
                      price_above_usual=False, accepted_at=None)]
    result = overview.shokh_month(purchases, flows, pocket='3000000')
    assert result['given'] == '4000000.00'
    assert result['direct'] == '200000.00'
    assert result['spent'] == '1000000.00'
    assert [row['reason'] for row in result['flagged']] == ['нет фото']
    assert result['top_items'][0] == {'item': 'Мясо', 'amount': '900000.00'}


def test_dividend_target_store_inherits_last_week_and_keeps_history(tmp_path):
    store = DividendTargetStore(tmp_path / 'a.sqlite3')
    assert store.get('2026-W39') is None
    store.set('2026-W38', '8000000', 'founder')
    inherited = store.get('2026-W39')
    assert inherited['amount'] == '8000000' and inherited['inherited'] is True
    store.set('2026-W39', '10 000 000', 'founder')
    assert store.get('2026-W39')['amount'] == '10000000'
    for broken in ('-1', 'много', '1.5'):
        with pytest.raises(DividendTargetError):
            store.set('2026-W39', broken, 'founder')


# ── iiko ──────────────────────────────────────────────────────────────────

def test_daily_orders_map_registers_to_venues_and_refuse_unknown_ones():
    rows = [row('2026-09-21', children=[row(None, RETRO_REGISTER, 110, 30000000),
                                        row(None, SCHOOL_REGISTER, 60, 4000000)])]
    assert daily_orders_from_olap(rows) == [
        dict(day='2026-09-21', direction='retro', orders=110, revenue='30000000'),
        dict(day='2026-09-21', direction='school', orders=60, revenue='4000000')]
    with pytest.raises(DataError):
        daily_orders_from_olap([row('2026-09-21', 'Новая касса', 1, 1)])


def test_chef_bills_join_every_waiter_of_one_order():
    rows = [row('2026-09-22', children=[row(None, 'order-1', children=[row(None, None, 7, children=[
        row(None, None, None, 'Алина', 300000, 90000.5), row(None, None, None, 'Музаффар', 150000, 40000)])])])]
    assert chef_bills_from_olap(rows) == [dict(day='2026-09-22', order_id='order-1', table=7,
                                               waiters=['Алина', 'Музаффар'], amount='450000',
                                               cost='130000.50')]


# ── API ───────────────────────────────────────────────────────────────────

def test_founder_sets_weekly_dividends_and_accountant_sees_the_target(tmp_path, today):
    with client(tmp_path) as c:
        empty = c.get('/api/founder/dividends/weekly').json()
        assert empty['target'] is None and empty['week'] == '2026-W39'
        saved = c.post('/api/founder/dividends/weekly', json={'week': '2026-W39', 'amount': '10000000'})
        assert saved.status_code == 200
        assert saved.json()['target'] == '10000000.00'
        assert saved.json()['target_source']['inherited'] is False
        day = c.get('/api/accountant/day', params={'date': TODAY.isoformat()}).json()
        assert day['dividends_week']['target'] == '10000000.00'
        assert day['dividends_week']['behind'] is True  # к четвергу ничего не отложено


@pytest.mark.parametrize('body', [
    {'week': '2026-W60', 'amount': '1'}, {'week': '2026-W37', 'amount': '1'},
    {'week': '2026-W39', 'amount': 'много'}, {'week': '2026-W39', 'amount': '-5'},
])
def test_dividend_target_refuses_bad_or_closed_weeks(tmp_path, today, body):
    with client(tmp_path) as c:
        assert c.post('/api/founder/dividends/weekly', json=body).status_code == 422


def test_payout_counts_only_money_put_into_the_safe(tmp_path, today):
    with client(tmp_path) as c:
        finance = c.app.state.accountant_finance
        finance.record_handover(MONDAY, Decimal('20000000'))
        finance.set_cash_opening(MONDAY, '0', 'Старт')
        finance.reserve_entry(MONDAY, 'dividends', 'opening', '0', 'Сейф', cashier_amount=Decimal('20000000'))
        finance.reserve_entry(MONDAY, 'dividends', 'transfer', '3000000', 'В сейф', cashier_amount=Decimal('20000000'))
        finance.add_expense(MONDAY, 'distribution_oxbridge', 'Oxbridge', '1000000', cashier_amount=Decimal('20000000'))
        data = c.get('/api/founder/dividends/weekly').json()
        assert data['collected'] == '3000000.00'


def test_week_marks_a_short_handover_and_keeps_iiko_gaps_empty(tmp_path, today):
    iiko = FakeIiko(fail_days={date(2026, 9, 22)})
    with client(tmp_path, iiko) as c:
        finance = c.app.state.accountant_finance
        # Расчёт кассира по фейку: «Демо» 9 млн, расходов нет — ровно 9 млн.
        finance.record_handover(MONDAY, Decimal('9000000'))
        finance.record_handover(date(2026, 9, 23), Decimal('8700000'))
        week = c.get('/api/founder/week').json()
    days = {day['date']: day for day in week['days']}
    assert week['week'] == '2026-W39'
    assert days['2026-09-21']['handover']['status'] == 'ok'
    assert days['2026-09-23']['handover'] == {'recorded': '8700000.00', 'expected': '9000000.00',
                                              'status': 'mismatch', 'difference': '-300000.00'}
    # iiko не ответил — день не превращается в нули.
    assert days['2026-09-22']['cashier'] is None
    assert days['2026-09-22']['cashier_error'] == 'iiko временно недоступен'
    assert days['2026-09-22']['handover']['status'] == 'unknown'
    assert days['2026-09-24']['handover']['status'] == 'pending'  # смена ещё идёт
    assert days['2026-09-25'] == {'date': '2026-09-25', 'weekday': 4, 'future': True, 'today': False}
    assert days['2026-09-21']['orders'] == {'retro': 100, 'school': 50}
    assert days['2026-09-24']['outlook']['estimate'] is True


def test_forecast_and_chef_account_report_iiko_errors_instead_of_zeros(tmp_path, today):
    with client(tmp_path, FakeIiko(error=DataError('iiko недоступен'))) as c:
        forecast = c.get('/api/founder/forecast').json()
        chef = c.get('/api/founder/chef-account').json()
        week = c.get('/api/founder/week').json()
    assert forecast == {'estimate': True, 'error': 'iiko недоступен', 'weekdays': None,
                        'month': None, 'today': None}
    assert chef == {'error': 'iiko недоступен'}
    assert week['orders_error'] == 'iiko недоступен'


def test_forecast_uses_eight_weeks_of_history(tmp_path, today):
    with client(tmp_path) as c:
        data = c.get('/api/founder/forecast').json()
    assert data['estimate'] is True and data['weeks'] == 8
    assert data['weekdays'][3]['retro'] == '33000000.00'
    assert data['today']['orders'] == {'retro': 103, 'school': 50}
    assert data['month']['fact_through'] == '2026-09-23'


def test_chef_account_lists_bills_with_the_limit(tmp_path, today):
    chef = [dict(day='2026-09-23', order_id='x', table=2, waiters=['Алина'], amount='520000', cost='1')]
    with client(tmp_path, FakeIiko(chef=chef)) as c:
        data = c.get('/api/founder/chef-account').json()
    assert data['limit'] == '400000.00'
    assert data['week_over'] == 1 and data['week'][0]['over'] is True


def test_month_export_is_an_xlsx_even_without_iiko(tmp_path, today):
    with client(tmp_path, FakeIiko(error=DataError('iiko недоступен'))) as c:
        c.app.state.accountant_finance.record_handover(MONDAY, Decimal('9000000'))
        response = c.get('/api/founder/export/month', params={'month': '2026-09'})
        assert c.get('/api/founder/export/month', params={'month': '2026-10'}).status_code == 422
    assert response.status_code == 200
    assert 'Retro-accountant-2026-09.xlsx' in response.headers['content-disposition']
    book = load_workbook(BytesIO(response.content))
    sheet = book['По дням']
    assert 'iiko недоступен' in sheet['A2'].value
    assert sheet.max_row == 4 + 24  # шапка и дни с 1 по 24 сентября
    assert sheet['F25'].value == 9000000  # получено от кассира 21.09
    assert sheet['B25'].value is None  # iiko нет — выручка пустая, а не ноль
    assert 'Расходы' in book.sheetnames


def test_cabinet_is_closed_to_other_roles(tmp_path):
    users = {'founder': ('password', 'founder'), 'director': ('password', 'director')}
    app = create_app(Settings(dashboard_panel_users=users, data_dir=tmp_path))
    app.state.iiko = FakeIiko()
    with TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)) as c:
        c.post('/api/session', json={'username': 'director', 'password': 'password'})
        assert c.get('/api/founder/dividends/weekly').status_code == 403
        assert c.post('/api/founder/dividends/weekly', json={'week': '2026-W39', 'amount': '1'}).status_code == 403
        c.post('/api/session/logout')
        c.post('/api/session', json={'username': 'founder', 'password': 'password'})
        assert c.get('/api/founder/dividends/weekly').status_code == 200
        assert c.get('/api/director/team').status_code == 403


def test_founder_pages_split_cabinet_and_analytics(tmp_path):
    app = create_app(Settings(data_dir=tmp_path))
    with TestClient(app, client=('127.0.0.1', 50000)) as c:
        cabinet = c.get('/founder').text
        analytics = c.get('/founder/analytics').text
    assert 'founder-cabinet.js' in cabinet and 'id="fo-week"' in cabinet
    assert 'href="/founder/analytics"' in cabinet
    assert 'founder-chat.js' in analytics  # прежняя аналитика с чатом-панелью на месте


# ── Регрессии из ревью ─────────────────────────────────────────────────────

def test_on_plan_morning_is_not_behind_and_suggestion_counts_today():
    monday = overview.daily_flows([dict(day='2026-09-21', type='reserve_transfer', amount='100000')])
    result = overview.dividend_week(MONDAY, Decimal('700000'), monday)
    # План на сегодня 100 000 уже отложен — больше не советуем.
    assert result['suggest_today'] == '0.00'
    assert result['behind'] is False
    assert overview.dividend_week(MONDAY, Decimal('700000'), {})['behind'] is False  # утро понедельника


def test_suggestion_never_exceeds_what_is_left():
    sunday = date(2026, 9, 27)
    flows = overview.daily_flows([dict(day='2026-09-21', type='reserve_transfer', amount='690000')])
    result = overview.dividend_week(sunday, Decimal('700000'), flows)
    assert result['left'] == '10000.00' and result['suggest_today'] == '10000.00'


def test_forecast_survives_a_weekday_with_zero_orders():
    history = overview.register_days([dict(day='2026-09-19', direction='retro', orders=0, revenue='100')])
    assert overview.weekday_forecast(history, TODAY)[5]['orders'] == 0


def test_week_label_must_be_canonical():
    for broken in ('2026-W 9', '2026-W+9'):
        with pytest.raises(ValueError):
            overview.parse_week(broken)
