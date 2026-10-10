"""ТЗ 09.10.2026, Б-08 и Б-09: доп. зарплата и дивиденды в сводках бухгалтера,
числа дэшборда C-01 и перенос старых строк журнала в понятные статьи.

Сценарии с базой гоняются на SQLite и, если задан RETRO_TEST_POSTGRES_URL, на
Postgres — как на проде.
"""
import json
import os
from contextlib import closing, contextmanager
from datetime import date
from decimal import Decimal
from io import BytesIO
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from legacy_app import create_app
from retro.config import Settings
from retro.modules.accountant import closing as month_closing
from retro.modules.accountant.expense_catalog import EXTRA_SALARY_ITEM, GROUPS, ITEMS
from retro.modules.accountant.ledger import FinanceStore, LedgerError, flow_json
from retro.modules.accountant.opening_migration import apply_october_opening
from retro.modules.accountant.reclassify import candidates, reclassify
from retro.modules.accountant.reserves import is_monthly_salary
from retro.modules.cashier.till import give_shokh, shokh_gives
from scripts import reclassify_expenses

POSTGRES_URL = os.getenv('RETRO_TEST_POSTGRES_URL', '')
DATABASES = ['sqlite', pytest.param('postgres', marks=pytest.mark.skipif(
    not POSTGRES_URL, reason='RETRO_TEST_POSTGRES_URL не задан'))]


def D(day):  # noqa: N802 — день октября 2026, как в ТЗ
    return date(2026, 10, day)


@contextmanager
def postgres_url():
    """Своя схема на тест: таблицы других тестов не мешают."""
    import psycopg
    from psycopg import sql
    schema = 'categories_' + uuid4().hex
    admin = psycopg.connect(POSTGRES_URL, autocommit=True)
    admin.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(schema)))
    parts = urlsplit(POSTGRES_URL)
    query = dict(parse_qsl(parts.query), options='-csearch_path=' + schema)
    try:
        yield urlunsplit(parts._replace(query=urlencode(query)))
    finally:
        admin.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(schema)))
        admin.close()


@pytest.fixture(params=DATABASES)
def store(request, tmp_path):
    if request.param == 'postgres':
        with postgres_url() as url:
            yield FinanceStore(url)
    else:
        yield FinanceStore(tmp_path / 'accountant.sqlite3')


def confirmed(store, day, amount):
    store.confirm_handover(D(day), amount, None, 'бухгалтер')


def october_to_c01(store):
    """Учёт с 02.10 (2 000 000), касса за 04.10 — 490 000: конец 05.10 = 2 490 000.
    06.10 — снимок C-01: касса за 05.10 = 11 014 000 подтверждена, сантехнические
    работы 200 000, кассир выдал Шоху 5 000 000 прямо из кассы."""
    for day in (2, 3, 4, 5):
        confirmed(store, day, '490000' if day == 5 else '0')
    apply_october_opening(store)
    confirmed(store, 6, '11014000')
    store.add_expense(D(6), 'admin_other', 'сантехнические работы', '200000')
    give_shokh(store, D(6), '5000000')


def flow(store, day):
    return store.daily_summary(D(day), None)['day_flow']


# ── Б-09: числа дэшборда C-01 ──────────────────────────────────────────────

def test_c01_cash_is_opening_plus_confirmed_receipts_minus_outflows(store):
    october_to_c01(store)
    day = flow(store, 6)
    assert (day['opening'], day['handover_counted'], day['handover_status']) == ('2490000', '11014000', 'confirmed')
    assert (day['salary'], day['monthly'], day['shoh'], day['other'], day['transfers']) == ('0', '0', '0', '200000', '0')
    # 2 490 000 + 11 014 000 − 0 − 0 − 200 000 = 13 304 000. Выдача Шоху из кассы
    # кассира уже вычтена из передачи: в расходах бухгалтера её нет.
    assert day['outflows'] == '200000' and day['closing'] == '13304000'
    assert [give['amount'] for give in shokh_gives(store, D(6))] == ['5000000']
    assert store.daily_summary(D(6), None)['cash_balance'] == Decimal('13304000')


# ── Б-08: доп. зарплата временному персоналу ───────────────────────────────

def test_extra_salary_item_is_a_salary_but_not_a_monthly_one():
    salary = next(items for code, _, items in GROUPS if code == 'salary')
    assert (EXTRA_SALARY_ITEM, 'Доп. зарплата и временный персонал') in salary
    assert ITEMS[EXTRA_SALARY_ITEM][0] == 'salary'
    assert not is_monthly_salary(EXTRA_SALARY_ITEM)
    assert is_monthly_salary('salary_monthly')


def test_extra_salary_goes_to_salaries_once_and_not_to_other_or_monthly_plan(store):
    october_to_c01(store)
    confirmed(store, 7, '1000000')
    # Долг по смене есть, а временную хостес всё равно можно выдать: общий
    # «ЗП персонал» такой долг не гасит, а доп. зарплата его и не трогает.
    with closing(store._open()) as connection, connection:
        connection.execute(
            'INSERT INTO accountant_accruals (work_day, employee_id, employee_name, group_name, '
            'attendance_status, rate, amount) VALUES (?, ?, ?, ?, ?, ?, ?)',
            (D(6).isoformat(), 1, 'Баходиров Ихтиер', 'Управление', 'on_time', '360000', '360000'))
    with pytest.raises(LedgerError, match='общий расход не погашает долг'):
        store.add_expense(D(7), 'salary_staff', 'смена', '100000')
    store.add_expense(D(7), EXTRA_SALARY_ITEM, 'хостес Карамат · временная', '150000')
    day = flow(store, 7)
    assert (day['salary'], day['monthly'], day['other'], day['outflows']) == ('150000', '0', '0', '150000')
    assert day['closing'] == str(Decimal('13304000') + Decimal('1000000') - Decimal('150000'))
    # План окладов месяца она не уменьшает.
    assert store.reserves(D(7))['monthly']['paid'] == '0'


def test_day_excel_counts_extra_salary_with_salaries(tmp_path):
    app = create_app(Settings(data_dir=tmp_path), expense_db_path=tmp_path / 'cashier.sqlite3',
                     accountant_db_path=tmp_path / 'accountant.sqlite3',
                     director_db_path=tmp_path / 'director.sqlite3', founder_db_path=tmp_path / 'founder.sqlite3')
    day = date(2026, 9, 16)
    with TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)) as c:
        finance = c.app.state.accountant_finance
        finance.record_handover(day, Decimal('5000000'))
        finance.set_cash_opening(day, '1000000', 'Пересчёт')
        catalog = c.get('/api/accountant/expenses/catalog').json()['groups']
        salary = next(group for group in catalog if group['code'] == 'salary')
        assert {'code': EXTRA_SALARY_ITEM, 'label': 'Доп. зарплата и временный персонал'} in salary['items']
        for code, note, amount in ((EXTRA_SALARY_ITEM, 'временные сотрудники', '300000'),
                                   ('admin_other', 'канцтовары', '240000'),
                                   ('distribution_dividends', 'учредителю', '1000000')):
            response = c.post('/api/accountant/expenses', json={'date': day.isoformat(), 'item_code': code,
                                                                 'note': note, 'amount': amount, 'paid_amount': None})
            assert response.status_code == 201, response.text
        response = c.get('/api/accountant/day/export', params={'date': day.isoformat()})
        assert response.status_code == 200, response.text
        book = load_workbook(BytesIO(response.content))
        total = {row[0].value: row[1].value for row in book['Итог'].iter_rows(min_row=3) if row[0].value}
        assert total['− Зарплаты сменным'] == 300000
        assert total['− Прочие расходы и сейф'] == 1240000
        assert (total['На начало дня'] + total['+ От кассира'] - total['− Зарплаты сменным']
                - total['− Прочие расходы и сейф']) == total['= Остаток на конец дня'] == 4460000
        kinds = {row[1].value: row[0].value for row in book['Операции'].iter_rows(min_row=4) if row[1].value}
        assert kinds['Доп. зарплата и временный персонал · временные сотрудники'] == 'Зарплата'
        assert kinds['Дивиденды напрямую из кассы (не из сейфа) · учредителю'] == 'Дивиденды'


# ── Б-08: дивиденды своей строкой, сданный отчёт не «меняется» ──────────────

def test_cash_dividends_are_named_apart_but_reports_keep_their_sums(store):
    october_to_c01(store)
    month_closing.submit_day_report(store, D(6), 'Лина')
    confirmed(store, 7, '70000000')
    store.add_expense(D(7), 'distribution_dividends', 'учредителю', '35000000')
    store.add_expense(D(7), 'admin_other', 'наклейки', '950000')
    store.reserve_entry(D(2), 'dividends', 'opening', '0', 'Сейф пуст')
    store.reserve_entry(D(7), 'dividends', 'transfer', '4000000', 'Дивиденды в сейф',
                        cashier_amount=Decimal('70000000'))
    day = flow(store, 7)
    # «Прочие» по-прежнему включают дивиденды из кассы — так сверяются сданные
    # раньше отчёты; экран показывает их своей строкой из `other_dividends`.
    assert (day['other'], day['other_dividends'], day['transfers']) == ('35950000', '35000000', '4000000')
    assert day['closing'] == str(Decimal('13304000') + Decimal('70000000') - Decimal('39950000'))
    report = month_closing.day_report(store, D(6), flow_json(store.day_flow(D(6))))
    assert report['changed'] is False and 'other_dividends' not in month_closing.REPORT_KEYS


# ── Б-08: перенос старых строк (scripts/reclassify_expenses.py) ─────────────

def c04_legacy(store):
    """07.10 по снимку C-04: временные и дивиденды записаны «Прочими расходами»."""
    october_to_c01(store)
    confirmed(store, 7, '70000000')
    extra = store.add_expense(D(7), 'admin_other', 'временные сотрудники', '2450000')
    dividends = store.add_expense(D(7), 'admin_other', 'дивиденды', '35000000')
    store.add_expense(D(7), 'proc_other', 'рыба сазан', '800000')
    store.add_expense(D(7), 'distribution_dividends', 'уже верно', '1000')
    return extra, dividends


def test_candidates_are_found_by_text_and_moving_keeps_amount_day_and_history(store):
    extra, dividends = c04_legacy(store)
    before = flow(store, 7)
    found = candidates(store.db)
    assert [(row['id'], row['day'], row['amount'], row['item_code'], row['target'], row['blocked'])
            for row in found] == [(extra, '2026-10-07', '2450000', 'admin_other', EXTRA_SALARY_ITEM, None),
                                  (dividends, '2026-10-07', '35000000', 'admin_other', 'distribution_dividends', None)]

    moved = reclassify(store.db, extra, EXTRA_SALARY_ITEM)
    assert (moved['day'], moved['amount'], moved['item_code']) == ('2026-10-07', '2450000', EXTRA_SALARY_ITEM)
    assert moved['description'] == 'Доп. зарплата и временный персонал · временные сотрудники'
    reclassify(store.db, dividends, 'distribution_dividends')
    after = flow(store, 7)
    # Касса та же; 2 450 000 ушли из «прочих» в зарплаты, дивиденды — своей строкой.
    assert after['closing'] == before['closing'] and after['outflows'] == before['outflows']
    assert Decimal(after['salary']) - Decimal(before['salary']) == Decimal('2450000')
    assert Decimal(before['other']) - Decimal(after['other']) == Decimal('2450000')
    assert after['other_dividends'] == '35001000'
    audit = [entry for entry in store.audit_entries(entity_type='movement', entity_id=extra)
             if entry['action'] == 'reclassify']
    assert len(audit) == 1
    assert (audit[0]['before']['item_code'], audit[0]['after']['item_code']) == ('admin_other', EXTRA_SALARY_ITEM)
    assert audit[0]['before']['amount'] == audit[0]['after']['amount'] == '2450000'
    assert candidates(store.db) == []
    with pytest.raises(LedgerError, match='уже в этой статье'):
        reclassify(store.db, extra, EXTRA_SALARY_ITEM)
    with pytest.raises(LedgerError, match='только в доп. зарплату или дивиденды'):
        reclassify(store.db, extra, 'admin_other')


def test_closed_month_and_debt_payments_are_not_moved(store):
    sep_29, sep_30 = date(2026, 9, 29), date(2026, 9, 30)
    store.record_handover(sep_29, Decimal('1000'))
    store.set_cash_opening(sep_29, '500', 'Пересчёт')
    old = store.add_expense(sep_29, 'admin_other', 'временные сотрудники', '200')
    store.record_handover(sep_30, Decimal('300'))
    month_closing.close_month(store, '2026-09', 'Лина', D(1))
    confirmed(store, 2, '5000')
    apply_october_opening(store)
    store.record_debt(D(2), 'admin_other', 'временные сотрудники за неделю', '3000', '1000')
    rows = {row['id']: row for row in candidates(store.db)}
    assert rows[old]['blocked'] == 'месяц закрыт'
    paid = next(row for row in rows.values() if row['id'] != old)
    assert paid['blocked'].startswith('это оплата долга')
    with pytest.raises(LedgerError, match='месяц закрыт'):
        reclassify(store.db, old, EXTRA_SALARY_ITEM)
    with pytest.raises(LedgerError, match='оплата долга'):
        reclassify(store.db, paid['id'], EXTRA_SALARY_ITEM)


def test_script_is_a_dry_run_by_default_and_applies_only_named_rows_after_backup(tmp_path, capsys):
    path = tmp_path / 'accountant.sqlite3'
    store = FinanceStore(path)
    extra, dividends = c04_legacy(store)
    database = str(path)

    assert reclassify_expenses.main(['--database', database]) == 0
    out = capsys.readouterr().out
    assert f'№{extra}  07.10.2026      2 450 000 сум  Прочие расходы → Доп. зарплата и временный персонал' in out
    assert f'№{dividends}  07.10.2026     35 000 000 сум' in out and 'Сухой прогон: ничего не изменено' in out
    assert len(candidates(store.db)) == 2

    assert reclassify_expenses.main(['--database', database, '--apply']) == 2
    assert reclassify_expenses.main(['--database', database, '--apply', '--id', str(extra)]) == 2
    assert reclassify_expenses.main(['--database', database, '--apply', '--id', '999999',
                                     '--backup-root', str(tmp_path / 'backups')]) == 2
    assert not (tmp_path / 'backups').exists() and len(candidates(store.db)) == 2
    capsys.readouterr()

    assert reclassify_expenses.main(['--database', database, '--apply', '--id', str(extra),
                                     '--backup-root', str(tmp_path / 'backups')]) == 0
    out = capsys.readouterr().out
    backup, = (tmp_path / 'backups').iterdir()
    assert (backup / 'accountant.sqlite3').is_file()
    assert json.loads((backup / 'manifest.json').read_text(encoding='utf-8'))['databases'][0]['integrity'] == 'ok'
    assert f'№{extra}: Прочие расходы → Доп. зарплата и временный персонал, 2 450 000 сум' in out
    assert 'изменён после сдачи' in out
    # Перенесена только названная строка; дивиденды ждут своей сверки.
    assert [row['id'] for row in candidates(store.db)] == [dividends]
    assert reclassify_expenses.main(['--database', str(tmp_path / 'нет.sqlite3')]) == 1


@pytest.mark.skipif(not POSTGRES_URL, reason='RETRO_TEST_POSTGRES_URL не задан')
def test_script_on_postgres_keeps_a_json_copy_of_movements_before_moving(tmp_path, capsys):
    with postgres_url() as url:
        store = FinanceStore(url)
        extra, _ = c04_legacy(store)
        assert reclassify_expenses.main(['--database', url, '--apply', '--id', str(extra),
                                         '--backup-root', str(tmp_path / 'backups')]) == 0
        backup, = (tmp_path / 'backups').iterdir()
        rows = json.loads((backup / 'accountant_movements.json').read_text(encoding='utf-8'))
        saved = next(row for row in rows if row['id'] == extra)
        # В копии — строка до переноса, в базе — после.
        assert (saved['item_code'], saved['amount']) == ('admin_other', '2450000')
        assert json.loads((backup / 'manifest.json').read_text(encoding='utf-8'))['rows'] == len(rows)
        assert [row['id'] for row in candidates(store.db)] != [] and extra not in [row['id'] for row in candidates(store.db)]
    assert 'перенесено' in capsys.readouterr().out
