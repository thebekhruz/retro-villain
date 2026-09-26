from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from fastapi.testclient import TestClient

from retro.app import create_app
from retro.config import Settings
from retro.modules.accountant.attendance import Entrance, export_entrances, is_late


def test_ten_am_boundary_uses_tashkent_time():
    tz = ZoneInfo('Asia/Tashkent')
    assert not is_late(datetime(2026, 9, 16, 10, 0, 0, tzinfo=tz))
    assert is_late(datetime(2026, 9, 16, 10, 0, 1, tzinfo=tz))
    assert not is_late(datetime(2026, 9, 16, 4, 59, tzinfo=ZoneInfo('UTC')))


def test_placeholder_export_is_dated_and_contains_no_people():
    with TestClient(create_app(Settings()), base_url='http://127.0.0.1',
                    client=('127.0.0.1', 50000)) as client:
        response = client.get('/api/accountant/entrances/export?date=2026-09-15')
        assert response.status_code == 200
        assert 'Retro-entrances-2026-09-15.xlsx' in response.headers['content-disposition']
        sheet = openpyxl.load_workbook(BytesIO(response.content)).active
        assert sheet['B3'].value.date() == date(2026, 9, 15)
        assert sheet['B4'].value == 0
        assert sheet['D4'].value == 0
        assert [sheet.cell(6, col).value for col in range(1, 5)] == ['№', 'Человек', 'Время входа', 'Статус']
        assert 'не подключены' in sheet['A7'].value
        assert client.get('/api/accountant/entrances/export?date=2099-01-01').status_code == 422


def test_accountant_page_owns_hikvision_preview_and_cashier_links_to_it():
    with TestClient(create_app(Settings()), base_url='http://127.0.0.1',
                    client=('127.0.0.1', 50000)) as client:
        accountant = client.get('/accountant')
        cashier = client.get('/')
        modules = client.get('/api/config').json()['modules']
    assert accountant.status_code == 200
    assert 'ДЕМО' not in accountant.text
    assert 'id="attendance-banner"' in accountant.text
    assert 'id="attendance-chip"' in accountant.text
    assert 'Опоздавшие сотрудники' in accountant.text
    assert 'id="late-details"' in accountant.text
    assert 'href="/accountant/employees"' in accountant.text
    assert 'Скачать опоздавших' in accountant.text
    assert 'name="item_code"' in accountant.text
    assert 'Подтвердить получение' not in accountant.text
    assert 'id="accountant-date"' in accountant.text
    assert 'id="entrances-download"' in accountant.text
    assert 'Зарплата к выплате' not in cashier.text
    # T-383: в разметке кассы чужих модулей нет — пункт «Бухгалтер» дорисует
    # меню, только если он открыт этой учётной записи (здесь — полный доступ).
    assert 'href="/accountant"' not in cashier.text
    assert any(item['id'] == 'accountant' and item['available'] for item in modules)
    assert 'Планирование смен' not in accountant.text
    assert 'scenario-groups' not in accountant.text


def test_attendance_frontend_has_real_source_and_unavailable_states():
    root = Path(__file__).parents[1]
    accountant_js = (root / 'retro/static/accountant.js').read_text(encoding='utf-8')
    employees_js = (root / 'retro/static/employees.js').read_text(encoding='utf-8')
    director_html = (root / 'retro/static/director.html').read_text(encoding='utf-8')

    for source in (accountant_js, employees_js):
        assert 'not_configured' in source
        assert '-DEMO.xlsx' not in source
        assert 'Проходы демонстрационные' not in source
    assert 'unavailable' in employees_js
    assert "unavailable: 'Нет данных'" in employees_js
    assert 'arrivalText(row.first_entry)' in employees_js
    assert "timeZone: 'Asia/Tashkent'" in employees_js
    assert 'демо Hikvision' not in director_html


def test_future_rows_mark_late_and_keep_names_as_text():
    tz = ZoneInfo('Asia/Tashkent')
    entries = (Entrance('Азиза', datetime(2026, 9, 15, 10, 0, tzinfo=tz)),
               Entrance('=HYPERLINK("evil")', datetime(2026, 9, 15, 10, 1, tzinfo=tz)))
    sheet = openpyxl.load_workbook(BytesIO(export_entrances(date(2026, 9, 15), entries))).active
    assert sheet['B4'].value == 2
    assert sheet['D4'].value == 1
    assert sheet['D7'].value == 'Вовремя'
    assert sheet['D8'].value == 'Опоздал'
    assert sheet['B8'].data_type == 's'
    assert sheet['B8'].value.startswith('=HYPERLINK')
    assert sheet['A8'].fill.fgColor.rgb.endswith('FCE6E9')
