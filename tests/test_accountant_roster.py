from decimal import Decimal

import pytest
from openpyxl import Workbook

from retro.integrations.hikvision import HikvisionPerson
from retro.modules.accountant.roster import RosterStore


def make_roster(path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'ЗП'
    sheet.append(['Номер', 'ФИО', 'Должность', 'Ставка'])
    for row, name, role, rate in [
        (5, 'Тест Повар', 'повар миллий', 250000),
        (6, 'Тест Уборка', 'техперсонал', None),
        (7, None, 'повар', 120000),
        (8, 'Тест Официант', 'официант', 170000),
    ]:
        sheet.cell(row, 1, row - 4)
        sheet.cell(row, 2, name)
        sheet.cell(row, 3, role)
        sheet.cell(row, 4, rate)
    sheet['D9'] = '=SUM(D5:D8)'
    book.save(path)


def test_import_only_named_people_preserves_blank_rates_and_source_rows(tmp_path):
    source = tmp_path / 'roster.xlsx'
    make_roster(source)
    store = RosterStore(tmp_path / 'demo.sqlite3')

    assert store.import_xlsx(source) == {'imported': 3, 'existing': 0}
    people = store.list()
    assert [(p.source_row, p.group_name, p.rate) for p in people] == [
        (5, 'Кухня', Decimal('250000')),
        (6, 'Уборка', None),
        (8, 'Обслуживание зала', Decimal('170000')),
    ]
    assert all(p.hikvision_id is None for p in people)


def test_reimport_does_not_overwrite_corrected_rate(tmp_path):
    source = tmp_path / 'roster.xlsx'
    make_roster(source)
    store = RosterStore(tmp_path / 'demo.sqlite3')
    store.import_xlsx(source)
    person = store.list()[0]
    store.update(person.id, rate='260000', group_name='Кухня', reason='Исправлена ставка')

    assert store.import_xlsx(source) == {'imported': 0, 'existing': 3}
    assert RosterStore(tmp_path / 'demo.sqlite3').list()[0].rate == Decimal('260000')


def test_replace_rejects_reusing_identity_for_another_person(tmp_path):
    source = tmp_path / 'roster.xlsx'
    make_roster(source)
    store = RosterStore(tmp_path / 'demo.sqlite3')
    store.import_xlsx(source)
    store.link_hikvision_people((
        HikvisionPerson('100', 'Повар Тест'),
    ))

    book = Workbook()
    sheet = book.active
    sheet.title = 'ЗП'
    sheet.cell(5, 1, 1)
    sheet.cell(5, 2, 'Другой Сотрудник')
    sheet.cell(5, 3, 'повар')
    sheet.cell(5, 4, 250000)
    book.save(source)

    with pytest.raises(ValueError, match='другому сотруднику'):
        store.import_xlsx(source, replace=True)
    assert store.list()[0].hikvision_id == '100'


@pytest.mark.parametrize('field,value', [
    ('card', 'bad'), ('cash', 'NaN'), ('advances', 'Infinity'),
    ('remaining', '-1'), ('salary', '1.001'),
])
def test_monthly_money_fields_reject_invalid_values(tmp_path, field, value):
    store = RosterStore(tmp_path / 'accountant.sqlite3')
    payload = dict(name='Сотрудник', role='Роль', salary='100', schedule='',
                   card='0', cash='0', advances='0', remaining='0')
    payload[field] = value

    with pytest.raises(ValueError):
        store.add_monthly(**payload)

    assert store.list_monthly() == []


def test_deleting_last_monthly_employee_keeps_roster_empty(tmp_path):
    store = RosterStore(tmp_path / 'accountant.sqlite3')
    employee = store.add_monthly(name='Сотрудник', role='Роль', salary='100')

    store.delete_monthly(employee.id)

    assert store.list_monthly() == []
