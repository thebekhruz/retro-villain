"""«Финансы дня» по макету ui_solid1: отметка без Hikvision по дням, оклад по
человеку, перечисления поставщикам, время записей и выгрузки Excel.

Новые запросы (`INSERT OR IGNORE`, `LIKE` с «%» в параметре, новая таблица)
гоняются и на Postgres, если задан RETRO_TEST_POSTGRES_URL, — как в CI.
"""

import os
import re
import sqlite3
import threading
import time
from contextlib import closing, contextmanager
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from retro.app import create_app
from retro.config import Settings
from retro.integrations.hikvision import HikvisionEvent, HikvisionPerson
from retro.modules.accountant.ledger import LedgerError, lock_day
from retro.modules.accountant.payroll import draft_payroll
from retro.modules.cashier.service import TZ, today_tashkent

DAY = date(2026, 9, 16)
BEFORE = DAY - timedelta(days=1)
POSTGRES_URL = os.getenv('RETRO_TEST_POSTGRES_URL', '')
TASHKENT_STAMP = re.compile(r'^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}\+05:00$')


def build_app(tmp_path, **settings):
    return create_app(Settings(data_dir=tmp_path, **settings),
                      expense_db_path=tmp_path / 'cashier.sqlite3',
                      accountant_db_path=tmp_path / 'accountant.sqlite3',
                      director_db_path=tmp_path / 'director.sqlite3',
                      founder_db_path=tmp_path / 'founder.sqlite3')


def client(tmp_path, **settings):
    return TestClient(build_app(tmp_path, **settings), base_url='http://127.0.0.1',
                      client=('127.0.0.1', 50000))


@contextmanager
def postgres_client(tmp_path):
    """Приложение целиком на Postgres, в своей схеме: таблицы других тестов не мешают."""
    import psycopg
    from psycopg import sql
    schema = 'parity_' + uuid4().hex
    admin = psycopg.connect(POSTGRES_URL, autocommit=True)
    admin.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(schema)))
    parts = urlsplit(POSTGRES_URL)
    query = dict(parse_qsl(parts.query))
    query['options'] = '-csearch_path=' + schema
    url = urlunsplit(parts._replace(query=urlencode(query)))
    try:
        app = create_app(Settings(database_url=url, data_dir=tmp_path))
        with TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)) as c:
            yield c
    finally:
        admin.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(schema)))
        admin.close()


@pytest.fixture(params=['sqlite', 'postgres'])
def any_db(request, tmp_path):
    if request.param == 'postgres':
        if not POSTGRES_URL:
            pytest.skip('RETRO_TEST_POSTGRES_URL не задан')
        with postgres_client(tmp_path) as c:
            yield c
    else:
        with client(tmp_path) as c:
            yield c


def cash(c, day=DAY, handover='5000000', opening='1000000'):
    finance = c.app.state.accountant_finance
    finance.record_handover(day, Decimal(handover))
    if opening is not None:
        finance.set_cash_opening(day, opening, 'Пересчёт')


def set_manual_since(c, employee_id, value):
    roster = c.app.state.accountant_roster
    with closing(roster._open()) as connection, connection:
        connection.execute('UPDATE accountant_employees SET manual_since = ? WHERE id = ?',
                           (value.isoformat() if value else None, employee_id))


def manual_person(c, name='Гульшан', role='техперсонал', rate='130000', group='Уборка',
                  since=date(2026, 9, 1)):
    """Сотрудник без Hikvision, которого отмечают вручную с `since`.

    Флаг включается «сегодня», а тестам нужны прошлые дни — дату сдвигаем.
    """
    roster = c.app.state.accountant_roster
    person = roster.add(name=name, role=role, rate=rate, group_name=group)
    roster.set_manual_attendance(person.id, True)
    set_manual_since(c, person.id, since)
    return next(item for item in roster.list() if item.id == person.id)


def staff_row(c, day, employee_id=None):
    rows = c.get('/api/accountant/staff', params={'date': day.isoformat()}).json()['employees']
    row = rows[0] if employee_id is None else next(r for r in rows if r['employee_id'] == employee_id)
    return row['status'], row['payable']


def mark(c, employee_id, day, present):
    return c.post('/api/accountant/manual-attendance', json={
        'date': day.isoformat(), 'employee_id': employee_id, 'present': present})


def day_json(c, day=DAY):
    response = c.get('/api/accountant/day', params={'date': day.isoformat()})
    assert response.status_code == 200, response.text
    return response.json()


def month_json(c, month='2026-09'):
    response = c.get('/api/accountant/payroll/month', params={'month': month})
    assert response.status_code == 200, response.text
    return response.json()


def pay_monthly(c, person_id, amount, day=DAY, **headers):
    return c.post('/api/accountant/monthly-payments', headers=headers,
                  json={'date': day.isoformat(), 'employee_id': person_id, 'amount': amount})


def movement_row(c, movement_id):
    finance = c.app.state.accountant_finance
    with closing(finance._open()) as connection:
        return finance._row_dict(connection, 'accountant_movements', movement_id)


def workbook(response):
    assert response.status_code == 200, response.text
    return load_workbook(BytesIO(response.content))


def rows_by_name(sheet, first_row=1):
    return {sheet.cell(row, 1).value: [sheet.cell(row, col).value for col in range(2, sheet.max_column + 1)]
            for row in range(first_row, sheet.max_row + 1) if sheet.cell(row, 1).value}


# ── Отметка «был / не был» без Hikvision ───────────────────────────────────

def test_manual_employee_is_present_by_default_and_absence_is_per_day(tmp_path):
    with client(tmp_path) as c:
        person = manual_person(c)
        row = c.get('/api/accountant/staff', params={'date': DAY.isoformat()}).json()['employees'][0]
        assert (row['status'], row['payable']) == ('manual_present', '130000')

        marked = c.post('/api/accountant/manual-attendance', json={
            'date': DAY.isoformat(), 'employee_id': person.id, 'present': False})
        assert marked.status_code == 200
        row = c.get('/api/accountant/staff', params={'date': DAY.isoformat()}).json()['employees'][0]
        assert (row['status'], row['payable']) == ('manual_absent', '0')
        # Отметка живёт один день: соседний день по-прежнему «был».
        other = c.get('/api/accountant/staff', params={'date': BEFORE.isoformat()}).json()['employees'][0]
        assert other['status'] == 'manual_present'

        # Смену с такими сотрудниками можно подтвердить без разовых разрешений,
        # и отсутствие начисляется нулём, а не ставкой.
        confirmed = c.post('/api/accountant/payroll/confirm',
                           json={'date': DAY.isoformat(), 'approver': 'Любовь'})
        assert confirmed.status_code == 200 and confirmed.json()['total'] == '0'
        accrual, = c.app.state.accountant_finance.accruals(DAY)
        assert (accrual['status'], accrual['amount']) == ('manual_absent', '0')

        # Подтверждённый день не переписать: отметка и деньги остаются как были.
        again = c.post('/api/accountant/manual-attendance', json={
            'date': DAY.isoformat(), 'employee_id': person.id, 'present': True})
        assert again.status_code == 409
        day = day_json(c)
        assert day['employees'][0]['status'] == 'manual_absent'
        assert day['employees'][0]['payable'] == '0'
        assert day['payroll']['manual_absent_count'] == 1


def test_hikvision_employee_cannot_be_marked_by_hand(tmp_path):
    with client(tmp_path) as c:
        person = c.app.state.accountant_roster.add(name='Жасур', role='официант', rate='180000',
                                                   group_name='Обслуживание зала')
        response = c.post('/api/accountant/manual-attendance', json={
            'date': DAY.isoformat(), 'employee_id': person.id, 'present': False})
        assert response.status_code == 422
        missing = c.post('/api/accountant/manual-attendance', json={
            'date': DAY.isoformat(), 'employee_id': 999, 'present': False})
        assert missing.status_code == 404


def test_manual_attendance_is_an_accountant_action(tmp_path):
    users = {'buh': ('pw', 'accountant'), 'kassa': ('pw', 'cashier'), 'shokh': ('pw', 'shokh')}
    with client(tmp_path, dashboard_panel_users=users) as c:
        person = manual_person(c)
        body = {'date': DAY.isoformat(), 'employee_id': person.id, 'present': False}
        for other in ('kassa', 'shokh'):
            c.post('/api/session', json={'username': other, 'password': 'pw'})
            assert c.post('/api/accountant/manual-attendance', json=body).status_code == 403
            c.post('/api/session/logout')
        c.post('/api/session', json={'username': 'buh', 'password': 'pw'})
        assert c.post('/api/accountant/manual-attendance', json=body).status_code == 200
        store = c.app.state.attendance_store
        with closing(store._open()) as connection:
            approver = connection.execute('SELECT approver FROM hikvision_manual_absences').fetchone()[0]
        assert approver == 'buh'


def test_ai_attendance_tools_can_filter_manual_statuses(tmp_path):
    from retro.modules.director.tools import DirectorChatTools
    from retro.modules.founder.tools import FounderChatTools
    with client(tmp_path) as c:
        person = manual_person(c)
        c.post('/api/accountant/manual-attendance', json={
            'date': DAY.isoformat(), 'employee_id': person.id, 'present': False})
        for tools in (FounderChatTools(c.app), DirectorChatTools(c.app)):
            absent = tools._attendance({'date': DAY.isoformat(), 'status': 'manual_absent'})
            assert [row['name'] for row in absent['employees']] == ['Гульшан']
            assert absent['counts']['manual_absent'] == 1 and absent['counts']['arrived'] == 0
            present = tools._attendance({'date': BEFORE.isoformat(), 'status': 'arrived'})
            assert [row['status'] for row in present['employees']] == ['manual_present']


# ── Оклады частями ─────────────────────────────────────────────────────────

def test_monthly_salary_part_is_linked_to_the_person_and_reaches_the_month_sheet(tmp_path):
    with client(tmp_path) as c:
        cash(c)
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        paid = pay_monthly(c, person.id, '2000000')
        assert paid.status_code == 201

        day = day_json(c)
        assert day['monthly_payments']['paid_by_employee'] == {str(person.id): '2000000'}
        today, = day['monthly_payments']['today']
        assert (today['name'], today['id']) == ('Азиз', paid.json()['id'])
        assert TASHKENT_STAMP.match(today['created_at'])
        # Это обычный расход кассы: он в журнале и уменьшает остаток.
        journal = [row for row in day['ledger']['movements'] if row['item_code'] == 'salary_monthly']
        assert journal[0]['description'] == 'Месячная заработная плата · Азиз'
        assert Decimal(day['ledger']['cash_balance']) == Decimal('4000000')
        # План окладов месяца видит ту же выплату — один раз.
        assert day['reserves']['monthly']['paid'] == '2000000'

        month = month_json(c)
        assert month['monthly_cells'] == {str(person.id): {DAY.isoformat(): '2000000'}}
        assert month['monthly_cell_ops'] == {
            str(person.id): {DAY.isoformat(): [{'id': paid.json()['id'], 'amount': '2000000'}]}}
        assert month['monthly_paid'] == '2000000'

        # Удалили строку из журнала — выплата исчезла и из ведомости.
        c.delete(f"/api/accountant/operations/movement/{journal[0]['id']}",
                 params={'date': DAY.isoformat()})
        month = month_json(c)
        assert month['monthly_cells'] == {} and month['monthly_cell_ops'] == {}


def test_unknown_monthly_employee_is_rejected(tmp_path):
    with client(tmp_path) as c:
        response = pay_monthly(c, 999, '1000')
        assert response.status_code == 404


def test_monthly_payment_needs_cashier_data_and_enough_cash(tmp_path):
    with client(tmp_path) as c:
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        # Без передачи кассира за день выдавать не из чего.
        assert pay_monthly(c, person.id, '1000000').status_code == 409
        cash(c)  # 1 000 000 + 5 000 000
        assert pay_monthly(c, person.id, '6000001').status_code == 422
        assert pay_monthly(c, person.id, '0').status_code == 422
        assert pay_monthly(c, person.id, '6000000').status_code == 201
        assert pay_monthly(c, person.id, '1').status_code == 422
        assert Decimal(day_json(c)['ledger']['cash_balance']) == 0


def test_monthly_payment_retry_with_the_same_key_pays_once(tmp_path):
    """Повтор после потерянного ответа не должен выдать оклад второй раз."""
    with client(tmp_path) as c:
        cash(c)
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        key = str(uuid4())
        first = pay_monthly(c, person.id, '2000000', **{'Idempotency-Key': key})
        second = pay_monthly(c, person.id, '2000000', **{'Idempotency-Key': key})
        assert first.status_code == second.status_code == 201
        assert second.json() == first.json()
        assert pay_monthly(c, person.id, '3000000', **{'Idempotency-Key': key}).status_code == 409
        assert day_json(c)['monthly_payments']['paid_by_employee'] == {str(person.id): '2000000'}


def test_monthly_cell_edit_on_a_past_day_keeps_the_person(tmp_path):
    """Ячейка ведомости правится PUT по id движения: сумма меняется, человек — нет."""
    with client(tmp_path) as c:
        cash(c, BEFORE)
        cash(c, DAY, handover='1000000', opening=None)
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        payment = pay_monthly(c, person.id, '2000000', day=BEFORE).json()['id']
        reference = movement_row(c, payment)['reference']

        def edit(**body):
            return c.put(f'/api/accountant/operations/movement/{payment}', json={
                'date': BEFORE.isoformat(), 'item_code': 'salary_monthly', **body})

        # Уменьшили — остаток вернулся в кассу, выплата осталась за Азизом.
        assert edit(amount='1500000').status_code == 200
        row = movement_row(c, payment)
        assert (row['amount'], row['reference'], row['description']) == (
            '1500000', reference, 'Месячная заработная плата · Азиз')
        assert month_json(c)['monthly_cells'] == {str(person.id): {BEFORE.isoformat(): '1500000'}}
        assert Decimal(day_json(c, BEFORE)['ledger']['cash_balance']) == Decimal('4500000')

        # Подпись с именем не переписывается пояснением: она обязана совпадать со ссылкой.
        assert edit(amount='1500000', note='Дилноза').status_code == 200
        assert movement_row(c, payment)['description'] == 'Месячная заработная плата · Азиз'

        # Другая статья тихо вычеркнула бы выплату из выданного Азизу — нельзя.
        moved = edit(item_code='admin_other', note='Канцтовары', amount='1500000')
        assert moved.status_code == 422 and 'Удалите выплату' in moved.json()['detail']
        assert edit(amount='0').status_code == 422
        assert movement_row(c, payment)['item_code'] == 'salary_monthly'

        # Каждая правка — в журнале аудита с «до» и «после».
        updates = [entry for entry in c.app.state.accountant_finance.audit_entries(
            entity_type='movement', entity_id=payment) if entry['action'] == 'update']
        assert updates[0]['before']['amount'] == '2000000' and updates[0]['after']['amount'] == '1500000'
        assert all(entry['after']['reference'] == reference for entry in updates)


def test_monthly_cell_edit_still_obeys_the_cash_balance(tmp_path):
    with client(tmp_path) as c:
        cash(c, BEFORE)  # 1 000 000 + 5 000 000
        cash(c, DAY, handover='1000000', opening=None)
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        payment = pay_monthly(c, person.id, '2000000', day=BEFORE).json()['id']
        # Следующий день опирается на перенесённый остаток: 4 000 000 + 1 000 000.
        spent = c.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'ops_rent', 'note': 'Аренда', 'amount': '4500000'})
        assert spent.status_code == 201

        def edit(amount):
            return c.put(f'/api/accountant/operations/movement/{payment}', json={
                'date': BEFORE.isoformat(), 'item_code': 'salary_monthly', 'amount': amount})

        # Больше, чем было в кассе в тот день.
        assert edit('7000000').status_code == 422
        # В тот день хватает, но завтрашний расход уйдёт в минус.
        too_much = edit('3000000')
        assert too_much.status_code == 422 and 'последующие' in too_much.json()['detail']
        assert movement_row(c, payment)['amount'] == '2000000'
        assert edit('2400000').status_code == 200
        assert Decimal(day_json(c)['ledger']['cash_balance']) == Decimal('100000')


def test_monthly_cell_can_be_cleared_on_a_past_day(tmp_path):
    with client(tmp_path) as c:
        cash(c, BEFORE)
        cash(c, DAY, handover='1000000', opening=None)
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        first = pay_monthly(c, person.id, '2000000', day=BEFORE).json()['id']
        second = pay_monthly(c, person.id, '500000', day=BEFORE).json()['id']
        ops = month_json(c)['monthly_cell_ops'][str(person.id)][BEFORE.isoformat()]
        assert [op['id'] for op in ops] == [first, second]
        assert c.delete(f'/api/accountant/operations/movement/{first}',
                        params={'date': DAY.isoformat()}).status_code == 422
        for op in ops:
            assert c.delete(f"/api/accountant/operations/movement/{op['id']}",
                            params={'date': BEFORE.isoformat()}).status_code == 204
        month = month_json(c)
        assert month['monthly_cells'] == {} and month['monthly_paid'] == '0'
        assert Decimal(day_json(c, BEFORE)['ledger']['cash_balance']) == Decimal('6000000')


def test_movement_moved_to_another_item_is_no_longer_salary(tmp_path):
    """Даже если статью сменили в обход API, чужая статья окладом не считается."""
    with client(tmp_path) as c:
        cash(c)
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        payment = pay_monthly(c, person.id, '2000000').json()['id']
        finance = c.app.state.accountant_finance
        with closing(finance._open()) as connection, connection:
            connection.execute("UPDATE accountant_movements SET item_code = 'admin_other' WHERE id = ?",
                               (payment,))
        assert finance.monthly_payments(DAY, DAY) == []
        assert day_json(c)['monthly_payments']['paid_by_employee'] == {}


def test_long_monthly_name_still_fits_the_journal(tmp_path):
    with client(tmp_path) as c:
        cash(c)
        name = 'Абдурахмон ' * 14
        person = c.app.state.accountant_roster.add_monthly(name=name, role='менеджер', salary='8000000')
        paid = pay_monthly(c, person.id, '100000')
        assert paid.status_code == 201
        description = movement_row(c, paid.json()['id'])['description']
        assert description.startswith('Месячная заработная плата · Абдурахмон') and len(description) <= 160


# ── Время записей ──────────────────────────────────────────────────────────

def test_every_recorded_journal_row_has_its_tashkent_time(tmp_path):
    with client(tmp_path) as c:
        cash(c, BEFORE)
        cash(c, DAY, handover='5000000', opening=None)
        finance = c.app.state.accountant_finance
        manual_person(c, 'Акмаль', 'охрана', '150000', 'Охрана')
        c.post('/api/accountant/payroll/confirm', json={'date': BEFORE.isoformat(), 'approver': 'Любовь'})
        accrual, = finance.accruals(BEFORE)
        assert c.post('/api/accountant/salary-payments', json={
            'accrual_id': accrual['id'], 'date': DAY.isoformat(), 'amount': '150000'}).status_code == 201
        assert c.post('/api/accountant/procurement', json={
            'date': DAY.isoformat(), 'recipient': 'Шох', 'purpose': 'закуп', 'amount': '700000'}).status_code == 201
        finance.reserve_entry(BEFORE, 'dividends', 'opening', '0', 'Сейф пуст')
        assert c.post('/api/accountant/reserves', json={
            'date': DAY.isoformat(), 'account': 'dividends', 'kind': 'transfer',
            'amount': '300000', 'note': 'Дивиденды'}).status_code == 201
        assert c.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'utilities', 'note': 'Свет', 'amount': '400000',
            'paid_amount': '100000'}).status_code == 201
        movements = day_json(c)['ledger']['movements']
        recorded = [row for row in movements if row['id'] is not None]
        assert {row['type'] for row in recorded} == {
            'salary_payment', 'procurement_advance', 'reserve_transfer', 'other_expense'}
        assert all(TASHKENT_STAMP.match(row['created_at']) for row in recorded), recorded
        # «Касса за день» и «Остаток на начало» — расчёт, у них нет момента записи.
        synthetic = [row for row in movements if row['id'] is None]
        assert {row['type'] for row in synthetic} == {'auto_cashier', 'opening'}
        assert all(row['created_at'] is None for row in synthetic)


@pytest.mark.skipif(not hasattr(time, 'tzset'), reason='нужен time.tzset')
def test_old_rows_without_zone_are_read_as_server_time(tmp_path, monkeypatch):
    """Старые строки писались без пояса во времени сервера (UTC на Railway)."""
    from retro.modules.accountant.ledger import local_timestamp
    monkeypatch.setenv('TZ', 'UTC')
    time.tzset()
    try:
        assert local_timestamp('2026-09-16T03:40:00.123456') == '2026-09-16T08:40:00+05:00'
        assert local_timestamp('2026-09-16T03:40:00+00:00') == '2026-09-16T08:40:00+05:00'
        assert local_timestamp('2026-09-16T08:40:00+05:00') == '2026-09-16T08:40:00+05:00'
        assert local_timestamp('synthetic') is None and local_timestamp(None) is None
        with client(tmp_path) as c:
            cash(c)
            finance = c.app.state.accountant_finance
            with closing(finance._open()) as connection, connection:
                connection.execute(
                    'INSERT INTO accountant_movements (day, kind, description, amount, item_code, created_at) '
                    "VALUES (?, 'other_expense', 'Старая запись', '1000', 'admin_other', ?)",
                    (DAY.isoformat(), '2026-09-16T03:40:00.500000'))
            finance.add_expense(DAY, 'admin_other', 'Новая запись', '1000', cashier_amount=Decimal('5000000'))
            rows = {row['description']: row for row in day_json(c)['ledger']['movements']}
            assert rows['Старая запись']['created_at'] == '2026-09-16T08:40:00+05:00'
            # Новые записи хранятся с поясом и от пояса сервера не зависят.
            with closing(finance._open()) as connection:
                stored = connection.execute(
                    "SELECT created_at FROM accountant_movements WHERE description LIKE 'Прочие расходы%'"
                ).fetchone()[0]
            assert stored.endswith('+05:00')
    finally:
        monkeypatch.undo()
        time.tzset()


# ── Перечисления поставщикам ───────────────────────────────────────────────

def add_transfer(c, key=None, **changes):
    body = {'date': DAY.isoformat(), 'supplier': 'ООО «Мясной двор»', 'item': 'Говядина, 20 кг',
            'point': 'Базар', 'amount': '1850000'}
    body.update(changes)
    headers = {'Idempotency-Key': key} if key else {}
    return c.post('/api/accountant/supplier-transfers', json=body, headers=headers)


def test_supplier_transfer_is_visible_but_moves_no_cash(tmp_path):
    with client(tmp_path) as c:
        cash(c)
        c.app.state.accountant_finance.reserve_entry(DAY, 'shoh', 'opening', '420000', 'Остаток у Шоха')
        before_day = day_json(c)
        before_home = c.get('/api/shokh/home', params={'date': DAY.isoformat()}).json()
        assert before_day['procurement_points'] == c.app.state.shokh.points()
        assert before_day['supplier_transfers'] == [] and before_day['supplier_transfers_total'] == '0'

        created = add_transfer(c)
        assert created.status_code == 201
        transfer = created.json()['transfer']
        assert {key: transfer[key] for key in ('day', 'supplier', 'item', 'point', 'amount')} == {
            'day': DAY.isoformat(), 'supplier': 'ООО «Мясной двор»', 'item': 'Говядина, 20 кг',
            'point': 'Базар', 'amount': '1850000'}
        assert transfer['id'] == created.json()['id'] and TASHKENT_STAMP.match(transfer['created_at'])

        after_day = day_json(c)
        assert after_day['supplier_transfers'] == [transfer]
        assert after_day['supplier_transfers_total'] == '1850000'
        # Касса, журнал и подотчёт Шоха не тронуты.
        assert after_day['ledger']['cash_balance'] == before_day['ledger']['cash_balance'] == '6000000'
        assert after_day['ledger']['movements'] == before_day['ledger']['movements']
        assert after_day['reserves'] == before_day['reserves']
        home = c.get('/api/shokh/home', params={'date': DAY.isoformat()}).json()
        assert home['transfers'] == [transfer] and home['transfers_total'] == '1850000'
        assert (home['pocket'], home['accounting_balance']) == (
            before_home['pocket'], before_home['accounting_balance'])
        # Кассовые потоки учредителя и недели его не видят: денег из кассы не ушло.
        flows = c.app.state.accountant_finance.cash_flows_between(DAY, DAY)
        assert [row['type'] for row in flows] == ['handover']
        # Другой день перечисления не видит.
        assert day_json(c, BEFORE)['supplier_transfers'] == []


def test_supplier_transfer_input_is_checked(tmp_path):
    with client(tmp_path) as c:
        assert add_transfer(c, point='Где-то').status_code == 422
        assert add_transfer(c, supplier='  ').status_code == 422
        assert add_transfer(c, amount='0').status_code == 422
        assert add_transfer(c, amount='-5').status_code == 422
        assert add_transfer(c, amount='12.345').status_code == 422
        assert add_transfer(c, date=(date.today() + timedelta(days=3)).isoformat()).status_code == 422
        assert day_json(c)['supplier_transfers'] == []
        empty = add_transfer(c, item='   ')
        assert empty.status_code == 201 and empty.json()['transfer']['item'] == '—'


def test_supplier_transfer_delete_checks_day_and_is_audited(tmp_path):
    with client(tmp_path) as c:
        transfer_id = add_transfer(c).json()['id']
        path = f'/api/accountant/supplier-transfers/{transfer_id}'
        assert c.delete(path, params={'date': BEFORE.isoformat()}).status_code == 422
        assert c.delete('/api/accountant/supplier-transfers/999', params={'date': DAY.isoformat()}).status_code == 404
        assert c.delete(path, params={'date': DAY.isoformat()}).status_code == 204
        assert day_json(c)['supplier_transfers'] == []
        assert c.delete(path, params={'date': DAY.isoformat()}).status_code == 404
        audit = c.app.state.accountant_finance.audit_entries(entity_type='supplier_transfer',
                                                             entity_id=transfer_id)
        assert [entry['action'] for entry in audit] == ['create', 'delete']
        assert audit[1]['before']['amount'] == '1850000' and audit[1]['after'] is None


def test_supplier_transfer_retry_with_the_same_key_records_once(tmp_path):
    with client(tmp_path) as c:
        key = str(uuid4())
        first, second = add_transfer(c, key=key), add_transfer(c, key=key)
        assert first.status_code == second.status_code == 201 and first.json() == second.json()
        assert len(day_json(c)['supplier_transfers']) == 1


def test_shokh_reads_transfers_but_only_the_accountant_records_them(tmp_path):
    users = {'buh': ('pw', 'accountant'), 'shokh': ('pw', 'shokh')}
    with client(tmp_path, dashboard_panel_users=users) as c:
        c.post('/api/session', json={'username': 'buh', 'password': 'pw'})
        assert add_transfer(c).status_code == 201
        c.post('/api/session/logout')
        c.post('/api/session', json={'username': 'shokh', 'password': 'pw'})
        home = c.get('/api/shokh/home', params={'date': DAY.isoformat()})
        assert home.status_code == 200 and home.json()['transfers'][0]['supplier'] == 'ООО «Мясной двор»'
        assert add_transfer(c).status_code == 403
        assert c.get('/api/accountant/day', params={'date': DAY.isoformat()}).status_code == 403


def test_founder_sees_transfers_as_procurement_spending(tmp_path):
    with client(tmp_path) as c:
        cash(c)
        c.post('/api/accountant/expenses', json={'date': DAY.isoformat(), 'item_code': 'proc_coal',
                                                 'note': 'Уголь', 'amount': '300000'})
        add_transfer(c)
        spending = c.get('/api/founder/spending', params={'date': DAY.isoformat()}).json()['expenses']
        categories = {item['label']: item['amount'] for item in spending['categories']}
        assert categories['Закуп · перечисления'] == '1850000.00'
        assert categories['Закуп · напрямую'] == '300000.00'
        assert spending['total'] == '2150000.00'
        # Выплата оклада — один раз, как «Зарплаты · оклады».
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        assert pay_monthly(c, person.id, '1000000').status_code == 201
        spending = c.get('/api/founder/spending', params={'date': DAY.isoformat()}).json()['expenses']
        categories = {item['label']: item['amount'] for item in spending['categories']}
        assert categories['Зарплаты · оклады'] == '1000000.00' and spending['total'] == '3150000.00'
        # Кассовые потоки учредителя (свободные деньги, неделя) перечисление не видят.
        flows = c.app.state.accountant_finance.cash_flows_between(DAY, DAY)
        assert sum(Decimal(row['amount']) for row in flows if row['type'] == 'other_expense') == 1300000
        assert {row['type'] for row in flows} == {'handover', 'other_expense'}
        book = workbook(c.get('/api/founder/export/month', params={'month': '2026-09'}))
        labels = [row[0].value for row in book['Расходы'].iter_rows(min_row=4)]
        assert 'Закуп · перечисления' in labels


# ── Выгрузки ───────────────────────────────────────────────────────────────

def test_day_export_pays_out_yesterdays_confirmed_shift(tmp_path):
    """Выбранная дата — день выплат: лист «Смена» — вчерашняя смена, как на экране."""
    with client(tmp_path) as c:
        finance = c.app.state.accountant_finance
        guard = manual_person(c, 'Акмаль', 'охрана', '150000', 'Охрана')
        cleaner = manual_person(c)
        cash(c, BEFORE, handover='3000000', opening='0')
        cash(c, DAY, handover='4000000', opening=None)
        assert c.post('/api/accountant/payroll/confirm', json={
            'date': BEFORE.isoformat(), 'approver': 'Любовь'}).status_code == 200
        accrual = next(item for item in finance.accruals(DAY) if item['employee_id'] == guard.id)
        assert c.post('/api/accountant/salary-payments', json={
            'accrual_id': accrual['id'], 'date': DAY.isoformat(), 'amount': '100000'}).status_code == 201
        # Отметка за сам день выплат к выдаваемой смене не относится.
        c.post('/api/accountant/manual-attendance', json={
            'date': DAY.isoformat(), 'employee_id': cleaner.id, 'present': False})

        book = workbook(c.get('/api/accountant/day/export', params={'date': DAY.isoformat()}))
        assert book.sheetnames == ['Смена', 'Операции', 'Итог']
        sheet = book['Смена']
        assert 'Смена 15.09.2026 — выдаётся 16.09.2026' in sheet['A2'].value
        assert 'подтверждена' in sheet['A2'].value and 'не подтверждена' not in sheet['A2'].value
        rows = rows_by_name(sheet, 5)
        assert rows['Акмаль'][:6] == ['охрана', '—', 'Был · вручную', 150000, 100000, 50000]
        assert rows['Гульшан'][:6] == ['техперсонал', '—', 'Был · вручную', 130000, 0, 130000]
        assert rows['Итого за смену'][3:6] == [280000, 100000, 180000]
        journal = rows_by_name(book['Операции'], 4)
        assert journal['Зарплата'][0].startswith('ЗП персонал · Акмаль · за 2026-09-15')


def test_day_export_shows_the_draft_of_an_unconfirmed_shift(tmp_path):
    with client(tmp_path) as c:
        manual_person(c, 'Акмаль', 'охрана', '150000', 'Охрана')
        cleaner = manual_person(c)
        c.post('/api/accountant/manual-attendance', json={
            'date': BEFORE.isoformat(), 'employee_id': cleaner.id, 'present': False})
        sheet = workbook(c.get('/api/accountant/day/export', params={'date': DAY.isoformat()}))['Смена']
        assert 'не подтверждена' in sheet['A2'].value
        rows = rows_by_name(sheet, 5)
        assert rows['Акмаль'][:6] == ['охрана', '—', 'Был · вручную', 150000, 0, 150000]
        assert rows['Гульшан'][:6] == ['техперсонал', '—', 'Не был · вручную', 0, 0, 0]


def test_day_export_journal_and_total_match_the_cash_card(tmp_path):
    with client(tmp_path) as c:
        cash(c)
        finance = c.app.state.accountant_finance
        person = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        assert pay_monthly(c, person.id, '2000000').status_code == 201
        finance.give_procurement(DAY, 'Шох', 'закуп за день', '700000', cashier_amount=Decimal('5000000'))
        finance.add_expense(DAY, 'admin_other', 'Канцтовары', '240000', cashier_amount=Decimal('5000000'))
        finance.add_income(DAY, 'income_other', 'Возврат', '50000')
        assert add_transfer(c, supplier='=HYPERLINK("http://example.com","x")').status_code == 201
        ledger = day_json(c)['ledger']

        book = workbook(c.get('/api/accountant/day/export', params={'date': DAY.isoformat()}))
        journal = book['Операции']
        lines = [[cell.value for cell in row] for row in journal.iter_rows(min_row=4, max_col=5)]
        assert ['Оклад', 'Месячная заработная плата · Азиз'] in [line[:2] for line in lines]
        assert ['Выдано Шоху', 'Шох: закуп за день'] in [line[:2] for line in lines]
        income = next(line for line in lines if line[0] == 'Приход')
        assert (income[3], income[4]) == (50000, None)
        cashier = next(line for line in lines if line[0] == 'Касса')
        assert (cashier[3], cashier[4]) == (5000000, None)
        assert next(line for line in lines if line[0] == 'Итого')[3:5] == [5050000, 2940000]
        # Текст, который ввёл человек, остаётся текстом, а не формулой Excel.
        supplier = next(row[1] for row in journal.iter_rows(min_row=4) if str(row[1].value).startswith('=HYPER'))
        assert supplier.data_type == 's'

        total = {row[0].value: row[1].value for row in book['Итог'].iter_rows(min_row=3) if row[0].value}
        assert total['На начало дня'] == 1000000 and total['+ От кассира'] == 5000000
        assert total['− Оклады частями'] == 2000000 and total['− Выдано Шоху на закуп'] == 700000
        assert total['− Прочие расходы и сейф'] == 240000
        assert total['= Остаток на конец дня'] == Decimal(ledger['cash_balance']) == 3110000
        assert (total['На начало дня'] + total['+ От кассира'] + total['+ Прочие поступления']
                - total['− Зарплаты сменным'] - total['− Оклады частями'] - total['− Выдано Шоху на закуп']
                - total['− Прочие расходы и сейф']) == total['= Остаток на конец дня']
        # Перечисление видно, но остаток кассы не уменьшает.
        assert total['Перечислено поставщикам (не из кассы)'] == 1850000


def test_exports_work_on_an_empty_database(tmp_path):
    with client(tmp_path) as c:
        day = workbook(c.get('/api/accountant/day/export', params={'date': DAY.isoformat()}))
        assert day.sheetnames == ['Смена', 'Операции', 'Итог']
        assert 'За смену записей нет.' in [row[0].value for row in day['Смена'].iter_rows(min_row=5)]
        total = {row[0].value: row[1].value for row in day['Итог'].iter_rows(min_row=3) if row[0].value}
        assert total['= Остаток на конец дня'] is None
        month = workbook(c.get('/api/accountant/payroll/month/export', params={'month': '2026-09'}))
        assert month.sheetnames == ['Ведомость']
        assert c.get('/api/accountant/payroll/month/export', params={'month': '2026-13'}).status_code == 422


def test_month_sheet_keeps_payments_of_removed_staff_and_salary_without_a_person(tmp_path):
    with client(tmp_path) as c:
        cash(c, handover='9000000')
        roster = c.app.state.accountant_roster
        aziz = roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
        dilnoza = roster.add_monthly(name='Дилноза', role='менеджер', salary='7000000')
        assert pay_monthly(c, aziz.id, '2000000').status_code == 201
        assert pay_monthly(c, dilnoza.id, '1500000').status_code == 201
        assert c.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'salary_monthly', 'note': 'Аванс без имени',
            'amount': '500000'}).status_code == 201
        roster.delete_monthly(dilnoza.id)
        assert month_json(c)['monthly_paid'] == '4000000'

        sheet = workbook(c.get('/api/accountant/payroll/month/export', params={'month': '2026-09'}))['Ведомость']
        rows = rows_by_name(sheet, 5)
        paid = sheet.max_column - 1  # колонки: …, Начислено, Выдано, Осталось
        column = lambda name: rows[name][paid - 2]
        assert column('Азиз') == 2000000 and rows['Азиз'][-1] == 6000000
        assert column(f'Сотрудник удалён · №{dilnoza.id}') == 1500000
        assert column('Оклады без сотрудника (общий расход)') == 500000
        assert column('Итого оклады') == 4000000
        day_column = 2 + DAY.day - 2  # индекс в списке значений после имени
        assert rows['Азиз'][day_column] == 2000000


# ── Оба диалекта ───────────────────────────────────────────────────────────

def test_new_backend_paths_behave_the_same_on_both_dialects(any_db):
    c = any_db
    cash(c, BEFORE)
    cash(c, DAY, handover='1000000', opening=None)

    # Отметка «не был»: повтор — INSERT OR IGNORE, снятие — DELETE.
    person = manual_person(c)
    body = {'date': BEFORE.isoformat(), 'employee_id': person.id, 'present': False}
    assert c.post('/api/accountant/manual-attendance', json=body).status_code == 200
    assert c.post('/api/accountant/manual-attendance', json=body).status_code == 200
    staff = c.get('/api/accountant/staff', params={'date': BEFORE.isoformat()}).json()
    assert staff['employees'][0]['status'] == 'manual_absent'
    assert c.post('/api/accountant/manual-attendance', json={**body, 'present': True}).status_code == 200
    staff = c.get('/api/accountant/staff', params={'date': BEFORE.isoformat()}).json()
    assert staff['employees'][0]['status'] == 'manual_present'
    assert c.post('/api/accountant/manual-attendance', json=body).status_code == 200

    # Оклад: LIKE с «%» в параметре, правка суммы и удаление по id из ведомости.
    monthly = c.app.state.accountant_roster.add_monthly(name='Азиз', role='менеджер', salary='8000000')
    key = str(uuid4())
    first = pay_monthly(c, monthly.id, '2000000', day=BEFORE, **{'Idempotency-Key': key})
    assert first.status_code == 201
    assert pay_monthly(c, monthly.id, '2000000', day=BEFORE, **{'Idempotency-Key': key}).json() == first.json()
    second = pay_monthly(c, monthly.id, '500000', day=DAY).json()['id']
    ops = month_json(c)['monthly_cell_ops'][str(monthly.id)]
    assert ops == {BEFORE.isoformat(): [{'id': first.json()['id'], 'amount': '2000000'}],
                   DAY.isoformat(): [{'id': second, 'amount': '500000'}]}
    assert c.put(f"/api/accountant/operations/movement/{first.json()['id']}", json={
        'date': BEFORE.isoformat(), 'item_code': 'salary_monthly', 'amount': '1500000'}).status_code == 200
    assert c.delete(f'/api/accountant/operations/movement/{second}',
                    params={'date': DAY.isoformat()}).status_code == 204
    month = month_json(c)
    assert month['monthly_cells'] == {str(monthly.id): {BEFORE.isoformat(): '1500000'}}
    assert movement_row(c, first.json()['id'])['description'] == 'Месячная заработная плата · Азиз'

    # Перечисление: своя таблица с автономером, аудит, удаление.
    transfer = add_transfer(c)
    assert transfer.status_code == 201 and transfer.json()['id']
    day = day_json(c)
    assert [row['id'] for row in day['supplier_transfers']] == [transfer.json()['id']]
    assert all(TASHKENT_STAMP.match(row['created_at'])
               for row in day['ledger']['movements'] if row['id'] is not None)
    assert c.delete(f"/api/accountant/supplier-transfers/{transfer.json()['id']}",
                    params={'date': DAY.isoformat()}).status_code == 204

    # Выгрузки собираются и на этих данных.
    assert workbook(c.get('/api/accountant/day/export', params={'date': DAY.isoformat()})).sheetnames == [
        'Смена', 'Операции', 'Итог']
    assert workbook(c.get('/api/accountant/payroll/month/export', params={'month': '2026-09'})).sheetnames == [
        'Ведомость']


def test_column_lookup_stays_in_its_own_schema():
    """Одноимённая таблица в другой схеме не должна подсовывать свои колонки:
    иначе ALTER новой колонки тихо пропускался, и запросы падали на Postgres."""
    if not POSTGRES_URL:
        pytest.skip('RETRO_TEST_POSTGRES_URL не задан')
    import psycopg
    from psycopg import sql
    from retro.db import Database, table_columns
    other, own = 'parity_' + uuid4().hex, 'parity_' + uuid4().hex
    admin = psycopg.connect(POSTGRES_URL, autocommit=True)
    try:
        for schema, column in ((other, 'manual_attendance'), (own, 'name')):
            admin.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(schema)))
            admin.execute(sql.SQL('CREATE TABLE {}.probe_columns ({} TEXT)').format(
                sql.Identifier(schema), sql.Identifier(column)))
        parts = urlsplit(POSTGRES_URL)
        query = dict(parse_qsl(parts.query), options='-csearch_path=' + own)
        database = Database(urlunsplit(parts._replace(query=urlencode(query))))
        with database.cursor() as connection:
            assert table_columns(connection, 'probe_columns') == {'name'}
    finally:
        for schema in (other, own):
            admin.execute(sql.SQL('DROP SCHEMA IF EXISTS {} CASCADE').format(sql.Identifier(schema)))
        admin.close()


def test_migration_to_postgres_carries_the_new_tables(tmp_path):
    """Перенос в Postgres сверяет строки и деньги — и для новых таблиц тоже."""
    if not POSTGRES_URL:
        pytest.skip('RETRO_TEST_POSTGRES_URL не задан')
    import subprocess
    import sys
    import psycopg
    from psycopg import sql
    from retro.modules.accountant.hikvision import AttendanceStore
    from retro.modules.accountant.ledger import FinanceStore
    source = tmp_path / 'accountant.sqlite3'
    finance = FinanceStore(source)
    finance.record_handover(DAY, Decimal('5000000'))
    finance.set_cash_opening(DAY, '1000000', 'Пересчёт')
    finance.pay_monthly(7, 'Азиз', DAY, '2000000', cashier_amount=Decimal('5000000'))
    finance.add_supplier_transfer(DAY, 'ООО «Мясной двор»', '', 'Базар', '1850000')
    AttendanceStore(source).set_manual_absence(3, DAY, True, 'Любовь')

    schema = 'parity_' + uuid4().hex
    admin = psycopg.connect(POSTGRES_URL, autocommit=True)
    admin.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(schema)))
    parts = urlsplit(POSTGRES_URL)
    target = urlunsplit(parts._replace(query=urlencode(
        dict(parse_qsl(parts.query), options='-csearch_path=' + schema))))
    try:
        done = subprocess.run(
            [sys.executable, 'scripts/migrate_to_postgres.py', '--target', target,
             '--accountant', str(source), '--apply'],
            capture_output=True, text=True, env={**os.environ, 'PYTHONPATH': '.'})
        assert done.returncode == 0, done.stdout + done.stderr
        assert '✓ accountant_supplier_transfers: 1 строк, 1850000 сум' in done.stdout
        assert '✓ hikvision_manual_absences: 1 строк' in done.stdout
        moved = FinanceStore(target)
        assert [(row['supplier'], row['item'], row['amount']) for row in moved.supplier_transfers(DAY)] == [
            ('ООО «Мясной двор»', '—', '1850000')]
        assert [(row['employee_id'], row['amount']) for row in moved.monthly_payments(DAY, DAY)] == [(7, '2000000')]
        assert AttendanceStore(target).manual_absences(DAY) == {3}
        # Счётчик номеров сдвинут: новая запись не сталкивается с перенесённой.
        assert moved.add_supplier_transfer(DAY, 'Хлебозавод', 'Хлеб', 'Базар', '90000') > 1
    finally:
        admin.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(schema)))
        admin.close()


# ── С какого дня «был» по умолчанию, явные отметки, аудит ───────────────────

def test_manual_default_starts_on_the_day_the_flag_was_set(tmp_path):
    today = today_tashkent()
    since = today - timedelta(days=3)
    with client(tmp_path) as c:
        roster = c.app.state.accountant_roster
        person = roster.add(name='Гульшан', role='техперсонал', rate='130000', group_name='Уборка')
        assert roster.set_manual_attendance(person.id, True).manual_since == today
        # Повторное включение дату не сдвигает.
        set_manual_since(c, person.id, since)
        assert roster.set_manual_attendance(person.id, True).manual_since == since

        assert staff_row(c, since) == ('manual_present', '130000')
        assert staff_row(c, today) == ('manual_present', '130000')
        # До включения «был» не подразумевается: за дни, когда человека ещё не
        # отмечали, оплата сама не начисляется.
        earlier = since - timedelta(days=1)
        assert staff_row(c, earlier) == ('manual_absent', '0')
        # «Был» можно поставить явно — и до даты включения тоже.
        assert mark(c, person.id, earlier, True).status_code == 200
        assert staff_row(c, earlier) == ('manual_present', '130000')
        # Явное «не был» сильнее умолчания и после даты включения.
        assert mark(c, person.id, since, False).status_code == 200
        assert staff_row(c, since) == ('manual_absent', '0')
        # «Нет привязки» у ручных не бывает, поэтому смену до даты можно подтвердить.
        assert c.post('/api/accountant/payroll/confirm', json={
            'date': (since - timedelta(days=2)).isoformat(), 'approver': 'Любовь'}).status_code == 200

        # Сняли флаг — дата стёрта; включили снова — отсчёт с сегодняшнего дня.
        off = roster.set_manual_attendance(person.id, False)
        assert (off.manual_attendance, off.manual_since) == (False, None)
        assert roster.set_manual_attendance(person.id, True).manual_since == today
        # Экран директора видит дату включения.
        team = c.get('/api/director/team').json()
        assert team['shift'][0]['manual_since'] == today.isoformat()


def test_before_the_flag_a_real_turnstile_entry_still_counts(tmp_path):
    with client(tmp_path) as c:
        roster = c.app.state.accountant_roster
        person = roster.add(name='Нодира Исмоилова', role='официант', rate='180000',
                            group_name='Обслуживание зала')
        roster.link_hikvision_people((HikvisionPerson('77', 'Нодира Исмоилова'),))
        c.app.state.attendance_store.ingest(HikvisionEvent(
            'retro-main-entry', 'e-1', '77', datetime(2026, 9, 15, 10, 20, tzinfo=TZ)), person.id)
        roster.set_manual_attendance(person.id, True)  # с сегодняшнего дня
        rows = {row['employee_id']: row for row in c.get(
            '/api/accountant/staff', params={'date': BEFORE.isoformat()}).json()['employees']}
        assert (rows[person.id]['status'], rows[person.id]['payable']) == ('late', '180000')
        assert rows[person.id]['first_entry'].startswith('2026-09-15T10:20')
        assert staff_row(c, BEFORE - timedelta(days=1), person.id) == ('manual_absent', '0')


def test_staff_flagged_before_the_date_existed_keep_being_present(tmp_path):
    """Уже отмеченные вручную до обновления: дата пустая — «был» во все дни,
    вчерашняя смена у них выдаётся как раньше; старые строки — отсутствия."""
    path = tmp_path / 'accountant.sqlite3'
    with closing(sqlite3.connect(path)) as connection, connection:
        connection.execute(
            'CREATE TABLE accountant_employees (id INTEGER PRIMARY KEY AUTOINCREMENT, '
            'source_row INTEGER NOT NULL UNIQUE, name TEXT NOT NULL, role TEXT NOT NULL, '
            'group_name TEXT NOT NULL, rate TEXT, hikvision_id TEXT UNIQUE, '
            'manual_attendance INTEGER NOT NULL DEFAULT 0)')
        connection.execute("INSERT INTO accountant_employees (source_row, name, role, group_name, rate, "
                           "manual_attendance) VALUES (1, 'Акмаль', 'охрана', 'Охрана', '150000', 1)")
        connection.execute(
            'CREATE TABLE hikvision_manual_absences (work_day TEXT NOT NULL, employee_id INTEGER NOT NULL, '
            'approver TEXT NOT NULL, created_at TEXT NOT NULL, PRIMARY KEY (work_day, employee_id))')
        connection.execute("INSERT INTO hikvision_manual_absences VALUES "
                           "('2026-09-15', 1, 'Любовь', '2026-09-16T09:00:00+05:00')")
    with client(tmp_path) as c:
        person, = c.app.state.accountant_roster.list()
        assert (person.manual_attendance, person.manual_since) == (True, None)
        assert staff_row(c, DAY) == ('manual_present', '150000')
        assert staff_row(c, date(2026, 1, 5)) == ('manual_present', '150000')
        assert staff_row(c, BEFORE) == ('manual_absent', '0')
        assert c.app.state.attendance_store.manual_marks(BEFORE) == {person.id: False}
        # Снять старое «не был» можно — это запись «был».
        assert mark(c, person.id, BEFORE, True).status_code == 200
        assert staff_row(c, BEFORE) == ('manual_present', '150000')


def test_every_mark_change_is_audited_and_a_repeat_writes_nothing(tmp_path):
    with client(tmp_path) as c:
        person = manual_person(c)
        for present in (False, False, True):
            assert mark(c, person.id, DAY, present).status_code == 200
        audit = c.app.state.accountant_finance.audit_entries(
            entity_type='manual_attendance', entity_id=f'{DAY.isoformat()}:{person.id}')
        assert [entry['action'] for entry in audit] == ['create', 'update']
        assert audit[0]['before'] is None
        assert (audit[0]['after']['present'], audit[0]['after']['approver']) == (False, 'local')
        assert (audit[1]['before']['present'], audit[1]['after']['present']) == (False, True)
        assert all(entry['changed_at'] for entry in audit)


# ── Отметка и подтверждение смены не проскакивают друг мимо друга ──────────

def test_confirm_rejects_marks_changed_after_the_shift_was_computed(any_db):
    c = any_db
    person = manual_person(c)
    finance, roster, attendance = (c.app.state.accountant_finance, c.app.state.accountant_roster,
                                   c.app.state.attendance)
    people = roster.list(DAY)
    snapshot = attendance.snapshot(DAY, people)
    rows = draft_payroll(DAY, people, set(), snapshot.rows)
    assert rows[0].payable == Decimal('130000')
    # «Не был» поставили уже после того, как смена посчитана.
    assert finance.mark_manual_attendance(person.id, DAY, False, 'Любовь') is True
    with pytest.raises(LedgerError, match='уже изменились'):
        finance.confirm_payroll(DAY, rows, 'Любовь', marks=snapshot.marks)
    assert finance.accruals(DAY) == []
    fresh = attendance.snapshot(DAY, people)
    assert finance.confirm_payroll(DAY, draft_payroll(DAY, people, set(), fresh.rows), 'Любовь',
                                   marks=fresh.marks)
    assert [item['amount'] for item in finance.accruals(DAY)] == ['0']


def test_a_mark_cannot_slip_into_a_shift_being_confirmed(any_db):
    """Пока идёт подтверждение смены, отметка ждёт — и потом получает отказ."""
    c = any_db
    person = manual_person(c)
    finance = c.app.state.accountant_finance
    outcome = {}

    def write_mark():
        try:
            outcome['changed'] = finance.mark_manual_attendance(person.id, DAY, False, 'Акмаль')
        except Exception as error:  # отказ проверяем в основном потоке
            outcome['error'] = error

    holder = finance._open()
    worker = threading.Thread(target=write_mark)
    try:
        holder.execute('BEGIN IMMEDIATE')
        lock_day(holder, DAY)
        holder.execute('INSERT INTO accountant_payroll_days (day, approver, confirmed_at) VALUES (?, ?, ?)',
                       (DAY.isoformat(), 'Любовь', '2026-09-17T09:00:00+05:00'))
        worker.start()
        worker.join(0.5)
        assert worker.is_alive(), 'отметка должна ждать, пока подтверждение не закончится'
        holder.commit()
    finally:
        holder.close()
    worker.join(10)
    assert not worker.is_alive()
    assert 'уже подтверждена' in str(outcome.get('error')), outcome
    assert c.app.state.attendance_store.manual_marks(DAY) == {}


# ── «Пришли / не пришли» одинаково везде ───────────────────────────────────

def test_director_counts_manual_marks_like_the_employees_page(tmp_path):
    from retro.modules.director.tools import DirectorChatTools
    from retro.modules.founder.tools import FounderChatTools
    with client(tmp_path) as c:
        manual_person(c, 'Акмаль', 'охрана', '150000', 'Охрана')
        absent = manual_person(c)
        assert mark(c, absent.id, DAY, False).status_code == 200
        attendance = c.get('/api/director/attendance', params={'date': DAY.isoformat()}).json()
        assert (attendance['arrived_count'], attendance['missing_count']) == (1, 1)
        team = c.get('/api/director/team', params={'date': DAY.isoformat()}).json()
        assert team['counts']['missing'] == 1
        for tools in (FounderChatTools(c.app), DirectorChatTools(c.app)):
            result = tools._attendance({'date': DAY.isoformat(), 'status': 'absent'})
            assert [row['name'] for row in result['employees']] == ['Гульшан']
            assert (result['counts']['arrived'], result['counts']['absent']) == (1, 1)


# ── Чистая прибыль учредителя: перечисления — как закуп напрямую ────────────

class FounderPnl:
    async def load_founder_analytics(self, start, end, granularity, directions):
        return {'pnl': {'net_profit': '10000000'}, 'scope_note': ''}


def test_founder_net_profit_subtracts_transfers_like_direct_procurement(tmp_path):
    """Закуп через Шоха — не расход (его накладные уже в себестоимости iiko),
    закуп, оплаченный поставщику напрямую — наличными или переводом, — расход."""
    with client(tmp_path) as c:
        c.app.state.iiko = FounderPnl()
        cash(c)
        finance = c.app.state.accountant_finance
        finance.add_expense(DAY, 'proc_coal', 'Уголь', '300000', cashier_amount=Decimal('5000000'))
        finance.give_procurement(DAY, 'Шох', 'закуп за день', '700000', cashier_amount=Decimal('5000000'))
        assert add_transfer(c).status_code == 201
        data = c.get('/api/founder/analytics', params={
            'start': DAY.isoformat(), 'end': DAY.isoformat()}).json()
        assert data['dashboard_expenses'] == {
            'cashier': '0', 'accountant_other': '2150000', 'accountant_salary': '0',
            'accountant': '2150000', 'total': '2150000'}
        assert data['net_profit_after_dashboard_expenses'] == '7850000'


# ── Касса → Шох и передача кассира в «Финансах дня» ─────────────────────────

def test_day_export_shows_till_gives_and_handover_without_touching_the_balance(any_db):
    from retro.modules.cashier.till import give_shokh
    c = any_db
    finance = c.app.state.accountant_finance
    finance.record_handover(DAY, Decimal('5000000'), source='cashier')
    finance.set_cash_opening(DAY, '1000000', 'Пересчёт')
    finance.reserve_entry(DAY, 'shoh', 'opening', '420000', 'Остаток у Шоха')
    finance.add_expense(DAY, 'admin_other', 'Канцтовары', '240000', cashier_amount=Decimal('5000000'))

    def export():
        book = workbook(c.get('/api/accountant/day/export', params={'date': DAY.isoformat()}))
        total = {row[0].value: row[1].value for row in book['Итог'].iter_rows(min_row=3) if row[0].value}
        return book, total

    _, before = export()
    give = give_shokh(finance, DAY, '300000')
    day = day_json(c)
    book, total = export()
    # Остаток бухгалтера тот же: выдача из кассы до него не доходит.
    assert total['= Остаток на конец дня'] == before['= Остаток на конец дня'] == Decimal(
        day['ledger']['cash_balance']) == 5760000
    assert total['Шоху из кассы · не из остатка'] == 300000
    assert total['Подотчёт Шоха на конец дня'] == 720000
    handed = next(label for label in total if str(label).startswith('Передача кассира'))
    assert handed == f"Передача кассира · получено {day['cashier_handover']['handed_at'][11:16]} · кассир"
    assert total[handed] == 5000000

    lines = [[cell.value for cell in row] for row in book['Операции'].iter_rows(min_row=4, max_col=5)]
    start = next(i for i, line in enumerate(lines) if line[0] == 'Выдано Шоху из кассы — не из остатка бухгалтера')
    assert lines[start + 2] == ['Кассир', 'Выдано Шоху на закуп', give['created_at'][11:16], None, 300000]
    assert lines[start + 3][0] == 'Итого' and lines[start + 3][4] == 300000
    # В приход и расход дня выдача не попала.
    day_total = next(line for line in lines if line[0] == 'Итого')
    assert day_total[3:5] == [5000000, 240000]


def test_founder_month_excel_counts_till_gives_to_shokh_once(tmp_path):
    from retro.modules.cashier.till import give_shokh
    with client(tmp_path) as c:
        cash(c)
        finance = c.app.state.accountant_finance
        finance.give_procurement(DAY, 'Шох', 'закуп за день', '700000', cashier_amount=Decimal('5000000'))
        give_shokh(finance, DAY, '300000')
        cabinet = c.get('/api/founder/spending', params={'date': DAY.isoformat()}).json()
        cabinet_categories = {item['label']: item['amount'] for item in cabinet['expenses']['categories']}
        assert cabinet_categories['Закуп · наличные Шоху'] == '1000000.00'
        book = workbook(c.get('/api/founder/export/month', params={'month': '2026-09'}))
        rows = {row[0].value: row[1].value for row in book['Расходы'].iter_rows(min_row=4) if row[0].value}
        # Один раз, как в кабинете: 700 000 от бухгалтера + 300 000 из кассы.
        assert rows['Закуп · наличные Шоху'] == 1000000
        assert rows['Итого'] == Decimal(cabinet['expenses']['total']) == 1000000
        # Кассовые потоки бухгалтера по дням выдачу из кассы не видят.
        by_day = book['По дням']
        heads = [cell.value for cell in by_day[4]]
        day_row = next(row for row in by_day.iter_rows(min_row=5) if row[0].value.date() == DAY)
        assert day_row[heads.index('Закуп')].value == 700000
