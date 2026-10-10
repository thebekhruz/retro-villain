"""T-434: временные сотрудники у бухгалтера — должность и период работы.

Пример ТЗ 10.10: хостес Карамат, временная, 150 000 за смену, работает с 08.10
по 10.10. В ведомости и списках дня она есть только в эти смены; выплату за
последнюю смену (10.10) делают назавтра — клетка 11.10 открыта. Отметка смены
12.10, выплата за смену 07.10 и доп. выплата за 07.10 отклоняются словами.
Временный без периода живёт как сменный. На SQLite и на Postgres
(RETRO_TEST_POSTGRES_URL).
"""

import os
import sqlite3
from contextlib import closing
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from test_accountant_design_parity import any_db  # noqa: F401 — фикстура

from retro.db import Database, table_columns
from retro.modules.accountant import extra_payouts, routes, salary_day
from retro.modules.accountant.hikvision import AttendanceService, AttendanceStore
from retro.modules.accountant.ledger import FinanceStore, LedgerError
from retro.modules.accountant.payroll import PayrollRow, draft_payroll
from retro.modules.accountant.roster import RosterStore
from retro.modules.accountant.work_period import label, normalize, outside_text
from retro.modules.cashier.service import TZ

TODAY = date(2026, 10, 12)
FROM, TO = date(2026, 10, 8), date(2026, 10, 10)
POSTGRES_URL = os.getenv('RETRO_TEST_POSTGRES_URL', '')
KARAMAT_TEXT = 'Карамат работает с 08.10 по 10.10 — '


def d(day: int) -> date:
    return date(2026, 10, day)


@pytest.fixture(autouse=True)
def today(monkeypatch):
    for module in (salary_day, routes, extra_payouts):
        monkeypatch.setattr(module, 'today_tashkent', lambda: TODAY)


@pytest.fixture(params=['sqlite', 'postgres'])
def database(request, tmp_path):
    if request.param == 'sqlite':
        yield tmp_path / 'accountant.sqlite3'
        return
    if not POSTGRES_URL:
        pytest.skip('RETRO_TEST_POSTGRES_URL не задан')
    import psycopg
    from psycopg import sql
    schema = 'temporary_' + uuid4().hex
    with psycopg.connect(POSTGRES_URL, autocommit=True) as admin:
        admin.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(schema)))
        try:
            parts = urlsplit(POSTGRES_URL)
            query = dict(parse_qsl(parts.query), options='-csearch_path=' + schema)
            yield Database(urlunsplit(parts._replace(query=urlencode(query))))
        finally:
            admin.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(schema)))


@pytest.fixture
def stores(database):
    roster = RosterStore(database)
    finance = FinanceStore(database)
    attendance = AttendanceStore(database)
    for offset in range(11):
        finance.record_handover(d(5) + timedelta(days=offset), Decimal(0))
    finance.set_cash_opening(d(5), '5000000', 'Начало')
    people = dict(
        ihtiyor=roster.add(name='Баходиров Ихтиер', role='Менеджер', rate='360000', group_name='Управление'),
        karamat=roster.add(name='Карамат', role='Хостес', rate='150000', group_name='Встреча гостей',
                           employment_type='temporary', work_from=FROM, work_to=TO, by='buh'),
        guest=roster.add(name='Гулноза', role='Хостес', rate='150000', group_name='Встреча гостей',
                         employment_type='temporary'))
    return roster, finance, attendance, people


def count(store, table):
    with closing(store._open()) as connection:
        return connection.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0]


def names(people):
    return [person.name for person in people]


# ── Период: проверка ввода и подписи ─────────────────────────────────────────

def test_period_words_and_validation():
    assert normalize('temporary', '2026-10-08', '2026-10-10') == ('temporary', FROM, TO)
    assert normalize('temporary', None, '') == ('temporary', None, None)
    assert normalize(None, None, None) == ('shift', None, None)
    assert (label(FROM, TO), label(FROM, None), label(None, TO), label(FROM, FROM), label(None, None)) == (
        '08.10–10.10', 'с 08.10', 'по 10.10', '08.10', None)
    assert outside_text('Карамат', FROM, TO, d(12), 'mark') == KARAMAT_TEXT + 'смену 12.10 отметить нельзя.'
    assert outside_text('Карамат', FROM, None, d(7), 'cell') == \
        'Карамат работает с 08.10 — выплату за смену 07.10 записать нельзя.'
    assert outside_text('Карамат', FROM, FROM, d(9), 'extra') == \
        'Карамат работает только 08.10 — доп. выплату за смену 09.10 записать нельзя.'
    with pytest.raises(ValueError, match=r'«по» \(07.10\) раньше, чем «с» \(08.10\)'):
        normalize('temporary', FROM, d(7))
    with pytest.raises(ValueError, match='только у временного'):
        normalize('shift', FROM, None)
    with pytest.raises(ValueError, match='сменный или временный'):
        normalize('monthly', None, None)
    with pytest.raises(ValueError, match='ГГГГ-ММ-ДД'):
        normalize('temporary', '08.10.2026', None)


# ── Реестр: списки дня только в периоде ──────────────────────────────────────

def test_karamat_is_in_day_lists_only_on_her_shifts(stores):
    roster, _, _, people = stores
    karamat = people['karamat']
    assert (karamat.employment_type, karamat.work_from, karamat.work_to) == ('temporary', FROM, TO)
    assert {key: karamat.json()[key] for key in ('employment_type', 'work_from', 'work_to', 'work_period')} == dict(
        employment_type='temporary', work_from='2026-10-08', work_to='2026-10-10', work_period='08.10–10.10')
    for day in (d(7), d(11), d(12)):
        assert 'Карамат' not in names(roster.list(day))
    for day in (d(8), d(9), d(10)):
        assert 'Карамат' in names(roster.list(day))
    # Весь реестр (без дня) — все; временный без периода — во все дни, как сменный.
    assert names(roster.list()) == ['Баходиров Ихтиер', 'Карамат', 'Гулноза']
    assert all('Гулноза' in names(roster.list(day)) for day in (d(1), d(9), d(12)))
    created = roster.history(karamat.id)[0]
    assert (created['action'], created['details'], created['changed_by']) == ('create', 'Временный · 08.10–10.10', 'buh')
    assert roster.history(people['guest'].id)[0]['details'] == 'Временный'
    assert roster.history(people['ihtiyor'].id)[0]['details'] == ''


def test_deleted_temporary_keeps_her_period_in_past_days(stores):
    roster, _, _, people = stores
    roster.delete(people['karamat'].id)
    assert 'Карамат' not in names(roster.list())
    # Прошлые дни помнят и её, и период: вне периода её нет и после удаления.
    assert 'Карамат' in names(roster.list(d(9)))
    assert 'Карамат' not in names(roster.list(d(7)))


def test_old_roster_gains_period_columns_in_both_tables(tmp_path):
    path = tmp_path / 'accountant.sqlite3'
    with sqlite3.connect(path) as connection:
        connection.execute('''CREATE TABLE accountant_employees (
            id INTEGER PRIMARY KEY AUTOINCREMENT, source_row INTEGER NOT NULL UNIQUE,
            name TEXT NOT NULL, role TEXT NOT NULL, group_name TEXT NOT NULL, rate TEXT,
            hikvision_id TEXT UNIQUE)''')
        connection.execute('''CREATE TABLE accountant_employee_versions (
            employee_id INTEGER NOT NULL, effective_day TEXT NOT NULL,
            source_row INTEGER NOT NULL, name TEXT NOT NULL, role TEXT NOT NULL,
            group_name TEXT NOT NULL, rate TEXT, hikvision_id TEXT,
            deleted INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(employee_id, effective_day))''')
        connection.execute("INSERT INTO accountant_employees (source_row, name, role, group_name, rate) "
                           "VALUES (1, 'Алиев Жасур', 'официант', 'Обслуживание зала', '180000')")
    roster = RosterStore(path)
    roster.add(name='Новый', role='официант', rate='180000', group_name='Обслуживание зала')
    RosterStore(path)  # повторная миграция ничего не ломает и не дублирует версии
    with closing(roster._open()) as connection:
        for table in ('accountant_employees', 'accountant_employee_versions'):
            assert {'work_from', 'work_to'} <= table_columns(connection, table)
        assert connection.execute('SELECT COUNT(*) FROM accountant_employee_versions').fetchone()[0] == 2
    person = roster.list()[0]
    assert (person.name, person.employment_type, person.work_from, person.work_to) == (
        'Алиев Жасур', 'shift', None, None)
    assert names(roster.list(d(9))) == ['Алиев Жасур', 'Новый']


# ── Правка: продлить, сократить, сменить тип ─────────────────────────────────

def test_shortening_is_refused_while_records_are_outside_extending_is_free(stores):
    roster, finance, attendance, people = stores
    karamat = people['karamat'].id
    finance.set_salary_day_cell(d(11), karamat, '150000', '0')   # смена 10.10 — последняя
    finance.add_extra_payout(employee_id=karamat, work_day=d(8), paid_day=d(9), amount='50000', note='Банкет')
    with pytest.raises(ValueError, match='начисления или выплаты за смены 10.10'):
        roster.update(karamat, rate='150000', reason='Ушла раньше', work_to=d(9))
    with pytest.raises(ValueError, match='доп. выплаты за смены 08.10'):
        roster.update(karamat, rate='150000', reason='Пришла позже', work_from=d(9))
    roster.check_period(karamat, 'shift', None, None)   # снять период — всегда можно
    with pytest.raises(ValueError, match='Сначала уберите их'):
        roster.check_period(karamat, 'temporary', d(9), d(9))
    # Ничего не записалось: период прежний.
    assert (roster.list()[1].work_from, roster.list()[1].work_to) == (FROM, TO)
    # Продлить можно: смена 11.10 теперь в периоде, в том числе в версиях прошлых дней.
    longer = roster.update(karamat, rate='150000', reason='Продлили на два дня', work_to=d(12), by='buh')
    assert longer.work_to == d(12) and 'Карамат' in names(roster.list(d(11)))
    assert longer.json()['work_period'] == '08.10–12.10'
    entry = roster.history(karamat)[0]
    assert (entry['reason'], entry['details'], entry['changed_by']) == (
        'Продлили на два дня', 'Период: 08.10–10.10 → 08.10–12.10', 'buh')
    # Правка без типа и периода (директор, прежний экран) период не трогает.
    roster.update(karamat, rate='160000', reason='Новая ставка')
    assert roster.list()[1].work_to == d(12)
    # Убрать выплату — и сократить уже можно; отметка «не был» за 12.10 снова мешает.
    finance.set_salary_day_cell(d(11), karamat, '0', '150000')
    roster.update(karamat, rate='160000', reason='Сократили', work_to=d(12))
    finance.mark_manual_attendance(karamat, d(12), False, 'buh')
    with pytest.raises(ValueError, match='отметки «был / не был» за 12.10'):
        roster.update(karamat, rate='160000', reason='Сократили', work_to=d(11))
    # Снова сменный — период снимается целиком: ограничений больше нет.
    shift = roster.update(karamat, rate='160000', reason='Взяли в штат', employment_type='shift',
                          work_from=None, work_to=None)
    assert (shift.employment_type, shift.work_from, shift.work_to, shift.json()['work_period']) == (
        'shift', None, None, None)
    assert roster.history(karamat)[0]['details'] == 'Тип: временный → сменный; Период: 08.10–12.10 → —'
    assert all('Карамат' in names(roster.list(day)) for day in (d(1), d(7), d(12)))
    with pytest.raises(ValueError, match='только у временного'):
        roster.update(karamat, rate='160000', reason='x', work_from=d(8))
    with pytest.raises(ValueError, match='раньше, чем'):
        roster.update(karamat, rate='160000', reason='x', employment_type='temporary', work_from=d(9), work_to=d(8))


# ── Сервер не пишет вне периода ──────────────────────────────────────────────

def test_money_and_marks_outside_the_period_are_refused_in_words(stores):
    roster, finance, _, people = stores
    karamat = people['karamat'].id
    before = {table: count(finance, table) for table in (
        'accountant_accruals', 'accountant_salary_payments', 'accountant_movements', 'accountant_extra_payouts',
        'hikvision_manual_absences', 'accountant_exceptions')}
    for action, text in [
            (lambda: finance.set_salary_day_cell(d(12), karamat, '150000', '0'),
             'выплату за смену 11.10 записать нельзя.'),
            (lambda: finance.set_salary_day_cell(d(8), karamat, '150000', '0'),
             'выплату за смену 07.10 записать нельзя.'),
            (lambda: finance.add_extra_payout(employee_id=karamat, work_day=d(7), paid_day=d(9), amount='150000',
                                              note='Подмена'), 'доп. выплату за смену 07.10 записать нельзя.'),
            (lambda: finance.mark_manual_attendance(karamat, d(12), True, 'buh'), 'смену 12.10 отметить нельзя.'),
            (lambda: finance.mark_manual_attendance(karamat, d(7), False, 'buh'), 'смену 07.10 отметить нельзя.'),
            (lambda: finance.grant_exception(karamat, d(12), 'Без турникета', 'buh'),
             'исключение за смену 12.10 дать нельзя.')]:
        with pytest.raises(LedgerError) as error:
            action()
        assert str(error.value) == KARAMAT_TEXT + text
    assert before == {table: count(finance, table) for table in before}
    # Свои смены — как у всех: 08.10 (клетка 09.10) и последняя 10.10 (клетка 11.10).
    for paid in (d(9), d(11)):
        assert finance.set_salary_day_cell(paid, karamat, '150000', '0')['changed'] is True
    item = finance.add_extra_payout(employee_id=karamat, work_day=d(10), paid_day=d(12), amount='50000',
                                    note='Банкет', confirm=True)
    with pytest.raises(LedgerError, match=KARAMAT_TEXT + 'доп. выплату за смену 07.10'):
        finance.update_extra_payout(item['id'], amount='50000', work_day=d(7), note='Банкет')
    assert finance.mark_manual_attendance(karamat, d(9), True, 'buh') is True
    # Временный без периода — как сменный: любой день.
    guest = people['guest'].id
    assert finance.set_salary_day_cell(d(6), guest, '150000', '0')['changed'] is True
    assert finance.set_salary_day_cell(d(12), guest, '150000', '0')['changed'] is True


def test_payroll_confirmation_leaves_her_out_outside_the_period(stores):
    roster, finance, attendance, people = stores
    service = AttendanceService(attendance, source='retro-main-entry', enabled=False, poll_seconds=60)
    day = d(11)
    staff = roster.list(day)
    assert 'Карамат' not in names(staff)
    rows = draft_payroll(day, staff, set(), service.snapshot(day, staff, now=datetime(2026, 10, 12, 9, tzinfo=TZ)).rows)
    finance.confirm_payroll(day, rows, 'buh')
    assert people['karamat'].id not in finance.accrued_employees(day)
    # Строка, посчитанная до сокращения периода, не начислится и в гонке.
    forged = [PayrollRow(people['karamat'].id, 'Карамат', 'Хостес', 'Встреча гостей', 'manual_present', None,
                         Decimal('150000'), Decimal('150000'), False)]
    with pytest.raises(LedgerError, match=KARAMAT_TEXT + 'смену 07.10 начислить нельзя'):
        finance.confirm_payroll(d(7), forged, 'buh')


# ── «Зарплата · день» ────────────────────────────────────────────────────────

def test_salary_day_shows_her_only_in_months_of_her_shifts_with_locked_cells(stores):
    roster, finance, _, people = stores
    roster.add(name='Сардор', role='официант', rate='180000', group_name='Обслуживание зала',
               employment_type='temporary', work_from=date(2026, 10, 25), work_to=date(2026, 10, 31))
    october = finance.salary_day_month(d(2), date(2026, 10, 31))
    karamat = next(p for p in october['people'] if p['name'] == 'Карамат')
    assert (karamat['temporary'], karamat['work_period'], karamat['work_from'], karamat['work_to']) == (
        True, '08.10–10.10', '2026-10-08', '2026-10-10')
    open_days = [day for day, cell in karamat['cells'].items() if cell['editable']]
    assert open_days == ['2026-10-09', '2026-10-10', '2026-10-11']
    outside = [day for day, cell in karamat['cells'].items() if cell.get('outside')]
    assert outside[:6] == ['2026-10-02', '2026-10-03', '2026-10-04', '2026-10-05', '2026-10-06', '2026-10-07']
    assert '2026-10-12' in outside and '2026-10-08' in outside
    assert all(cell['amount'] == '0' for cell in karamat['cells'].values())
    # У сменного и временного без периода — как раньше: ни одной «вне периода».
    for name in ('Баходиров Ихтиер', 'Гулноза'):
        person = next(p for p in october['people'] if p['name'] == name)
        assert not any(cell.get('outside') for cell in person['cells'].values())
        assert 'work_period' not in person
    # Ноябрь: смена 31.10 Сардора выплачивается 01.11 — его строка есть; Карамат — нет.
    november = finance.salary_day_month(date(2026, 11, 1), date(2026, 11, 30))
    assert [p['name'] for p in november['people']] == ['Баходиров Ихтиер', 'Гулноза', 'Сардор']
    sardor = next(p for p in november['people'] if p['name'] == 'Сардор')
    assert [day for day, cell in sardor['cells'].items() if not cell.get('outside')] == ['2026-11-01']
    # Сентябрь — ни одной смены периода: строки Карамат нет.
    september = finance.salary_day_month(date(2026, 9, 1), date(2026, 9, 30))
    assert 'Карамат' not in [p['name'] for p in september['people']]


# ── Через API: «Сотрудники», отказы словами, правка, Excel ───────────────────

@pytest.fixture
def api(any_db):  # noqa: F811 — фикстура
    c = any_db
    finance = c.app.state.accountant_finance
    for offset in range(11):
        finance.record_handover(d(5) + timedelta(days=offset), Decimal('500000'))
    finance.set_cash_opening(d(5), '5000000', 'Начало')
    roster = c.app.state.accountant_roster
    c.ihtiyor = roster.add(name='Баходиров Ихтиер', role='Менеджер', rate='360000', group_name='Управление')
    return c


def test_api_adds_karamat_and_the_staff_page_sees_her_only_in_her_days(api):
    created = api.post('/api/accountant/employees', json=dict(
        name='Карамат', role='Хостес', rate='150000', group='Встреча гостей', manual_attendance=True,
        employment_type='temporary', work_from='2026-10-08', work_to='2026-10-10'))
    assert created.status_code == 201, created.text
    karamat = created.json()['employee']
    assert (karamat['employment_type'], karamat['work_period'], karamat['manual_attendance']) == (
        'temporary', '08.10–10.10', True)
    inside = api.get('/api/accountant/staff', params={'date': '2026-10-09'}).json()
    row = next(item for item in inside['employees'] if item['name'] == 'Карамат')
    assert (row['employment_type'], row['work_from'], row['work_to'], row['work_period']) == (
        'temporary', '2026-10-08', '2026-10-10', '08.10–10.10')
    assert inside['outside_period'] == [] and inside['roster_count'] == 2
    outside = api.get('/api/accountant/staff', params={'date': '2026-10-12'}).json()
    assert [item['name'] for item in outside['employees']] == ['Баходиров Ихтиер']
    # Вне периода — не «не пришла» и не в счётчиках; карточка доступна для правки.
    assert outside['roster_count'] == 1 and outside['payroll']['manual_absent_count'] == 0
    assert [(item['name'], item['status'], item['work_period']) for item in outside['outside_period']] == [
        ('Карамат', 'outside', '08.10–10.10')]
    # «Финансы дня»: смена 12.10 без неё.
    day = api.get('/api/accountant/day', params={'date': '2026-10-12'}).json()
    assert [item['name'] for item in day['employees']] == ['Баходиров Ихтиер']
    # Без периода и сменный — ничего нового.
    plain = api.post('/api/accountant/employees', json=dict(name='Гулноза', role='Хостес', rate='150000',
                                                           group='Встреча гостей'))
    assert (plain.json()['employee']['employment_type'], plain.json()['employee']['work_period']) == ('shift', None)
    bad = api.post('/api/accountant/employees', json=dict(
        name='Нигора', role='Хостес', group='Встреча гостей', employment_type='temporary',
        work_from='2026-10-10', work_to='2026-10-08'))
    assert (bad.status_code, bad.json()['detail']) == (422, 'Период работы: «по» (08.10) раньше, чем «с» (10.10).')


def test_api_refusals_name_her_period(api):
    karamat = api.post('/api/accountant/employees', json=dict(
        name='Карамат', role='Хостес', rate='150000', group='Встреча гостей', manual_attendance=True,
        employment_type='temporary', work_from='2026-10-08', work_to='2026-10-10')).json()['employee']
    checks = [
        (api.post('/api/accountant/manual-attendance', json=dict(date='2026-10-12', employee_id=karamat['id'],
                                                                 present=True)),
         'смену 12.10 отметить нельзя.'),
        (api.put('/api/accountant/salary-day/cell', json=dict(date='2026-10-12', employee_id=karamat['id'],
                                                             amount='150000', expected_amount='0')),
         'выплату за смену 11.10 записать нельзя.'),
        (api.post('/api/accountant/salary-day/extra', json=dict(
            employee_id=karamat['id'], work_day='2026-10-07', paid_day='2026-10-09', amount='150000', note='Подмена')),
         'доп. выплату за смену 07.10 записать нельзя.'),
        (api.post('/api/accountant/exceptions', json=dict(date='2026-10-12', employee_id=karamat['id'],
                                                          reason='Без турникета', approver='buh')),
         'исключение за смену 12.10 дать нельзя.')]
    for response, text in checks:
        assert (response.status_code, response.json()['detail']) == (422, KARAMAT_TEXT + text)
    # Своя смена — пишется.
    ok = api.put('/api/accountant/salary-day/cell', json=dict(date='2026-10-11', employee_id=karamat['id'],
                                                             amount='150000', expected_amount='0'))
    assert ok.status_code == 200, ok.text
    extra = api.post('/api/accountant/salary-day/extra', json=dict(
        employee_id=karamat['id'], work_day='2026-10-08', paid_day='2026-10-09', amount='50000', note='Банкет'))
    assert extra.status_code == 201, extra.text
    moved = api.put(f'/api/accountant/salary-day/extra/{extra.json()["id"]}', json=dict(
        work_day='2026-10-07', amount='50000', note='Банкет', expected_amount='50000'))
    assert (moved.status_code, moved.json()['detail']) == (422, KARAMAT_TEXT + 'доп. выплату за смену 07.10 записать нельзя.')
    assert api.post('/api/accountant/manual-attendance', json=dict(
        date='2026-10-09', employee_id=karamat['id'], present=True)).status_code == 200


def test_api_edit_extends_refuses_shortening_and_keeps_the_card_whole(api):
    karamat = api.post('/api/accountant/employees', json=dict(
        name='Карамат', role='Хостес', rate='150000', group='Встреча гостей',
        employment_type='temporary', work_from='2026-10-08', work_to='2026-10-10')).json()['employee']
    url = f'/api/accountant/employees/{karamat["id"]}'
    assert api.put('/api/accountant/salary-day/cell', json=dict(
        date='2026-10-11', employee_id=karamat['id'], amount='150000', expected_amount='0')).status_code == 200
    # Сократить по 09.10 нельзя — выплата за смену 10.10. Остальное в карточке не сохранилось.
    refused = api.patch(url, json=dict(name='Карамат', role='Хостес', rate='200000', group='Встреча гостей',
                                       reason='Ушла раньше', hikvision_id='77', employment_type='temporary',
                                       work_from='2026-10-08', work_to='2026-10-09'))
    assert refused.status_code == 422
    assert 'начисления или выплаты за смены 10.10' in refused.json()['detail']
    card = next(item for item in api.app.state.accountant_roster.list() if item.id == karamat['id'])
    assert (str(card.rate), card.hikvision_id, card.work_to) == ('150000', None, d(10))
    # Продлить — можно; только конец, начало остаётся.
    longer = api.patch(url, json=dict(rate='150000', reason='Продлили', work_to='2026-10-14'))
    assert longer.status_code == 200, longer.text
    assert (longer.json()['employee']['work_from'], longer.json()['employee']['work_period']) == (
        '2026-10-08', '08.10–14.10')
    history = api.get(f'{url}/history').json()['history']
    assert (history[0]['reason'], history[0]['details']) == ('Продлили', 'Период: 08.10–10.10 → 08.10–14.10')
    # Снова сменный: даты не прислали — период снимается сам.
    shift = api.patch(url, json=dict(rate='150000', reason='В штат', employment_type='shift'))
    assert shift.status_code == 200, shift.text
    assert (shift.json()['employee']['employment_type'], shift.json()['employee']['work_period']) == ('shift', None)


def test_salary_day_month_and_excel_follow_the_period(api):
    karamat = api.post('/api/accountant/employees', json=dict(
        name='Карамат', role='Хостес', rate='150000', group='Встреча гостей',
        employment_type='temporary', work_from='2026-10-08', work_to='2026-10-10')).json()['employee']
    api.put('/api/accountant/salary-day/cell', json=dict(date='2026-10-09', employee_id=karamat['id'],
                                                        amount='150000', expected_amount='0'))
    month = api.get('/api/accountant/salary-day/month', params={'month': '2026-10'}).json()
    row = next(p for p in month['people'] if p['id'] == karamat['id'])
    assert row['cells']['2026-10-12'] == dict(amount='0', work_day='2026-10-11', work_days=[], editable=False, rate='150000',
                                              outside=True)
    # Посещаемость клетке вне периода не ставится: её нет в списке дня смены.
    assert 'attendance' not in row['cells']['2026-10-08']
    response = api.get('/api/accountant/salary-day/export', params={'month': '2026-10'})
    sheet = load_workbook(BytesIO(response.content))['Ведомость']
    rows = [[cell.value for cell in line] for line in sheet.iter_rows()]
    line = next(index for index, values in enumerate(rows, 1) if values[1] == 'Карамат')
    assert rows[line - 1][2] == 'Хостес · временный · 08.10–10.10'
    head = rows[3]
    column = {value.split('\n')[0]: index + 1 for index, value in enumerate(head) if value and value.startswith('Выплата')}
    assert sheet.cell(line, column['Выплата 09.10']).value == 150000
    assert sheet.cell(line, column['Выплата 12.10']).value is None
    assert sheet.cell(line, column['Выплата 12.10']).fill.fgColor.rgb.endswith('F2F4EF')
    assert sheet.cell(line, column['Выплата 10.10']).fill.fgColor.rgb in (None, '00000000')
