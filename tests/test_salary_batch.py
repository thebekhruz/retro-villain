"""Deferred shift payouts: actual cash date, atomicity, partials and corrections.

Every case runs on SQLite and the isolated CI Postgres schema.
"""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from test_salary_payout_dates import c, seed, oct_, spent_on, cell, sheet, table, cash_days  # noqa: F401


def preview(c, first='2026-10-07', last='2026-10-08', paid='2026-10-10'):
    response = c.get('/api/accountant/salary-day/batch', params=dict(first=first, last=last, paid_day=paid))
    assert response.status_code == 200, response.text
    return response.json()


def item(row, amount=None):
    return dict(employee_id=row['employee_id'], work_day=row['work_day'], amount=amount or row['remaining'], version=row['version'])


def selected(c, employee_id, **kwargs):
    return [r for r in preview(c, **kwargs)['rows'] if r['employee_id'] == employee_id]


def pay(c, items, paid='2026-10-10', **kwargs):
    return c.post('/api/accountant/salary-day/batch', json=dict(paid_day=paid, items=items), **kwargs)


def test_two_shifts_one_actual_day_late_entry_consistent_everywhere_and_export(c):
    person = seed(c)['ikhtiyor']
    rows = selected(c, person.id)
    assert all(r['editable'] for r in rows)
    c.clock['today'] = oct_(12)  # entered later; cash still went out on the 10th
    response = pay(c, [item(r) for r in rows])
    assert response.status_code == 201, response.text
    assert response.json()['total'] == '720000'
    assert spent_on(c, oct_(8)) == spent_on(c, oct_(9)) == 0
    assert spent_on(c, oct_(10)) == Decimal('720000')
    stored = next(p for p in sheet(c)['people'] if p['id'] == person.id)['cells']
    assert stored['2026-10-10']['work_days'] == ['2026-10-07', '2026-10-08']
    assert len(stored['2026-10-10']['payment_details']) == 2
    assert stored['2026-10-10']['batch_managed'] and not stored['2026-10-10']['editable']
    assert stored['2026-10-08']['paid_elsewhere'] == ['2026-10-10']
    assert cell(c, oct_(10), person, '0', '720000').status_code == 409
    book = load_workbook(BytesIO(c.get('/api/accountant/salary-day/export?month=2026-10').content))
    header = [v.value for v in book['Ведомость'][4]]
    col = header.index('Выплата 10.10\nза 07.10, 08.10')
    exported = next(r for r in book['Ведомость'].iter_rows(min_row=5) if r[1].value == person.name)[col]
    assert exported.value == 720000
    assert '07.10.2026: 360000' in exported.comment.text
    assert '08.10.2026: 360000' in exported.comment.text


def test_duplicate_request_replay_and_new_key_do_not_pay_twice(c):
    person = seed(c)['ikhtiyor']
    items = [item(r) for r in selected(c, person.id)]
    headers = {'Idempotency-Key': str(uuid4())}
    first = pay(c, items, headers=headers)
    second = pay(c, items, headers=headers)
    assert first.status_code == second.status_code == 201, (first.text, second.text)
    assert first.json() == second.json()
    assert pay(c, items).status_code == 409
    assert len(table(c, 'accountant_salary_payments')) == 2


def test_partial_then_rest_and_edit_cancel_restore_only_that_shift(c):
    person = seed(c)['ikhtiyor']
    first = pay(c, [item(selected(c, person.id)[0], '100000')])
    assert first.status_code == 201, first.text
    row = selected(c, person.id)[0]
    assert (row['paid'], row['remaining']) == ('100000', '260000')
    rest = pay(c, [item(row)], paid='2026-10-09')
    assert rest.status_code == 201, rest.text
    assert not selected(c, person.id)[0]['editable']
    payment = first.json()['payments'][0]
    edit = c.put(f'/api/accountant/salary-day/batch-payments/{payment["id"]}',
                 json=dict(amount='50000', expected_amount='100000'))
    assert edit.status_code == 200, edit.text
    assert selected(c, person.id)[0]['remaining'] == '50000'
    over = c.put(f'/api/accountant/salary-day/batch-payments/{payment["id"]}',
                 json=dict(amount='100001', expected_amount='50000'))
    assert over.status_code == 422
    canceled = c.put(f'/api/accountant/salary-day/batch-payments/{payment["id"]}',
                     json=dict(amount='0', expected_amount='50000'))
    assert canceled.status_code == 200, canceled.text
    assert selected(c, person.id)[0]['remaining'] == '100000'
    assert selected(c, person.id)[0]['attendance']['time'] == '09:38'
    assert spent_on(c, oct_(10)) == 0
    assert spent_on(c, oct_(9)) == Decimal('260000')


def test_stale_or_invalid_last_row_rolls_back_entire_batch(c):
    person = seed(c)['ikhtiyor']
    rows = selected(c, person.id)
    response = pay(c, [item(rows[0]), item(rows[1], '360001')])
    assert response.status_code == 422
    assert not table(c, 'accountant_salary_payments')
    assert not table(c, 'accountant_accruals')
    assert cell(c, oct_(9), person, '360000').status_code == 200
    response = pay(c, [item(r) for r in rows])
    assert response.status_code == 409
    assert len(table(c, 'accountant_salary_payments')) == 1
    assert spent_on(c, oct_(10)) == 0


def test_preview_and_submit_require_known_attendance(c):
    people = seed(c)
    unknown = selected(c, people['selvina'].id)
    assert not any(r['editable'] for r in unknown)
    response = pay(c, [item(unknown[0])])
    assert response.status_code == 422
    c.app.state.attendance_store.set_manual_mark(people['selvina'].id, oct_(7), True, 'Бухгалтер')
    confirmed = selected(c, people['selvina'].id)[0]
    assert confirmed['editable']
    c.app.state.attendance_store.set_manual_mark(people['selvina'].id, oct_(7), False, 'Бухгалтер')
    assert pay(c, [item(confirmed)]).status_code == 409
    assert not table(c, 'accountant_salary_payments')


def test_insufficient_cash_and_duplicate_shift_are_atomic(c):
    person = seed(c, opening='500000', handover='0')['ikhtiyor']
    rows = selected(c, person.id)
    assert pay(c, [item(rows[0]), item(rows[0])]).status_code == 422
    assert pay(c, [item(r) for r in rows]).status_code == 422
    assert not table(c, 'accountant_salary_payments')
    assert not table(c, 'accountant_accruals')


def test_concurrent_payouts_of_same_shifts_have_one_winner(c):
    person = seed(c)['ikhtiyor']
    items = [item(r) for r in selected(c, person.id)]
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: pay(c, items).status_code, range(2)))
    assert sorted(results) == [201, 409]
    assert spent_on(c, oct_(10)) == Decimal('720000')


def test_cross_month_and_invalid_dates(c):
    person = seed(c)['ikhtiyor']
    c.clock['today'] = date(2026, 11, 2)
    cash_days(c, oct_(11), date(2026, 11, 2), '0')
    for work in (oct_(30), oct_(31)):
        c.app.state.attendance_store.set_manual_mark(person.id, work, True, 'Бухгалтер')
    rows = selected(c, person.id, first='2026-10-30', last='2026-10-31', paid='2026-11-02')
    response = pay(c, [item(r) for r in rows], paid='2026-11-02')
    assert response.status_code == 201, response.text
    november = c.get('/api/accountant/salary-day/month?month=2026-11').json()
    stored = next(p for p in november['people'] if p['id'] == person.id)['cells']['2026-11-02']
    assert stored['amount'] == '720000'
    assert stored['work_days'] == ['2026-10-30', '2026-10-31']
    for first, last, paid in [('2026-10-04', '2026-10-05', '2026-11-02'),
                              ('2026-10-08', '2026-10-07', '2026-11-02'),
                              ('2026-10-30', '2026-10-31', '2026-10-30'),
                              ('2026-11-02', '2026-11-02', '2026-11-02'),
                              ('2026-10-30', '2026-10-31', '2026-11-03')]:
        assert c.get('/api/accountant/salary-day/batch', params=dict(first=first,last=last,paid_day=paid)).status_code == 422


def test_closed_period_blocks_create_and_correction(c):
    person = seed(c)['ikhtiyor']
    response = pay(c, [item(selected(c, person.id)[0])])
    assert response.status_code == 201, response.text
    payment_id = response.json()['payments'][0]['id']
    finance = c.app.state.accountant_finance
    with closing(finance._open()) as connection, connection:
        connection.execute('INSERT INTO accountant_month_closures (month,last_day,closed_by,closed_at,closing_balance,snapshot) VALUES (?,?,?,?,?,?)',
                           ('2026-10', '2026-10-31', 'Бухгалтер', '2026-11-01', '0', '{}'))
    assert not any(r['editable'] for r in selected(c, person.id))
    assert c.put(f'/api/accountant/salary-day/batch-payments/{payment_id}',
                 json=dict(amount='0',expected_amount='360000')).status_code == 409


def test_current_roster_and_rate_version_are_rechecked(c):
    person = seed(c)['ikhtiyor']
    rows = selected(c, person.id)
    finance = c.app.state.accountant_finance
    with closing(finance._open()) as connection, connection:
        connection.execute('UPDATE accountant_employee_versions SET rate=? WHERE employee_id=?', ('400000', person.id))
    assert pay(c, [item(rows[0])]).status_code == 409
    fresh = selected(c, person.id)
    assert fresh[0]['remaining'] == '400000'
    assert pay(c, [item(fresh[0])]).status_code == 201


def test_panels_start_collapsed_and_new_assets_are_private(c):
    html = c.get('/accountant/salary-day').text
    from html.parser import HTMLParser
    found = {}
    class Parser(HTMLParser):
        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            if tag == 'details':
                found[attrs.get('id')] = attrs
    Parser().feed(html)
    assert {'salary-batch', 'salary-sheet'} <= found.keys()
    assert all('open' not in found[key] for key in ('salary-batch', 'salary-sheet'))
    from retro.app import STATIC_PANELS
    assert STATIC_PANELS['salary-batch.js'] == frozenset({'accountant'})


def test_public_batch_range_accepts_first_shift_without_moving_accounting_start():
    from retro.accounting_access import check_dates
    from fastapi import HTTPException
    check_dates([('first', '2026-10-05'), ('last', '2026-10-05'), ('paid_day', '2026-10-06')],
                salary=True, salary_shift_range=True)
    with pytest.raises(HTTPException):
        check_dates([('paid_day', '2026-10-05')], salary=True, salary_shift_range=True)
    with pytest.raises(HTTPException):
        check_dates([('first', '2026-10-04')], salary=True, salary_shift_range=True)


def test_paid_standard_shift_and_temporary_period_cannot_be_paid_again(c):
    people = seed(c)
    person = people['ikhtiyor']
    assert cell(c, oct_(8), person, '360000').status_code == 200
    rows = selected(c, person.id)
    assert not rows[0]['editable'] and rows[0]['reason'] == 'Уже выдано'
    assert rows[1]['editable']
    with closing(c.app.state.accountant_finance._open()) as connection, connection:
        connection.execute('UPDATE accountant_employees SET employment_type=?,work_from=?,work_to=? WHERE id=?',
                           ('temporary','2026-10-08','2026-10-08',person.id))
    outside = selected(c, person.id)[0]
    assert not outside['editable'] and 'работает' in outside['reason']
    assert pay(c, [item(outside)]).status_code == 422


def test_standard_and_deferred_payments_in_same_cell_can_both_be_corrected(c):
    person = seed(c)['ikhtiyor']
    assert cell(c, oct_(9), person, '360000').status_code == 200
    response = pay(c, [item(selected(c, person.id)[0])], paid='2026-10-09')
    assert response.status_code == 201, response.text
    stored = next(p for p in sheet(c)['people'] if p['id'] == person.id)['cells']['2026-10-09']
    assert stored['amount'] == '720000'
    assert all(r['editable'] for r in stored['payment_details'])
    standard = next(r for r in stored['payment_details'] if r['work_day'] == '2026-10-08')
    assert not standard['restores_balance']
    path = f'/api/accountant/salary-day/batch-payments/{standard["id"]}'
    assert c.put(path, json=dict(amount='400000', expected_amount='360000')).status_code == 200
    assert spent_on(c, oct_(9)) == Decimal('760000')
    assert selected(c, person.id)[1]['remaining'] == '0'
    assert c.put(path, json=dict(amount='0', expected_amount='400000')).status_code == 200
    assert spent_on(c, oct_(9)) == Decimal('360000')
    assert selected(c, person.id)[1]['remaining'] == '360000'
