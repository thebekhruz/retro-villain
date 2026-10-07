"""Кассир по макету 5a: передача бухгалтеру, выдача Шоху из кассы, доллары в сейф
и смена кассы из iiko.

Деньги проверяются сквозняком: одна запись кассира должна одинаково читаться у
кассира, у бухгалтера (остаток, «Баланс Шох», резервы), у Шоха на телефоне и у
учредителя — без двойного списания. Всё, что пишет в базу, гоняется и на
Postgres, если задан RETRO_TEST_POSTGRES_URL, — как в CI.
"""

import os
from contextlib import closing, contextmanager
from dataclasses import replace
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from uuid import uuid4

import openpyxl
import pytest
from fastapi.testclient import TestClient

from legacy_app import create_app
from retro.config import Settings
from retro.accounting_period import ACCOUNTING_START
from retro.modules.cashier.archive import decode_snapshot
from retro.modules.cashier.service import TZ, build_snapshot, shift_status, today_tashkent
from retro.modules.cashier.till import (add_usd_deposit, give_shokh, migrate_legacy_usd,
                                         till_totals, usd_day)
from retro.modules.founder import cabinet

POSTGRES_URL = os.getenv('RETRO_TEST_POSTGRES_URL', '')
DAY = today_tashkent() - timedelta(days=1)
RECEIVED = DAY + timedelta(days=1)
DEMO_CASH, PREPAY = Decimal('1000000'), Decimal('200000')


def row(*values):
    return {f'field{i}': {'value': value} for i, value in enumerate(values)}


def snapshot_for(day=DAY, *, cash=DEMO_CASH, **changes):
    snapshot = build_snapshot(day, [row(day.isoformat(), 3, int(cash) + 500000)],
                              [row('Демо', int(cash)), row('UzCard', 500000)])
    return replace(snapshot, cash_prepayment=PREPAY, new_prepayment=PREPAY, **changes)


def settings(tmp_path, **extra):
    return Settings(data_dir=tmp_path, manual_handover_only=True, **extra)


def build_app(tmp_path, **extra):
    return create_app(settings(tmp_path, **extra),
                      expense_db_path=tmp_path / 'cashier.sqlite3',
                      accountant_db_path=tmp_path / 'accountant.sqlite3',
                      director_db_path=tmp_path / 'director.sqlite3',
                      founder_db_path=tmp_path / 'founder.sqlite3')


def client_for(app, **kwargs):
    return TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000), **kwargs)


@contextmanager
def postgres_app(tmp_path):
    import psycopg
    from psycopg import sql
    schema = 'cashier_' + uuid4().hex
    admin = psycopg.connect(POSTGRES_URL, autocommit=True)
    admin.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(schema)))
    parts = urlsplit(POSTGRES_URL)
    query = dict(parse_qsl(parts.query))
    query['options'] = '-csearch_path=' + schema
    url = urlunsplit(parts._replace(query=urlencode(query)))
    try:
        yield lambda: create_app(settings(tmp_path, database_url=url))
    finally:
        admin.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(schema)))
        admin.close()


@pytest.fixture(params=['sqlite', 'postgres'])
def make_app(request, tmp_path):
    """Фабрика приложения: повторный вызов — «перезапуск» на тех же базах."""
    if request.param == 'postgres':
        if not POSTGRES_URL:
            pytest.skip('RETRO_TEST_POSTGRES_URL не задан')
        with postgres_app(tmp_path) as factory:
            yield factory
    else:
        yield lambda: build_app(tmp_path)


@pytest.fixture
def c(make_app):
    with client_for(make_app()) as client:
        yield client


def expected_for(c, day=DAY):
    snapshot = c.app.state.cache.latest_for_day(day)
    totals = till_totals(c.app.state, day)
    return snapshot.payments[0].amount + snapshot.cash_prepayment + totals.receipts - totals.cash_out


def till_day(c, day=DAY, *, expense='100000', receipt='50000', give='300000'):
    """Касса дня: снимок iiko на экране, расход, поступление и выдача Шоху."""
    snapshot = snapshot_for(day)
    c.app.state.cache.put(snapshot)
    if expense:
        assert c.post('/api/cashier/expenses', json={'date': day.isoformat(), 'description': 'Такси',
                                                     'amount': expense}).status_code == 201
    if receipt:
        assert c.post('/api/cashier/receipts', json={'date': day.isoformat(), 'description': 'Долг',
                                                     'amount': receipt}).status_code == 201
    if give:
        assert c.post('/api/cashier/shokh', json={'date': day.isoformat(),
                                                  'amount': give}).status_code == 201
    return snapshot


def hand_over(c, snapshot, expected, day=DAY):
    return c.post('/api/cashier/handover', json={'date': day.isoformat(), 'snapshot_id': snapshot.id,
                                                 'expected': str(expected)})


def accountant_day(c, day=RECEIVED):
    response = c.get('/api/accountant/day', params={'date': day.isoformat()})
    assert response.status_code == 200, response.text
    return response.json()


# ── Смена кассы ──────────────────────────────────────────────────────────

SHIFT_OPEN = dict(cashRegNumber=1, openDate=f'{DAY.isoformat()}T10:56:11', closeDate=None,
                  isOpen=True, sessionStatus='OPEN')
SHIFT_CLOSED = dict(cashRegNumber=1, openDate=f'{DAY.isoformat()}T09:50:07',
                    closeDate=f'{DAY.isoformat()}T23:13:16', isOpen=False, sessionStatus='UNACCEPTED')
SCHOOL_OPEN = dict(cashRegNumber=2, openDate=f'{DAY.isoformat()}T08:04:00', closeDate=None,
                   isOpen=True, sessionStatus='OPEN')


def test_shift_status_comes_only_from_the_retro_register_of_that_day():
    assert shift_status(DAY, [SHIFT_OPEN, SCHOOL_OPEN]).json() == dict(
        open=True, opened_at=f'{DAY.isoformat()}T10:56:11', closed_at=None)
    # Касса школы открыта, касса Retro закрыта — смена Retro закрыта.
    closed = shift_status(DAY, [SHIFT_CLOSED, SCHOOL_OPEN])
    assert closed.open is False and closed.closed_at == f'{DAY.isoformat()}T23:13:16'
    # Переоткрыли после закрытия — смена открыта.
    assert shift_status(DAY, [SHIFT_CLOSED, SHIFT_OPEN]).open is True
    # Смена другого дня, чужая касса и мусор статуса не дают.
    yesterday = dict(SHIFT_OPEN, openDate=(DAY - timedelta(days=1)).isoformat() + 'T10:00:00')
    assert shift_status(DAY, [yesterday, SCHOOL_OPEN, 'bad']) is None
    assert shift_status(DAY, None) is None
    # Закрыта, но без времени закрытия — не выдумываем.
    assert shift_status(DAY, [dict(SHIFT_CLOSED, closeDate=None)]) is None


def test_shift_status_survives_the_day_archive_and_old_rows_read_as_unknown():
    import json
    snapshot = snapshot_for(shift=shift_status(DAY, [SHIFT_CLOSED]))
    payload = json.dumps(snapshot.json())
    assert decode_snapshot(payload).shift == snapshot.shift
    legacy = snapshot.json()
    legacy.pop('shift')
    assert decode_snapshot(json.dumps(legacy)).shift is None


def test_cashier_day_payload_carries_shift_and_handover_state(c):
    snapshot = snapshot_for(shift=shift_status(DAY, [SHIFT_OPEN]))

    class Iiko:
        async def load_prepayments(self, day):
            return ()

        async def load(self, day):
            return snapshot

    c.app.state.iiko = Iiko()
    c.app.state.settings = replace(c.app.state.settings, login='l', password='p', store_id=1)
    data = c.get('/api/cashier/day', params={'date': DAY.isoformat()}).json()
    assert data['shift']['open'] is True
    assert data['handover'] is None
    assert c.get('/api/cashier/day', params={'date': DAY.isoformat(), 'demo': 'true'}).json()['handover'] is None


# ── Передача бухгалтеру ──────────────────────────────────────────────────

def test_handover_is_computed_on_the_server_and_becomes_the_accountant_income(c):
    snapshot = till_day(c)
    expected = DEMO_CASH + PREPAY + Decimal('50000') - Decimal('100000') - Decimal('300000')
    assert expected_for(c) == expected == Decimal('850000')
    # Сумма из браузера не записывается, если разошлась с серверной.
    changed = hand_over(c, snapshot, '999999')
    assert changed.status_code == 409 and 'изменилась' in changed.json()['detail']
    assert c.app.state.accountant_finance.handover_for_day(RECEIVED) is None

    response = hand_over(c, snapshot, '850000.00')
    assert response.status_code == 201, response.text
    handover = response.json()['handover']
    assert handover['amount'] == '850000' and handover['source'] == 'cashier'
    assert datetime.fromisoformat(handover['handed_at']).utcoffset() == timedelta(hours=5)
    assert c.get('/api/cashier/handover', params={'date': DAY.isoformat()}).json()['handover'] == handover

    day = accountant_day(c)
    assert day['expected_cashier'] == '850000'  # ручной режим: запись и есть приход
    assert day['cashier_handover']['amount'] == '850000'
    assert day['cashier_handover']['handed_at'] == handover['handed_at']
    assert day['cashier_handover']['source'] == 'cashier'
    assert day['cashier_handover']['expected'] == '850000'


def test_before_handover_accountant_sees_expected_but_no_income(c):
    till_day(c)
    day = accountant_day(c)
    assert day['expected_cashier'] is None
    assert day['cashier_handover'] == dict(amount=None, handed_at=None, source=None,
                                           confirmed_at=None, confirmed_by=None, expected_amount=None,
                                           shortfall=None, checked=False, calculation=None,
                                           expected_changed=False, cashier_active=None,
                                           expected='850000', expected_at=day['cashier_handover']['expected_at'])
    assert day['cashier_handover']['expected_at'] is not None


def test_repress_updates_the_amount_and_same_amount_keeps_the_time(c):
    snapshot = till_day(c)
    first = hand_over(c, snapshot, '850000').json()['handover']
    again = hand_over(c, snapshot, '850000').json()['handover']
    assert again == first  # та же сумма — «получено» в то же время
    c.post('/api/cashier/expenses', json={'date': DAY.isoformat(), 'description': 'Хлеб', 'amount': '20000'})
    updated = hand_over(c, snapshot, '830000')
    assert updated.status_code == 201
    assert updated.json()['handover']['amount'] == '830000'
    audit = c.app.state.accountant_finance.audit_entries(entity_type='handover', entity_id=RECEIVED.isoformat())
    assert [entry['action'] for entry in audit] == ['create', 'update']
    assert audit[-1]['after']['source'] == 'cashier'


def test_cashier_cannot_overwrite_or_cancel_the_accountant_record(c):
    snapshot = till_day(c)
    assert c.post('/api/accountant/handover', json={'date': RECEIVED.isoformat(), 'amount': '700000',
                                                    'note': 'пересчитала сама'}).status_code == 201
    refused = hand_over(c, snapshot, '850000')
    assert refused.status_code == 409 and 'бухгалтер' in refused.json()['detail']
    assert c.delete('/api/cashier/handover', params={'date': DAY.isoformat()}).status_code == 409
    assert c.app.state.accountant_finance.handover_for_day(RECEIVED) == Decimal('700000')
    assert c.get('/api/cashier/handover', params={'date': DAY.isoformat()}).json()['handover']['source'] == 'accountant'


def test_undo_removes_the_income_until_the_accountant_confirms_it(c):
    """Передача кассира — «ожидается»: тратить её нельзя, пока бухгалтер не
    подтвердил сумму (ТЗ 02.10). После подтверждения кассир её не отменит."""
    snapshot = till_day(c)
    assert hand_over(c, snapshot, '850000').status_code == 201
    assert c.delete('/api/cashier/handover', params={'date': DAY.isoformat()}).status_code == 204
    assert c.app.state.accountant_finance.handover_for_day(RECEIVED) is None
    assert c.delete('/api/cashier/handover', params={'date': DAY.isoformat()}).status_code == 204  # уже нет
    assert hand_over(c, snapshot, '850000').status_code == 201
    finance = c.app.state.accountant_finance
    set_start_balance(finance, RECEIVED, '0')
    expense = {'date': RECEIVED.isoformat(), 'item_code': 'admin_other', 'note': 'Канцтовары', 'amount': '800000'}
    unconfirmed = c.post('/api/accountant/expenses', json=expense)
    assert unconfirmed.status_code == 422 and 'недостаточно' in unconfirmed.json()['detail']
    assert c.post('/api/accountant/handover/confirm', json={'date': RECEIVED.isoformat(),
                                                            'amount': '850000'}).status_code == 200
    spent = c.post('/api/accountant/expenses', json=expense)
    assert spent.status_code == 201, spent.text
    blocked = c.delete('/api/cashier/handover', params={'date': DAY.isoformat()})
    assert blocked.status_code == 409 and 'отменить передачу нельзя' in blocked.json()['detail']
    lower = hand_over(c, snapshot, '750000')
    assert lower.status_code == 409
    assert finance.handover_for_day(RECEIVED) == Decimal('850000')


def test_accountant_confirms_received_cash_and_then_cashier_cannot_undo(c):
    """«Получено» бухгалтера: меньше расчёта — недостача видна, остаток от
    полученного, а кассир свою передачу больше не отменит и не перепишет."""
    snapshot = till_day(c)
    assert hand_over(c, snapshot, '850000').status_code == 201
    state = c.get('/api/cashier/handover', params={'date': DAY.isoformat()}).json()['handover']
    assert state['confirmed_at'] is None and state['shortfall'] is None
    confirmed = c.post('/api/accountant/handover/confirm', json={'date': RECEIVED.isoformat(), 'amount': '550000'})
    assert confirmed.status_code == 200, confirmed.text
    handover = confirmed.json()['handover']
    assert handover['amount'] == '550000' and handover['expected_amount'] == '850000'
    assert Decimal(handover['shortfall']) == Decimal('300000') and handover['confirmed_at']
    assert c.app.state.accountant_finance.handover_for_day(RECEIVED) == Decimal('550000')
    day = accountant_day(c)
    assert day['expected_cashier'] == '550000'
    assert Decimal(day['cashier_handover']['shortfall']) == Decimal('300000')
    # Кассир видит подтверждение и не может ни отменить, ни переписать передачу.
    assert c.get('/api/cashier/handover', params={'date': DAY.isoformat()}).json()['handover']['confirmed_at']
    undo = c.delete('/api/cashier/handover', params={'date': DAY.isoformat()})
    assert undo.status_code == 409 and 'подтвердил' in undo.json()['detail']
    again = hand_over(c, snapshot, '850000')
    assert again.status_code == 409 and 'подтвердил' in again.json()['detail']
    # Повторное подтверждение не теряет расчёт: сверяется с тем же 850 000.
    fixed = c.post('/api/accountant/handover/confirm', json={'date': RECEIVED.isoformat(), 'amount': '850000'})
    assert fixed.json()['handover']['expected_amount'] == '850000'
    assert Decimal(fixed.json()['handover']['shortfall']) == 0
    audit = c.app.state.accountant_finance.audit_entries(entity_type='handover', entity_id=RECEIVED.isoformat())
    assert [entry['action'] for entry in audit][-2:] == ['confirm', 'confirm']


def test_accountant_can_confirm_before_the_cashier_button_against_the_iiko_calculation(c):
    till_day(c)
    confirmed = c.post('/api/accountant/handover/confirm', json={'date': RECEIVED.isoformat(), 'amount': '800000'})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()['handover']['expected_amount'] == '850000'
    assert Decimal(confirmed.json()['handover']['shortfall']) == Decimal('50000')


def test_correcting_the_accountants_own_record_is_not_a_shortfall(c):
    """Своя ручная запись бухгалтера — не расчёт: исправление сверяется с iiko."""
    till_day(c)
    assert c.post('/api/accountant/handover', json={'date': RECEIVED.isoformat(), 'amount': '900000',
                                                    'note': 'пересчитала'}).status_code == 201
    fixed = c.post('/api/accountant/handover/confirm', json={'date': RECEIVED.isoformat(), 'amount': '850000'})
    assert fixed.json()['handover']['expected_amount'] == '850000'
    assert Decimal(fixed.json()['handover']['shortfall']) == 0


# ── T-399: недостача кассы — одно число везде ───────────────────────────

def day_export_lines(c, day=RECEIVED):
    """Excel дня бухгалтера: {подпись строки: сумма} со всех листов."""
    response = c.get('/api/accountant/day/export', params={'date': day.isoformat()})
    assert response.status_code == 200, response.text
    book = openpyxl.load_workbook(BytesIO(response.content))
    return {row[0]: row[1] for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)
            if row and isinstance(row[0], str) and len(row) > 1}


def founder_handover(c, day=DAY):
    week = c.get('/api/founder/week', params={'date': day.isoformat()})
    assert week.status_code == 200, week.text
    return {item['date']: item for item in week.json()['days']}[day.isoformat()]['handover']


def test_shortfall_is_one_number_on_every_screen(c):
    """Недостача = текущий расчёт кассы − полученное. Сообщение после
    «Подтвердить» (ответ сервера), карточка 2a, кассир, учредитель и Excel
    показывают одно и то же число."""
    snapshot = till_day(c)  # расчёт 850 000: Демо + предоплата + поступление − расход − Шох
    assert hand_over(c, snapshot, '850000').status_code == 201
    confirmed = c.post('/api/accountant/handover/confirm', json={'date': RECEIVED.isoformat(), 'amount': '550000'})
    assert confirmed.status_code == 200, confirmed.text
    toast = confirmed.json()['handover']
    assert (toast['shortfall'], toast['calculation'], toast['checked']) == ('300000', '850000', True)
    card = accountant_day(c)['cashier_handover']
    assert (card['shortfall'], card['calculation'], card['expected_changed']) == ('300000', '850000', False)
    cashier = c.get('/api/cashier/handover', params={'date': DAY.isoformat()}).json()['handover']
    assert cashier['shortfall'] == '300000'
    founder = founder_handover(c)
    assert founder['shortfall'] == '300000' and founder['status'] == 'mismatch'
    assert Decimal(founder['expected']) == Decimal('850000')
    lines = day_export_lines(c)
    assert lines['⚠ Получено меньше расчёта на'] == 300000


def test_manual_income_is_checked_against_the_cashier_calculation_at_once(c):
    """Приход, записанный бухгалтером вручную, сверяется с расчётом кассы сразу —
    без «Изменить → Подтвердить»."""
    till_day(c)
    recorded = c.post('/api/accountant/incomes', json={'date': RECEIVED.isoformat(), 'item_code': 'income_cashier',
                                                       'note': 'Касса', 'amount': '450000'})
    assert recorded.status_code == 201, recorded.text
    card = accountant_day(c)['cashier_handover']
    assert card['source'] == 'accountant' and card['confirmed_at'] is None
    assert (card['checked'], card['calculation'], card['shortfall']) == (True, '850000', '400000')
    assert founder_handover(c)['shortfall'] == '400000'
    lines = day_export_lines(c)
    assert lines['Передача кассира · расчёт'] == 850000
    assert lines['⚠ Получено меньше расчёта на'] == 400000


def test_manual_income_is_not_checked_when_the_cashier_did_not_use_the_panel(c):
    """Прод ведёт приход вручную, модуль кассира только внедряется: в день, когда
    кассир в панели ничего не делал, расчёт iiko не знает реальных расходов
    кассы — ручной приход не сверяется нигде (2a, учредитель, Excel)."""
    c.app.state.cache.put(snapshot_for())  # расчёт iiko есть: 1 200 000
    recorded = c.post('/api/accountant/incomes', json={'date': RECEIVED.isoformat(), 'item_code': 'income_cashier',
                                                       'note': 'Касса', 'amount': '450000'})
    assert recorded.status_code == 201, recorded.text
    card = accountant_day(c)['cashier_handover']
    assert card['cashier_active'] is False and card['expected'] == '1200000'
    assert (card['checked'], card['calculation'], card['shortfall']) == (False, None, None)
    founder = founder_handover(c)
    assert founder['status'] == 'unchecked' and founder['shortfall'] is None
    lines = day_export_lines(c)
    assert '⚠ Получено меньше расчёта на' not in lines
    assert 'Кассир в панели не работал — сверки с расчётом нет.' in lines
    # Кассир отметил хоть одну операцию — сверка появляется сразу.
    assert c.post('/api/cashier/expenses', json={'date': DAY.isoformat(), 'description': 'Хлеб',
                                                 'amount': '50000'}).status_code == 201
    card = accountant_day(c)['cashier_handover']
    assert (card['cashier_active'], card['checked'], card['shortfall']) == (True, True, '700000')
    assert founder_handover(c)['shortfall'] == '700000'


def test_manual_income_without_a_cashier_calculation_has_no_check(c):
    """iiko нет — сверять не с чем: проверки «получено меньше расчёта» нет,
    даже когда кассир в панели работал."""
    assert c.post('/api/cashier/expenses', json={'date': DAY.isoformat(), 'description': 'Хлеб',
                                                 'amount': '50000'}).status_code == 201
    recorded = c.post('/api/accountant/incomes', json={'date': RECEIVED.isoformat(), 'item_code': 'income_cashier',
                                                       'note': 'Касса', 'amount': '450000'})
    assert recorded.status_code == 201, recorded.text
    card = accountant_day(c)['cashier_handover']
    assert card['checked'] is True and card['calculation'] is None and card['shortfall'] is None
    assert '⚠ Получено меньше расчёта на' not in day_export_lines(c)


def test_cashier_change_after_confirmation_asks_the_accountant_again(c):
    """Кассир изменил день после подтверждения (поступление, выдача Шоху) — его
    не блокируем; бухгалтер видит «касса изменилась», сверка идёт с ТЕКУЩИМ
    расчётом, повторное подтверждение запоминает новый."""
    snapshot = till_day(c)
    assert hand_over(c, snapshot, '850000').status_code == 201
    assert c.post('/api/accountant/handover/confirm',
                  json={'date': RECEIVED.isoformat(), 'amount': '850000'}).status_code == 200
    card = accountant_day(c)['cashier_handover']
    assert (card['expected_changed'], card['shortfall']) == (False, '0')
    # После подтверждения кассир принял ещё 100 000 — к передаче стало 950 000.
    extra = c.post('/api/cashier/receipts', json={'date': DAY.isoformat(), 'description': 'Возврат долга',
                                                  'amount': '100000'})
    assert extra.status_code == 201, extra.text
    card = accountant_day(c)['cashier_handover']
    assert card['expected_changed'] is True
    assert (card['expected_amount'], card['calculation'], card['shortfall']) == ('850000', '950000', '100000')
    assert founder_handover(c)['shortfall'] == '100000'
    lines = day_export_lines(c)
    assert lines['⚠ Касса изменилась после подтверждения · было'] == 850000
    assert lines['⚠ Получено меньше расчёта на'] == 100000
    # Кассир выдал Шоху ещё 100 000 — расчёт снова 850 000, но при подтверждении
    # он уже был таким: ничего не изменилось, недостачи нет.
    assert c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': '100000'}).status_code == 201
    card = accountant_day(c)['cashier_handover']
    assert (card['expected_changed'], card['shortfall']) == (False, '0')
    # Ещё поступление — бухгалтер подтверждает снова: расчёт на момент подтверждения — новый.
    c.post('/api/cashier/receipts', json={'date': DAY.isoformat(), 'description': 'Чай', 'amount': '40000'})
    again = c.post('/api/accountant/handover/confirm', json={'date': RECEIVED.isoformat(), 'amount': '890000'})
    assert again.status_code == 200, again.text
    handover = again.json()['handover']
    assert (handover['expected_amount'], handover['expected_changed'], handover['shortfall']) == ('890000', False, '0')


def test_cashier_undo_is_allowed_while_the_money_is_not_spent(c):
    """Зарплату выдают утром, кассу передают вечером: операции дня сами по
    себе не запрещают кассиру отменить неподтверждённую передачу."""
    finance = c.app.state.accountant_finance
    before = DAY
    finance.record_handover(before, Decimal('1000000'))
    set_start_balance(finance, before, '1000000')
    snapshot = till_day(c)
    assert hand_over(c, snapshot, '850000').status_code == 201
    spent = c.post('/api/accountant/expenses', json={'date': RECEIVED.isoformat(), 'item_code': 'admin_other',
                                                     'note': 'Канцтовары', 'amount': '100000'})
    assert spent.status_code == 201, spent.text
    assert c.delete('/api/cashier/handover', params={'date': DAY.isoformat()}).status_code == 204
    assert finance.handover_for_day(RECEIVED) is None


def test_handover_needs_the_fresh_snapshot_on_screen(c):
    live = today_tashkent()
    till_day(c, live)
    demo = replace(snapshot_for(live), id=uuid4().hex, demo=True)
    stale = replace(snapshot_for(live), id=uuid4().hex, stale=True)
    old = replace(snapshot_for(live), id=uuid4().hex, fetched_at=datetime.now(TZ) - timedelta(minutes=6))
    for snapshot in (demo, stale, old):
        c.app.state.cache.put(snapshot)
    assert hand_over(c, demo, '850000', live).status_code == 422
    assert hand_over(c, stale, '850000', live).status_code == 409
    assert 'старше 5 минут' in hand_over(c, old, '850000', live).json()['detail']
    assert c.post('/api/cashier/handover', json={'date': DAY.isoformat(), 'snapshot_id': 'a' * 32}).status_code == 409
    assert c.app.state.accountant_finance.handover_for_day(RECEIVED) is None
    # Прошлый день с окончательным снимком (после 06:00 следующего дня) — передавать можно.
    past = DAY - timedelta(days=3)
    final = replace(snapshot_for(past), fetched_at=datetime.now(TZ) - timedelta(days=1))
    c.app.state.cache.put(final)
    assert hand_over(c, final, str(DEMO_CASH + PREPAY), past).status_code == 201


def test_negative_handover_is_not_recorded(c):
    snapshot = till_day(c, give='2000000')
    response = hand_over(c, snapshot, str(expected_for(c)))
    assert response.status_code == 422 and 'передавать нечего' in response.json()['detail']
    assert c.app.state.accountant_finance.handover_for_day(RECEIVED) is None


def test_repeated_request_with_the_same_key_is_not_recorded_twice(c):
    snapshot = till_day(c)
    body = {'date': DAY.isoformat(), 'snapshot_id': snapshot.id, 'expected': '850000'}
    key = {'Idempotency-Key': str(uuid4())}
    first = c.post('/api/cashier/handover', json=body, headers=key)
    c.post('/api/cashier/expenses', json={'date': DAY.isoformat(), 'description': 'Хлеб', 'amount': '20000'})
    replay = c.post('/api/cashier/handover', json=body, headers=key)
    assert replay.status_code == first.status_code == 201
    assert replay.json() == first.json()


def test_cashier_role_hands_over_but_never_reaches_the_accountant(tmp_path):
    users = {'kassa': ('secret', 'cashier'), 'buh': ('secret', 'accountant')}
    app = build_app(tmp_path, dashboard_panel_users=users)
    snapshot = snapshot_for()
    app.state.cache.put(snapshot)
    with client_for(app) as c:
        cashier, accountant = ('kassa', 'secret'), ('buh', 'secret')
        body = {'date': DAY.isoformat(), 'snapshot_id': snapshot.id, 'expected': str(DEMO_CASH + PREPAY)}
        assert c.post('/api/cashier/handover', json=body, auth=cashier).status_code == 201
        assert c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': '1'}, auth=cashier).status_code == 201
        assert c.get('/api/accountant/day', params={'date': DAY.isoformat()}, auth=cashier).status_code == 403
        assert c.post('/api/accountant/handover', json={'date': RECEIVED.isoformat(), 'amount': '1', 'note': 'x'},
                      auth=cashier).status_code == 403
        assert c.post('/api/cashier/handover', json=body, auth=accountant).status_code == 403
        assert c.get('/api/accountant/day', params={'date': DAY.isoformat()}, auth=accountant).status_code == 200


# ── Выдать Шоху из кассы: одна запись, три экрана, без двойного списания ──

def set_start_balance(finance, through, amount):
    # Рабочий остаток задаётся на 2 октября; в тесте все промежуточные дни
    # явно закрыты нулевыми передачами, а не считаются известными автоматически.
    first = min(through, ACCOUNTING_START)
    day = first
    while day <= through:
        if finance.handover_for_day(day) is None:
            finance.record_handover(day, Decimal('0'))
        day += timedelta(days=1)
    finance.set_cash_opening(first, amount, 'Пересчёт')


def accountant_cash(c, day=DAY, *, handover='5000000'):
    finance = c.app.state.accountant_finance
    finance.record_handover(day, Decimal(handover))
    set_start_balance(finance, day, '0')
    finance.reserve_entry(min(day, ACCOUNTING_START), 'shoh', 'opening', '420000', 'Остаток у Шоха')


def test_till_give_raises_shokh_balance_but_not_accountant_spending(c):
    accountant_cash(c)
    assert c.post('/api/accountant/procurement', json={'date': DAY.isoformat(), 'recipient': 'Шох',
                                                       'purpose': 'закуп за день', 'amount': '3000000'}).status_code == 201
    given = c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': '500000'})
    assert given.status_code == 201, given.text
    assert given.json()['total'] == '500000'
    assert Decimal(given.json()['pocket']) == Decimal('3920000')

    day = accountant_day(c, DAY)
    shoh = day['reserves']['shoh']
    assert Decimal(shoh['balance']) == Decimal('420000') + Decimal('3000000') + Decimal('500000')
    till = [entry for entry in shoh['entries'] if entry.get('source') == 'cashier']
    assert [(entry['kind'], entry['amount']) for entry in till] == [('deposit', '500000')]
    assert till[0]['created_at'].endswith('+05:00')
    # Деньги бухгалтера: минус только его собственная выдача.
    assert Decimal(day['ledger']['cash_balance']) == Decimal('5000000') - Decimal('3000000')
    assert [m['type'] for m in day['ledger']['movements'] if m['type'] == 'procurement_advance'] == ['procurement_advance']
    assert Decimal(day['ledger']['cash_flow']['other_outflows']) == Decimal('3000000')
    assert day['cashier_shokh_gives']['total'] == '500000'
    assert day['cashier_shokh_gives']['gives'][0]['amount'] == '500000'
    # Счёт Шохруха у бухгалтера — те же деньги (одна формула, shokh.store.pocket_position).
    assert Decimal(day['shoh_pocket']['pocket']) == Decimal('3920000')
    # Кассир: выдача уменьшает передачу вместе с расходами.
    assert till_totals(c.app.state, DAY).cash_out == Decimal('500000')


def test_founder_counts_till_gives_once(c):
    accountant_cash(c)
    c.post('/api/accountant/procurement', json={'date': DAY.isoformat(), 'recipient': 'Шох',
                                                'purpose': 'закуп', 'amount': '3000000'})
    c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': '500000'})
    spending = cabinet.founder_spending(c.app.state, DAY)
    categories = {row['label']: row['amount'] for row in spending['expenses']['categories']}
    assert categories['Закуп · наличные Шоху'] == '3500000.00'
    assert spending['shokh']['given'] == '3500000.00'
    assert spending['shokh']['given_from_till'] == '500000.00'
    # Денежные потоки бухгалтера — только его выдача: передача кассира уже меньше.
    flows = c.app.state.accountant_finance.cash_flows_between(DAY, DAY)
    assert sum(Decimal(row['amount']) for row in flows if row['type'] == 'procurement_advance') == Decimal('3000000')
    assert not [row for row in flows if row['amount'] == '500000']


def test_founder_handover_check_matches_the_cashier_press(c):
    import asyncio
    snapshot = till_day(c)
    assert hand_over(c, snapshot, '850000').status_code == 201

    class Iiko:
        async def load(self, day):
            return snapshot

    c.app.state.iiko = Iiko()
    c.app.state.settings = replace(c.app.state.settings, login='l', password='p', store_id=1)

    class Request:
        app = c.app
        state = type('S', (), {'request_id': 'test'})()

    cashier, error = asyncio.run(cabinet.cashier_day(Request(), DAY))
    assert error is None and cashier['expected_handover'] == '850000.00'


def test_removing_a_till_give_that_shokh_already_spent_is_refused(c):
    accountant_cash(c)
    give = c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': '500000'}).json()['give']
    shokh = c.app.state.shokh
    purchase = shokh.add_purchase(DAY, datetime.now(TZ), point='Базар', item='Мясо', unit='кг',
                                  quantity='10', price='90000')
    shokh.accept_with_finance(purchase['id'], DAY, datetime.now(TZ), c.app.state.accountant_finance)
    refused = c.delete(f'/api/cashier/shokh/{give["id"]}', params={'date': DAY.isoformat()})
    assert refused.status_code == 409 and 'отчитался' in refused.json()['detail']
    assert c.get('/api/cashier/shokh', params={'date': DAY.isoformat()}).json()['total'] == '500000'


def test_till_give_is_validated_and_removed_by_day(c):
    for amount in ('0', '-5', 'abc', '1.234'):
        assert c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': amount}).status_code == 422
    give = c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': '250000'}).json()['give']
    other_day = (DAY - timedelta(days=1)).isoformat()
    assert c.delete(f'/api/cashier/shokh/{give["id"]}', params={'date': other_day}).status_code == 404
    assert c.delete(f'/api/cashier/shokh/{give["id"]}', params={'date': DAY.isoformat()}).status_code == 204
    assert c.get('/api/cashier/shokh', params={'date': DAY.isoformat()}).json()['gives'] == []
    audit = c.app.state.accountant_finance.audit_entries(entity_type='cashier_shokh_give')
    assert [entry['action'] for entry in audit] == ['create', 'delete']


def test_till_give_does_not_block_undoing_the_handover(c):
    snapshot = till_day(c, give=None)
    assert hand_over(c, snapshot, '1150000').status_code == 201
    c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': '100000'})
    assert c.delete('/api/cashier/handover', params={'date': DAY.isoformat()}).status_code == 204


def test_cashier_export_lists_the_till_give_as_an_expense(c):
    snapshot = till_day(c)
    shokh = c.get('/api/cashier/shokh', params={'date': DAY.isoformat()}).json()
    response = c.get('/api/cashier/export', params={'date': DAY.isoformat(), 'snapshot_id': snapshot.id,
                                                   'shokh_revision': shokh['revision']})
    assert response.status_code == 200
    book = openpyxl.load_workbook(BytesIO(response.content))
    assert book['отчет']['B28'].value == 400000   # такси + Шох
    assert book['отчет']['B29'].value == 850000   # к передаче — как на экране
    names = [book['Расходы'].cell(row, 1).value for row in range(8, 10)]
    assert names[0] == 'Такси' and names[1].startswith('Шоху на закуп · ')
    c.post('/api/cashier/shokh', json={'date': DAY.isoformat(), 'amount': '1'})
    assert c.get('/api/cashier/export', params={'date': DAY.isoformat(), 'snapshot_id': snapshot.id,
                                               'shokh_revision': shokh['revision']}).status_code == 409


# ── Доллары сразу в сейф ─────────────────────────────────────────────────

def test_usd_deposits_feed_the_accountant_safe(c):
    finance = c.app.state.accountant_finance
    finance.reserve_entry(min(DAY - timedelta(days=1), ACCOUNTING_START), 'usd', 'opening', '3330', 'Пересчёт сейфа')
    first = c.post('/api/cashier/usd-deposits', json={'date': DAY.isoformat(), 'amount': '120'})
    assert first.status_code == 201, first.text
    assert first.json()['safe_balance'] == '3450'
    second = c.post('/api/cashier/usd-deposits', json={'date': DAY.isoformat(), 'amount': '30,5'}).json()
    assert second['total'] == '150.5' and len(second['deposits']) == 2
    assert all(row['created_at'].endswith('+05:00') and not row['legacy'] for row in second['deposits'])
    usd = accountant_day(c, DAY)['reserves']['usd']
    assert Decimal(usd['balance']) == Decimal('3480.5')
    assert {entry['source'] for entry in usd['entries']} == {'cashier'}
    # Передача сумов от долларов не меняется.
    assert till_totals(c.app.state, DAY).cash_out == 0
    # Удалить можно, пока доллары не выданы из сейфа.
    deposit_id = second['deposit']['id']
    assert c.delete(f'/api/cashier/usd-deposits/{deposit_id}', params={'date': DAY.isoformat()}).status_code == 204
    assert usd_day(finance, DAY)['total'] == '120'
    finance.reserve_entry(DAY, 'usd', 'withdrawal', '3450', 'Выдано собственнику')
    blocked = c.delete(f'/api/cashier/usd-deposits/{first.json()["deposit"]["id"]}', params={'date': DAY.isoformat()})
    assert blocked.status_code == 409 and 'выданы из сейфа' in blocked.json()['detail']
    for amount in ('0', '-1', '1.234', 'много'):
        assert c.post('/api/cashier/usd-deposits', json={'date': DAY.isoformat(), 'amount': amount}).status_code == 422


def test_usd_without_safe_count_is_unknown_not_zero(c):
    data = c.post('/api/cashier/usd-deposits', json={'date': DAY.isoformat(), 'amount': '50'}).json()
    assert data['safe_balance'] is None and data['total'] == '50'


def test_old_daily_usd_sums_become_one_deposit_without_changing_the_safe(make_app):
    app = make_app()
    old_day = (DAY - timedelta(days=5)).isoformat()
    with closing(app.state.usd_rates._open()) as connection, connection:
        connection.execute('INSERT INTO cashier_usd_balances(day, amount) VALUES (?, ?)', (old_day, '1250.50'))
        connection.execute('INSERT INTO cashier_usd_balances(day, amount) VALUES (?, ?)',
                           ((DAY - timedelta(days=4)).isoformat(), '0'))
    finance = app.state.accountant_finance
    finance.reserve_entry(DAY - timedelta(days=10), 'usd', 'opening', '1000', 'Пересчёт')
    assert migrate_legacy_usd(app.state.usd_rates, finance) == 2
    day = DAY - timedelta(days=5)
    migrated = usd_day(finance, day)
    assert [(row['amount'], row['legacy'], row['created_at']) for row in migrated['deposits']] == [('1250.50', True, None)]
    assert migrated['safe_balance'] == '1000'  # сейф бухгалтера задним числом не изменился
    assert usd_day(finance, DAY - timedelta(days=4))['deposits'] == []
    # «Перезапуск»: ничего не дублируется, удалённый перенос не воскресает.
    again = make_app()
    assert usd_day(again.state.accountant_finance, day)['total'] == '1250.50'
    with client_for(again) as c:
        deposit_id = migrated['deposits'][0]['id']
        assert c.delete(f'/api/cashier/usd-deposits/{deposit_id}', params={'date': day.isoformat()}).status_code == 204
    third = make_app()
    assert usd_day(third.state.accountant_finance, day)['deposits'] == []
    assert migrate_legacy_usd(third.state.usd_rates, third.state.accountant_finance) == 0


def test_founder_chat_reads_usd_and_till_gives_from_one_place(c):
    add_usd_deposit(c.app.state.accountant_finance, DAY, '75')
    give_shokh(c.app.state.accountant_finance, DAY, '10000')
    assert usd_day(c.app.state.accountant_finance, DAY)['total'] == '75'
    assert till_totals(c.app.state, DAY).shokh == Decimal('10000')


# ── QA 5a: числа карточки передачи — с сервера; авто-строка зарплаты ─────

def test_screen_gets_ready_handover_numbers_from_the_server(c):
    """Функционал §1: формула одна, экран получает готовые числа."""
    snapshot = till_day(c)
    summary = c.get('/api/cashier/summary', params={'date': DAY.isoformat(),
                                                    'snapshot_id': snapshot.id}).json()
    assert summary['handover'] == '850000' and Decimal(summary['handover']) == expected_for(c)
    assert (summary['demo_cash'], summary['cash_prepayment']) == ('1000000', '200000')
    assert (summary['expenses'], summary['shokh'], summary['cash_out']) == ('100000', '300000', '400000')
    assert summary['receipts'] == '50000'
    # Касса за день: продажи + новые предоплаты (регистра нет) + ручные поступления.
    assert Decimal(summary['total_inflow']) == snapshot.revenue + PREPAY + Decimal('50000')
    # Доллары в сейф в передачу не входят.
    assert c.post('/api/cashier/usd-deposits', json={'date': DAY.isoformat(), 'amount': '100'}).status_code == 201
    again = c.get('/api/cashier/summary', params={'date': DAY.isoformat(), 'snapshot_id': snapshot.id}).json()
    assert again['handover'] == '850000'
    stale = c.get('/api/cashier/summary', params={'date': DAY.isoformat(), 'snapshot_id': 'f' * 32})
    assert stale.status_code == 409


def test_cashier_salary_row_is_automatic_daily_and_not_deletable(tmp_path):
    from retro.modules.cashier.expenses import seed_cashier_expense
    start = DAY - timedelta(days=3)
    seed_cashier_expense(tmp_path / 'cashier.sqlite3', start, start, 'Зарплата кассира', Decimal('350000'))
    with client_for(build_app(tmp_path)) as c:
        for day in (start, DAY):  # после засеянного дня строка появляется сама
            listed = c.get('/api/cashier/expenses', params={'date': day.isoformat()}).json()
            assert [(e['description'], e['amount'], e.get('automatic')) for e in listed['expenses']] == [
                ('Зарплата кассира', '350000', True)]
            refused = c.delete(f"/api/cashier/expenses/{listed['expenses'][0]['id']}",
                               params={'date': day.isoformat()})
            assert refused.status_code == 409 and 'удалить нельзя' in refused.json()['detail']
        before = c.get('/api/cashier/expenses', params={'date': (start - timedelta(days=1)).isoformat()}).json()
        assert before['expenses'] == []  # до начала политики строки нет
        # Ручная строка кассира удаляется как раньше; авто-строка входит в передачу.
        manual = c.post('/api/cashier/expenses', json={'date': DAY.isoformat(), 'description': 'Такси',
                                                       'amount': '20000'}).json()
        assert 'automatic' not in manual
        assert c.delete(f"/api/cashier/expenses/{manual['id']}", params={'date': DAY.isoformat()}).status_code == 204
        snapshot = snapshot_for(DAY)
        c.app.state.cache.put(snapshot)
        summary = c.get('/api/cashier/summary', params={'date': DAY.isoformat(), 'snapshot_id': snapshot.id}).json()
        assert summary['expenses'] == '350000'
        assert Decimal(summary['handover']) == DEMO_CASH + PREPAY - Decimal('350000')
    # Без настроенной политики ничего не создаётся.
    with client_for(build_app(tmp_path / 'clean')) as c:
        assert c.get('/api/cashier/expenses', params={'date': DAY.isoformat()}).json()['expenses'] == []
