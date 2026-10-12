"""Checks belong to work dates; new cash payouts belong to server today."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from test_salary_payout_dates import (c, seed, oct_, cash_days, shift_cell, shift_sheet,
                                      cell, day_json, spent_on, table)  # noqa: F401


def test_several_checks_for_one_employee_spend_today_once_and_export_actual_date(c):
    person = seed(c, cash_through=oct_(12))['ikhtiyor']
    c.clock['today'] = oct_(12)
    for work in (oct_(9), oct_(10)):
        assert shift_cell(c, work, person, '360000').status_code == 200
    assert spent_on(c, oct_(12)) == Decimal('720000')
    for day in (oct_(9), oct_(10), oct_(11)):
        assert spent_on(c, day) == 0
    cells = next(p for p in shift_sheet(c)['people'] if p['id'] == person.id)['cells']
    for work in ('2026-10-09', '2026-10-10'):
        assert cells[work]['amount'] == '360000'
        assert cells[work]['paid_days'] == ['2026-10-12']
        assert cells[work]['editable']
    response = c.get('/api/accountant/salary-day/export?month=2026-10&basis=shift')
    book = load_workbook(BytesIO(response.content))['Ведомость']
    headers = [value.value for value in book[4]]
    row = next(r for r in book.iter_rows(min_row=5) if r[1].value == person.name)
    for work in ('09.10', '10.10'):
        exported = row[headers.index('Смена ' + work)]
        assert exported.value == 360000
        assert '12.10.2026' in exported.comment.text
    assert len(table(c, 'accountant_salary_payments')) == 2


def test_today_shift_is_payable_and_uses_today_cash_not_next_day(c):
    person = seed(c)['ikhtiyor']
    response = shift_cell(c, oct_(10), person, '360000')
    assert response.status_code == 200, response.text
    assert response.json()['date'] == response.json()['work_day'] == '2026-10-10'
    assert spent_on(c, oct_(10)) == Decimal('360000')
    assert shift_cell(c, oct_(11), person, '360000').status_code == 422


def test_old_check_edits_and_retry_keep_original_cash_date_then_new_check_uses_today(c):
    person = seed(c, cash_through=oct_(12))['ikhtiyor']
    # A next-day payment from the old UI is editable alongside a delayed payout.
    assert cell(c, oct_(10), person, '360000').status_code == 200
    assert shift_cell(c, oct_(7), person, '360000').status_code == 200
    c.clock['today'] = oct_(12)
    assert shift_cell(c, oct_(7), person, '360000').json()['changed'] is False
    assert shift_cell(c, oct_(7), person, '250000', '360000').json()['date'] == '2026-10-10'
    assert shift_cell(c, oct_(9), person, '200000', '360000').json()['date'] == '2026-10-10'
    assert spent_on(c, oct_(10)) == Decimal('450000')
    assert spent_on(c, oct_(12)) == 0
    assert shift_cell(c, oct_(7), person, '0', '250000').status_code == 200
    assert spent_on(c, oct_(10)) == Decimal('200000')
    assert shift_cell(c, oct_(7), person, '360000').json()['date'] == '2026-10-12'
    assert spent_on(c, oct_(12)) == Decimal('360000')


def test_historical_shift_uses_only_current_cash_and_insufficient_cash_is_atomic(c):
    person = seed(c, opening='0', handover='0')['ikhtiyor']
    finance = c.app.state.accountant_finance
    finance.record_handover(oct_(10), Decimal('360000'))
    assert shift_cell(c, oct_(6), person, '360000').status_code == 200
    before = table(c, 'accountant_accruals'), table(c, 'accountant_salary_payments')
    refused = shift_cell(c, oct_(7), person, '360000')
    assert refused.status_code == 422, refused.text
    assert (table(c, 'accountant_accruals'), table(c, 'accountant_salary_payments')) == before
    assert day_json(c, oct_(6))['ledger']['cash_balance'] == '0'
    assert spent_on(c, oct_(10)) == Decimal('360000')


def test_concurrent_checks_do_not_duplicate_or_overwrite_a_shift(c):
    person = seed(c)['ikhtiyor']
    with ThreadPoolExecutor(max_workers=2) as pool:
        statuses = list(pool.map(lambda amount: shift_cell(c, oct_(7), person, amount).status_code,
                                 ['100000', '200000']))
    assert sorted(statuses) == [200, 409]
    assert len(table(c, 'accountant_salary_payments')) == 1
    with ThreadPoolExecutor(max_workers=2) as pool:
        statuses = list(pool.map(lambda work: shift_cell(c, work, person, '100000').status_code,
                                 [oct_(8), oct_(9)]))
    assert statuses == [200, 200]
    assert len(table(c, 'accountant_salary_payments')) == 3


def test_first_shift_stays_available_and_closed_work_month_cannot_be_changed(c):
    person = seed(c)['ikhtiyor']
    assert shift_cell(c, oct_(5), person, '360000').status_code == 200
    assert shift_cell(c, oct_(4), person, '360000').status_code == 422
    c.clock['today'] = date(2026, 11, 2)
    cash_days(c, oct_(11), c.clock['today'])
    with closing(c.app.state.accountant_finance._open()) as connection, connection:
        connection.execute('INSERT INTO accountant_month_closures '
                           '(month,last_day,closed_by,closed_at,closing_balance,snapshot) VALUES (?,?,?,?,?,?)',
                           ('2026-10', '2026-10-31', 'Бухгалтер', '2026-11-01', '0', '{}'))
    assert shift_cell(c, oct_(7), person, '360000').status_code == 409
    assert shift_cell(c, oct_(5), person, '0', '360000').status_code == 409
