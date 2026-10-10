"""Доп. выплаты и «Скачать Excel» в «Зарплата · день» (ТЗ 09.10, Б-05 и Б-07).

Контрольные примеры ТЗ: хостес Карамат (временная, 150 000) получает разовую
выплату без своей клетки; Баходиров Ихтиер (менеджер, 360 000) — выплату в
клетке, и доп. выплата за тот же день требует явного подтверждения. Деньги —
один расход в день выплаты: ведомость, журнал «Финансов дня», строка
«Зарплаты» дэшборда и Excel сходятся; в долги выплата не попадает.
На SQLite и на Postgres (RETRO_TEST_POSTGRES_URL).
"""

import json
import os
import subprocess
from contextlib import closing
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from retro.app import create_app
from retro.config import Settings
from retro.db import Database
from retro.modules.accountant import extra_payouts, routes, salary_day
from retro.modules.accountant.extra_payouts import ExtraPayoutChanged, ExtraPayoutConfirm
from retro.modules.accountant.ledger import FinanceStore, LedgerError
from retro.modules.accountant.roster import RosterStore
from retro.modules.accountant.salary_day_export import matches_query

TODAY = date(2026, 10, 10)
PAID = date(2026, 10, 9)
SHIFT = date(2026, 10, 8)
POSTGRES_URL = os.getenv('RETRO_TEST_POSTGRES_URL', '')
ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture(autouse=True)
def today(monkeypatch):
    for module in (salary_day, routes, extra_payouts):
        monkeypatch.setattr(module, 'today_tashkent', lambda: TODAY)


@pytest.fixture(params=['sqlite', 'postgres'])
def database(request, tmp_path):
    if request.param == 'sqlite':
        yield tmp_path / 'finance.sqlite3'
        return
    if not POSTGRES_URL:
        pytest.skip('RETRO_TEST_POSTGRES_URL не задан')
    import psycopg
    from psycopg import sql
    schema = 'extra_payouts_' + uuid4().hex
    with psycopg.connect(POSTGRES_URL, autocommit=True) as admin:
        admin.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(schema)))
        try:
            parts = urlsplit(POSTGRES_URL)
            query = dict(parse_qsl(parts.query), options='-csearch_path=' + schema)
            yield Database(urlunsplit(parts._replace(query=urlencode(query))))
        finally:
            admin.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(schema)))


def staff(roster, finance):
    """Люди из ТЗ. Признак «временный» ведёт кабинет менеджера (T-432) —
    здесь поле добавлено руками, как его добавит миграция."""
    people = dict(
        ihtiyor=roster.add(name='Баходиров Ихтиер', role='Менеджер', rate='360000', group_name='Управление'),
        jahongir=roster.add(name='Каримов Жахонгир', role='Менеджер', rate='360000', group_name='Управление'),
        selvina=roster.add(name='Абдулганиева Сельвина', role='Хостес', rate='360000', group_name='Встреча гостей'),
        karamat=roster.add(name='Карамат', role='Хостес', rate='150000', group_name='Встреча гостей'))
    with closing(finance._open()) as connection, connection:
        connection.execute('ALTER TABLE accountant_employees ADD COLUMN employment_type TEXT')
        connection.execute("UPDATE accountant_employees SET employment_type = 'temporary' WHERE id = ?",
                           (people['karamat'].id,))
    return people


@pytest.fixture
def stores(database):
    roster = RosterStore(database)
    finance = FinanceStore(database)
    for offset in range(9):
        finance.record_handover(date(2026, 10, 2) + timedelta(days=offset), Decimal(0))
    finance.set_cash_opening(date(2026, 10, 2), '2000000', 'Начало')
    return finance, roster, staff(roster, finance)


def add(finance, person, amount='150000', note='Подмена хостес', **values):
    values = dict(dict(work_day=SHIFT, paid_day=PAID, confirm=False, by='buh'), **values)
    return finance.add_extra_payout(employee_id=person.id, amount=amount, note=note, **values)


def count(finance, table):
    with closing(finance._open()) as connection:
        return connection.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0]


def test_karamat_payout_is_one_expense_on_the_payout_day_and_reaches_every_total(stores):
    finance, _, people = stores
    item = add(finance, people['karamat'])
    assert (item['name'], item['role'], item['temporary']) == ('Карамат', 'Хостес', True)
    assert (item['work_day'], item['paid_day'], item['amount'], item['created_by']) == (
        '2026-10-08', '2026-10-09', '150000', 'buh')
    # Ведомость: запись периода, человек помечен временным; клетка своей суммы не меняет.
    data = finance.salary_day_month(date(2026, 10, 2), date(2026, 10, 31))
    assert [(x['employee_id'], x['amount'], x['editable']) for x in data['extras']] == [
        (people['karamat'].id, '150000', True)]
    karamat = next(p for p in data['people'] if p['id'] == people['karamat'].id)
    assert karamat['temporary'] is True and karamat['cells'][PAID.isoformat()]['amount'] == '0'
    # Журнал и касса 09.10: один расход статьи «Доп. выплаты»; 08.10 не тронут.
    summary = finance.daily_summary(PAID, None)
    extra_rows = [m for m in summary['movements'] if m.get('item_code') == 'salary_extra_payout']
    assert [(m['type'], m['amount']) for m in extra_rows] == [('other_expense', '150000')]
    assert summary['cash_balance'] == Decimal('1850000')
    flow = finance.day_flow(PAID)
    # Строка «Зарплаты · сменные и оклады» дэшборда = salary + monthly.
    assert (flow['salary'], flow['monthly'], flow['other'], flow['outflows']) == (
        Decimal('150000'), Decimal(0), Decimal(0), Decimal('150000'))
    assert finance.day_flow(SHIFT)['outflows'] == Decimal(0)
    # Не долг и не оклад.
    assert summary['salary_debt'] == Decimal(0)
    assert finance.reserves(PAID)['monthly']['paid'] == '0'
    assert finance.extra_payouts(PAID, PAID)[0]['note'] == 'Подмена хостес'


def test_cell_payment_the_same_day_asks_for_confirmation_and_nothing_is_written(stores):
    finance, _, people = stores
    finance.set_salary_day_cell(PAID, people['ihtiyor'].id, '360000', '0')
    before = (count(finance, 'accountant_movements'), count(finance, 'accountant_finance_audit'))
    with pytest.raises(ExtraPayoutConfirm, match='уже отмечена выплата в клетке: 09.10 — 360 000 сум'):
        add(finance, people['ihtiyor'], '50000', 'Премия за банкет')
    assert (count(finance, 'accountant_movements'), count(finance, 'accountant_finance_audit')) == before
    add(finance, people['ihtiyor'], '50000', 'Премия за банкет', confirm=True)
    # Ведомость: 360 000 в клетке + 50 000 доп. — 410 000 за 09.10.
    data = finance.salary_day_month(date(2026, 10, 2), date(2026, 10, 31))
    ihtiyor = next(p for p in data['people'] if p['id'] == people['ihtiyor'].id)
    assert ihtiyor['cells'][PAID.isoformat()]['amount'] == '360000'
    assert sum(Decimal(x['amount']) for x in data['extras'] if x['employee_id'] == ihtiyor['id']) == 50000
    assert finance.day_flow(PAID)['salary'] == Decimal('410000')
    # Повтор той же записи — снова вопрос, а не вторая выдача.
    with pytest.raises(ExtraPayoutConfirm, match='Такая доп. выплата уже записана'):
        add(finance, people['ihtiyor'], '50000', 'Ещё раз')
    assert count(finance, 'accountant_extra_payouts') == 1
    # Сельвину в клетке не отмечали — её выплата пишется сразу.
    add(finance, people['selvina'], '100000', 'Доплата')
    assert count(finance, 'accountant_extra_payouts') == 2


def test_same_shift_paid_in_the_cell_on_another_day_also_warns(stores):
    finance, _, people = stores
    finance.set_salary_day_cell(SHIFT + timedelta(days=1), people['jahongir'].id, '360000', '0')
    with pytest.raises(ExtraPayoutConfirm, match='смена 08.10'):
        add(finance, people['jahongir'], '20000', 'Такси', paid_day=TODAY)


def test_edit_and_delete_keep_audit_and_history_survives_archive(stores):
    finance, roster, people = stores
    item = add(finance, people['karamat'])
    with pytest.raises(ExtraPayoutChanged):
        finance.update_extra_payout(item['id'], amount='200000', work_day=SHIFT, note='x', expected_amount='1')
    unchanged = finance.update_extra_payout(item['id'], amount='150000', work_day=SHIFT, note='Подмена хостес')
    assert unchanged['changed'] is False
    edited = finance.update_extra_payout(item['id'], amount='200000', work_day=date(2026, 10, 7),
                                         note='Две подмены', expected_amount='150000', by='buh2')
    assert (edited['amount'], edited['work_day'], edited['note'], edited['updated_by']) == (
        '200000', '2026-10-07', 'Две подмены', 'buh2')
    assert edited['paid_day'] == PAID.isoformat()
    assert finance.day_flow(PAID)['salary'] == Decimal('200000')
    movement = next(m for m in finance.daily_summary(PAID, None)['movements'] if m.get('item_code') == 'salary_extra_payout')
    assert movement['description'] == 'Доп. выплата · Карамат · смена 07.10 · Две подмены'
    # Временная ушла: запись и строка в ведомости остаются по имени.
    roster.delete(people['karamat'].id)
    data = finance.salary_day_month(date(2026, 10, 2), date(2026, 10, 31))
    gone = next(p for p in data['people'] if p['id'] == people['karamat'].id)
    assert (gone['name'], gone['archived'], gone['temporary']) == ('Карамат', True, True)
    assert data['extras'][0]['editable'] is True
    finance.delete_extra_payout(item['id'], by='buh')
    assert finance.extra_payouts(PAID, PAID) == []
    assert finance.day_flow(PAID)['outflows'] == Decimal(0)
    audit = finance.audit_entries(entity_type='extra_payout')
    assert [entry['action'] for entry in audit] == ['create', 'update', 'delete']
    assert audit[-1]['before']['amount'] == '200000' and audit[-1]['after']['deleted_by'] == 'buh'
    assert [entry['action'] for entry in finance.audit_entries(entity_type='movement')
            if entry['entity_id'] == str(item['movement_id'])] == ['create', 'update', 'delete']


def test_days_cash_and_closed_month_are_checked_like_the_cell(database):
    roster = RosterStore(database)
    finance = FinanceStore(database)
    people = staff(roster, finance)
    for offset in range(9):
        finance.record_handover(date(2026, 10, 2) + timedelta(days=offset), Decimal(0))
    finance.set_cash_opening(date(2026, 10, 2), '100000', 'Начало')
    karamat = people['karamat']
    for values, text in [(dict(paid_day=TODAY + timedelta(days=1)), 'будущим днём'),
                         (dict(paid_day=date(2026, 10, 1), work_day=date(2026, 9, 30)), '02.10.2026'),
                         (dict(work_day=TODAY), 'позже дня выплаты'),
                         (dict(work_day=date(2026, 9, 30)), '01.10.2026'),
                         (dict(), 'недостаточно')]:
        with pytest.raises(LedgerError, match=text):
            add(finance, karamat, **values)
    with pytest.raises(LedgerError, match='не найден'):
        finance.add_extra_payout(employee_id=999, work_day=SHIFT, paid_day=PAID, amount='1', note='x')
    with pytest.raises(LedgerError, match='назначение'):
        add(finance, karamat, '1000', ' ')
    assert count(finance, 'accountant_extra_payouts') == count(finance, 'accountant_movements') == 0
    item = add(finance, karamat, '100000')
    # Расход съел весь остаток: увеличить нельзя, уменьшить можно.
    with pytest.raises(LedgerError, match='недостаточно'):
        finance.update_extra_payout(item['id'], amount='100001', work_day=SHIFT, note='Подмена хостес')
    finance.update_extra_payout(item['id'], amount='90000', work_day=SHIFT, note='Подмена хостес')
    with closing(finance._open()) as connection, connection:
        connection.execute('INSERT INTO accountant_month_closures VALUES (?,?,?,?,?,?)',
                           ('2026-10', '2026-10-31', '0', '2026-11-01', 'buh', '{}'))
    for action in (lambda: add(finance, karamat, '1000'),
                   lambda: finance.update_extra_payout(item['id'], amount='1000', work_day=SHIFT, note='x'),
                   lambda: finance.delete_extra_payout(item['id'])):
        with pytest.raises(LedgerError, match='Месяц закрыт'):
            action()


def test_journal_and_debts_cannot_create_or_rewrite_an_extra_payout(stores):
    finance, _, people = stores
    item = add(finance, people['karamat'])
    for action in (lambda: finance.update_movement(item['movement_id'], PAID, 'admin_other', 'Временные', '1'),
                   lambda: finance.delete_operation('movement', item['movement_id'], PAID)):
        with pytest.raises(LedgerError, match='Зарплата · день'):
            action()
    other = finance.add_expense(PAID, 'admin_other', 'Канцтовары', '1000')
    for action in (lambda: finance.add_expense(PAID, 'salary_extra_payout', 'Без сотрудника', '1000'),
                   lambda: finance.record_debt(PAID, 'salary_extra_payout', 'Без сотрудника', '1000', '0'),
                   lambda: finance.update_movement(other, PAID, 'salary_extra_payout', 'Канцтовары', '1000')):
        with pytest.raises(LedgerError, match='с сотрудником и датой смены'):
            action()
    assert finance.extra_payouts(PAID, PAID)[0]['amount'] == '150000'


# ── Через API: идемпотентность, вопрос о повторе, журнал дня, Excel ──────────

@pytest.fixture
def client(tmp_path):
    app = create_app(Settings(data_dir=tmp_path, manual_handover_only=True))
    with TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)) as test_client:
        finance = test_client.app.state.accountant_finance
        # Начальный остаток 02.10 приложение ставит само (opening_migration).
        for offset in range(9):
            finance.record_handover(date(2026, 10, 2) + timedelta(days=offset), Decimal('500000'))
        test_client.people = staff(test_client.app.state.accountant_roster, finance)
        yield test_client


def post(c, person, key=None, **values):
    body = dict(employee_id=person.id, work_day=SHIFT.isoformat(), paid_day=PAID.isoformat(),
                amount='150000', note='Подмена хостес') | values
    return c.post('/api/accountant/salary-day/extra', json=body,
                  headers={'Idempotency-Key': key} if key else {})


def test_api_retry_with_the_same_key_writes_once_and_overlap_is_a_question(client):
    karamat, ihtiyor = client.people['karamat'], client.people['ihtiyor']
    key = str(uuid4())
    first, again = post(client, karamat, key), post(client, karamat, key)
    assert first.status_code == again.status_code == 201
    assert first.json() == again.json()
    assert len(client.app.state.accountant_finance.extra_payouts(PAID, PAID)) == 1
    client.put('/api/accountant/salary-day/cell', json=dict(
        date=PAID.isoformat(), employee_id=ihtiyor.id, amount='360000', expected_amount='0'))
    question = post(client, ihtiyor, amount='50000', note='Премия')
    assert question.status_code == 409
    assert question.json()['detail']['confirm'] is True
    assert 'уже отмечена выплата в клетке' in question.json()['detail']['message']
    assert post(client, ihtiyor, amount='50000', note='Премия', confirm=True).status_code == 201
    # Журнал дня: обе доп. выплаты пришли списком для строки «Доп. выплаты · 2 чел.».
    day = client.get('/api/accountant/day', params={'date': PAID.isoformat()}).json()
    assert sorted(x['name'] for x in day['extra_payouts']) == ['Баходиров Ихтиер', 'Карамат']
    assert Decimal(day['ledger']['day_flow']['salary']) == Decimal('560000')
    assert Decimal(day['ledger']['salary_debt']) == 0
    # Правка и удаление.
    extra = next(x for x in day['extra_payouts'] if x['name'] == 'Карамат')
    edited = client.put(f'/api/accountant/salary-day/extra/{extra["id"]}', json=dict(
        work_day=SHIFT.isoformat(), amount='160000', note='Подмена', expected_amount='150000'))
    assert edited.status_code == 200 and edited.json()['amount'] == '160000'
    stale = client.put(f'/api/accountant/salary-day/extra/{extra["id"]}', json=dict(
        work_day=SHIFT.isoformat(), amount='170000', note='Подмена', expected_amount='150000'))
    assert stale.status_code == 409
    assert client.delete(f'/api/accountant/salary-day/extra/{extra["id"]}').status_code == 204
    assert client.delete(f'/api/accountant/salary-day/extra/{extra["id"]}').status_code == 404


def sheet_values(sheet):
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_excel_matches_the_screen_totals_and_prints_with_headers(client):
    people = client.people
    put = lambda person, day, amount: client.put('/api/accountant/salary-day/cell', json=dict(
        date=day.isoformat(), employee_id=person.id, amount=amount, expected_amount='0'))
    put(people['ihtiyor'], date(2026, 10, 8), '360000')
    put(people['ihtiyor'], PAID, '360000')
    put(people['jahongir'], PAID, '300000')
    assert post(client, people['karamat']).status_code == 201
    month = client.get('/api/accountant/salary-day/month', params={'month': '2026-10'}).json()
    response = client.get('/api/accountant/salary-day/export', params={'month': '2026-10'})
    assert response.status_code == 200
    assert 'Retro-salary-day-2026-10.xlsx' in response.headers['content-disposition']
    book = load_workbook(BytesIO(response.content))
    sheet = book['Ведомость']
    rows = sheet_values(sheet)
    assert rows[0][0] == 'RETRO MILLIY · Зарплата · день · октябрь 2026'
    assert 'Вся ведомость, 4 сотрудника.' in rows[1][0] and '02.10.2026 – 10.10.2026' in rows[1][0]
    head = rows[3]
    assert head[:4] == ['№', 'Сотрудник', 'Должность', 'Ставка']
    assert head[4] == 'Выплата 02.10\nсмена 01.10' and head[11] == 'Выплата 09.10\nсмена 08.10'
    assert head[-2:] == ['Доп. выплаты', 'Итого'] and len(head) == 4 + 9 + 2
    body = {row[1]: row for row in rows[4:8]}
    assert body['Баходиров Ихтиер'][2] == 'Менеджер' and body['Баходиров Ихтиер'][3] == 360000
    assert body['Карамат'][2] == 'Хостес · временный'
    # Итоги — те же, что у экрана: клетки + доп. выплаты по дню выплаты.
    for person in month['people']:
        line = body[person['name']]
        cells = sum(Decimal(person['cells'][day]['amount']) for day in month['days'] if day <= month['today'])
        extra = sum(Decimal(x['amount']) for x in month['extras'] if x['employee_id'] == person['id'])
        assert Decimal(str(line[-1])) == cells + extra
        assert Decimal(str(line[-2] or 0)) == extra
    total, extras_row = rows[8], rows[9]
    assert total[1] == 'Итого за день' and extras_row[1] == 'в т.ч. доп. выплаты'
    assert total[11] == 360000 + 300000 + 150000 and extras_row[11] == 150000
    assert total[-1] == 360000 * 2 + 300000 + 150000
    assert isinstance(total[-1], (int, float)) and sheet.cell(9, 12).number_format == '#,##0'
    # Печать: альбом, по ширине страницы, шапка на каждой странице, имена закреплены.
    assert sheet.page_setup.orientation == 'landscape'
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert sheet.print_title_rows == '$1:$4' and sheet.freeze_panes == 'E5'
    extras = sheet_values(book['Доп. выплаты'])
    assert extras[3][:7] == ['№', 'Дата выплаты', 'Дата смены', 'Сотрудник', 'Должность', 'Временный', 'Сумма']
    assert extras[4][:8] == [1, '09.10.2026', '08.10.2026', 'Карамат', 'Хостес', 'да', 150000, 'Подмена хостес']
    assert extras[5][1] == 'Итого' and extras[5][6] == 150000


def test_excel_selection_is_named_in_the_header(client):
    client.app.state.accountant_finance  # noqa: B018 — стенд уже собран фикстурой
    assert post(client, client.people['karamat']).status_code == 201
    response = client.get('/api/accountant/salary-day/export',
                          params={'month': '2026-10', 'group': 'Встреча гостей', 'q': 'karamat'})
    assert 'selection' in response.headers['content-disposition']
    rows = sheet_values(load_workbook(BytesIO(response.content))['Ведомость'])
    assert 'Выборка: группа «Встреча гостей», поиск «karamat» — 1 из 4 сотрудников.' in rows[1][0]
    assert rows[4][1] == 'Карамат' and rows[5][1] == 'Итого за день'


def test_search_in_the_file_is_the_same_as_on_screen():
    cases = [('Alijon', 'Алижон'), ('ihtiyor', 'Ихтиёр'), ('Shoxrux', 'Шохрух'), ('Shokhrukh', 'Шохрух'),
             ('повар', 'Повар миллий'), ('ПОВАР', 'повар'), ('selvina', 'Сельвина'), ('karamat', 'Карамат'),
             ('жахонгир', 'Каримов Жахонгир'), ('zhahongir', 'Жахонгир'), ('tandir', 'Повар тандыр'),
             ('xyz', 'Карамат'), ('', 'Карамат'), ('ё', 'Ёлка'), ('qo‘zi', 'Қўзи')]
    script = ("const L=require(process.argv[1]);const cases=JSON.parse(process.argv[2]);"
              "console.log(JSON.stringify(cases.map(([q,f])=>L.matchesQuery(q,f))))")
    result = subprocess.run(['node', '-e', script, str(ROOT / 'retro/static/employees-logic.js'), json.dumps(cases)],
                            check=True, capture_output=True, text=True)
    assert json.loads(result.stdout) == [matches_query(query, field) for query, field in cases]
