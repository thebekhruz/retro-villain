from datetime import date, datetime, timedelta
from decimal import Decimal
from dataclasses import replace
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from retro.app import create_app
from retro.config import Settings
from retro.integrations.hikvision import HikvisionEvent, HikvisionPerson
from retro.modules.cashier.service import DataError, Payment, TZ, demo_snapshot, today_tashkent
from retro.modules.cashier.expenses import seed_cashier_expense


DAY = date(2026, 9, 16)


def test_safe_workflow_and_historical_balances_through_api(tmp_path):
    with demo_client(tmp_path) as client:
        client.app.state.cache.put(replace(demo_snapshot(DAY), demo=False,
                                           payments=(Payment('Демо', Decimal('1350000')),)))
        payload = dict(date=DAY.isoformat(), account='dividends', kind='opening', amount='0', note='Начало')
        assert client.post('/api/accountant/reserves', json=payload).status_code == 201
        payload.update(kind='transfer', amount='300000')
        assert client.post('/api/accountant/reserves', json=payload).status_code == 201
        payload.update(kind='withdrawal', amount='100000')
        assert client.post('/api/accountant/reserves', json=payload).status_code == 201
        data = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert data['reserves']['dividends']['balance'] == '200000'
        assert data['ledger']['cash_balance'] == '700000'
        assert client.post('/api/accountant/reserves', json=dict(payload, date='2099-01-01')).status_code == 422
        assert client.post('/api/accountant/reserves', json=dict(payload, kind='deposit')).status_code == 422
        assert client.post('/api/accountant/reserves', json=payload,
                           headers={'host': 'public.trycloudflare.com'}).status_code == 403


def test_monthly_plan_and_usd_do_not_require_iiko_but_cash_transfer_does(tmp_path):
    with demo_client(tmp_path) as client:
        payload = dict(date=DAY.isoformat(), account='usd', kind='opening', amount='50', note='Фактические USD')
        assert client.post('/api/accountant/reserves', json=payload).status_code == 201
        assert client.post('/api/accountant/reserves', json=dict(payload, account='dividends', amount='0')).status_code == 201
        assert client.post('/api/accountant/reserves', json=dict(payload, account='dividends', kind='transfer')).status_code == 409
        assert client.post('/api/accountant/monthly-plan', json={
            'date': DAY.isoformat(), 'amount': '1000000', 'note': 'Оклады из кассы'}).status_code == 201
        data = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert data['reserves']['usd']['balance'] == '50'
        assert data['reserves']['monthly']['balance'] == '1000000'
        assert data['ledger']['cash_balance'] is None


def test_monthly_salary_total_is_calculated_from_monthly_register(tmp_path):
    with demo_client(tmp_path) as client:
        response = client.post('/api/accountant/monthly-employees', json={
            'name': 'Администратор', 'role': 'Управление', 'salary': '5000000',
            'schedule': '5/2', 'card': '3000000', 'cash': '2000000',
            'advances': '0', 'remaining': '0'})
        assert response.status_code == 201
        data = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert data['reserves']['monthly']['total'] == '5000000'
        assert data['monthly_employees'][0]['name'] == 'Администратор'
        invalid = client.patch(
            f"/api/accountant/monthly-employees/{response.json()['employee']['id']}",
            json=dict(response.json()['employee'], cash='NaN'))
        assert invalid.status_code == 422


def test_manual_cashier_income_is_used_without_iiko(tmp_path):
    with demo_client(tmp_path) as client:
        client.app.state.settings = replace(client.app.state.settings, manual_handover_only=True)
        response = client.post('/api/accountant/handover', json={
            'date': DAY.isoformat(), 'amount': '750000', 'note': 'Передано кассиром'})
        assert response.status_code == 201
        data = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert data['expected_cashier'] == '750000'
        assert data['ledger']['cash_balance'] == '750000'


def test_manual_cashier_income_can_be_updated_and_deleted(tmp_path):
    with demo_client(tmp_path) as client:
        client.app.state.settings = replace(client.app.state.settings, manual_handover_only=True)
        payload = {'date': DAY.isoformat(), 'amount': '750000', 'note': 'Передано'}
        assert client.post('/api/accountant/handover', json=payload).status_code == 201
        assert client.put('/api/accountant/handover/' + DAY.isoformat(),
                          json={'date': DAY.isoformat(), 'amount': '800000', 'note': 'Исправлено'}).status_code == 200
        assert client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['expected_cashier'] == '800000'
        assert client.delete('/api/accountant/handover/' + DAY.isoformat()).status_code == 204
        assert client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['expected_cashier'] is None


def test_accountant_income_and_expense_can_be_updated(tmp_path):
    with demo_client(tmp_path) as client:
        client.app.state.settings = replace(client.app.state.settings, manual_handover_only=True)
        assert client.post('/api/accountant/handover', json={
            'date': DAY.isoformat(), 'amount': '1000000', 'note': 'Передано'}).status_code == 201
        income = client.post('/api/accountant/incomes', json={
            'date': DAY.isoformat(), 'item_code': 'income_other',
            'note': 'Старое назначение', 'amount': '200000'}).json()
        expense = client.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'ops_rent',
            'note': 'Старая аренда', 'amount': '300000'}).json()
        assert client.put(f"/api/accountant/operations/movement/{income['id']}", json={
            'date': DAY.isoformat(), 'item_code': 'income_other',
            'note': 'Новое назначение', 'amount': '250000'}).status_code == 200
        assert client.put(f"/api/accountant/operations/movement/{expense['id']}", json={
            'date': DAY.isoformat(), 'item_code': 'ops_rent',
            'note': 'Новая аренда', 'amount': '150000'}).status_code == 200
        movements = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['ledger']['movements']
        assert next(item for item in movements if item['id'] == income['id'])['amount'] == '250000'
        assert next(item for item in movements if item['id'] == expense['id'])['description'] == 'Аренда помещения · Новая аренда'


def test_manual_mode_does_not_carry_balance_into_historical_dates(tmp_path):
    with demo_client(tmp_path) as client:
        client.app.state.settings = replace(client.app.state.settings, manual_handover_only=True)
        first = DAY - timedelta(days=1)
        client.post('/api/accountant/handover', json={'date': first.isoformat(), 'amount': '900000', 'note': 'Вчера'})
        client.post('/api/accountant/handover', json={'date': DAY.isoformat(), 'amount': '100000', 'note': 'Сегодня'})
        historical = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert historical['ledger']['cash_flow']['opening_balance'] == '0'


def test_verified_accountant_start_reconciles_september_report_and_carries_forward(tmp_path, monkeypatch):
    monkeypatch.setattr('retro.modules.accountant.routes.today_tashkent', lambda: date(2026, 9, 18))
    with demo_client(tmp_path) as client:
        client.app.state.settings = replace(client.app.state.settings, manual_handover_only=True)
        finance_day = date(2026, 9, 17)
        next_day = finance_day + timedelta(days=1)
        assert client.post('/api/accountant/handover', json={
            'date': '2026-09-15', 'amount': '15089000', 'note': 'Старый тестовый день'}).status_code == 201
        assert client.post('/api/accountant/handover', json={
            'date': finance_day.isoformat(), 'amount': '15992000',
            'note': 'Касса за 16.09'}).status_code == 201
        assert client.post('/api/accountant/cash-opening', json={
            'date': finance_day.isoformat(), 'amount': '104000',
            'note': 'Остаток на 17.09 по отчёту Лины'}).status_code == 201
        assert client.post('/api/accountant/expenses', json={
            'date': finance_day.isoformat(), 'item_code': 'salary_staff',
            'note': 'ЗП персонал', 'amount': '15862000'}).status_code == 201

        day = client.get('/api/accountant/day', params={'date': finance_day.isoformat()}).json()
        assert day['ledger']['cash_flow']['opening_balance'] == '104000'
        assert day['ledger']['cash_flow']['received_from_cashier'] == '15992000'
        assert day['ledger']['cash_flow']['closing_balance'] == '234000'
        assert day['ledger']['movements'][1]['description'] == 'Касса за 16.09.2026'

        tomorrow = client.get('/api/accountant/day', params={'date': next_day.isoformat()}).json()
        assert tomorrow['ledger']['cash_flow']['opening_balance'] == '234000'
        assert tomorrow['ledger']['cash_balance'] is None


def demo_client(tmp_path):
    seed_cashier_expense(
        tmp_path / 'cashier.sqlite3', date(2026, 1, 1), date(2026, 12, 31),
        'Зарплата', Decimal('350000'))
    app = create_app(Settings(), expense_db_path=tmp_path / 'cashier.sqlite3',
                     accountant_db_path=tmp_path / 'accountant-demo.sqlite3')
    source = tmp_path / 'roster.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'ЗП'
    for row in range(5, 25):
        sheet.cell(row, 1, row - 4)
        sheet.cell(row, 2, f'Сотрудник {row}')
        sheet.cell(row, 3, 'техперсонал' if row % 2 else 'официант')
        sheet.cell(row, 4, 250000)
    book.save(source)
    app.state.accountant_roster.import_xlsx(source)
    roster = app.state.accountant_roster.list()
    app.state.accountant_roster.link_hikvision_people(tuple(
        HikvisionPerson(f'test-{employee.id}', employee.name)
        for employee in roster if employee.source_row % 19 != 0))
    for employee in app.state.accountant_roster.list():
        if employee.hikvision_id is not None:
            app.state.attendance_store.ingest(HikvisionEvent(
                'retro-main-entry', f'serial-{employee.id}', employee.hikvision_id,
                datetime(2026, 9, 16, 9, employee.id % 50, tzinfo=TZ)), employee.id)
    app.state.attendance_store.record_success(
        'retro-main-entry', at=datetime(2026, 9, 17, 0, 5, tzinfo=TZ),
        cursor_at=datetime(2026, 9, 17, 0, 5, tzinfo=TZ),
        covered_from=datetime(2026, 9, 16, 0, 0, tzinfo=TZ),
        covered_through=datetime(2026, 9, 17, 0, 0, tzinfo=TZ))
    return TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000))


def test_partial_expense_becomes_debt_and_payment_rolls_forward(tmp_path):
    with demo_client(tmp_path) as client:
        first_day = DAY - timedelta(days=1)
        next_day = DAY
        for day in (first_day, next_day):
            client.app.state.cache.put(replace(demo_snapshot(day), demo=False,
                                               payments=(Payment('Демо', Decimal('1350000')),)))
        first = client.post('/api/accountant/expenses', json={
            'date': first_day.isoformat(), 'item_code': 'salary_technical',
            'note': 'Смена Малики', 'amount': '1200000', 'paid_amount': '300000'})
        assert first.status_code == 201
        one = client.get('/api/accountant/day', params={'date': first_day.isoformat()}).json()['ledger']
        assert one['cash_balance'] == '700000'
        assert one['manual_debt_total'] == '900000'
        assert one['salary_recorded_on_day'] == '300000'
        debt_id = one['manual_debts'][0]['id']
        assert client.post('/api/accountant/debts/pay', json={
            'date': next_day.isoformat(), 'debt_id': debt_id, 'amount': '800000'}).status_code == 201
        two = client.get('/api/accountant/day', params={'date': next_day.isoformat()}).json()['ledger']
        assert two['cash_flow']['opening_balance'] == '700000'
        assert two['cash_balance'] == '900000'
        assert two['manual_debt_total'] == '100000'
        assert two['salary_recorded_on_day'] == '800000'
        assert client.post('/api/accountant/debts/pay', json={
            'date': next_day.isoformat(), 'debt_id': debt_id, 'amount': '200000'}).status_code == 422


def test_missing_handover_in_rollforward_stays_unknown(tmp_path):
    with demo_client(tmp_path) as client:
        first_day = DAY - timedelta(days=2)
        for day in (first_day, DAY):
            client.app.state.cache.put(replace(demo_snapshot(day), demo=False,
                                               payments=(Payment('Демо', Decimal('1350000')),)))
        client.get('/api/accountant/day', params={'date': first_day.isoformat()})
        missing = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert missing['ledger']['cash_balance'] is None
        assert missing['ledger']['cash_flow']['missing_day'] == (DAY - timedelta(days=1)).isoformat()


def test_verified_initial_cash_balance_is_carried_once(tmp_path):
    with demo_client(tmp_path) as client:
        first_day = DAY - timedelta(days=1)
        for day in (first_day, DAY):
            client.app.state.cache.put(replace(demo_snapshot(day), demo=False,
                                               payments=(Payment('Демо', Decimal('1350000')),)))
        client.get('/api/accountant/day', params={'date': first_day.isoformat()})
        response = client.post('/api/accountant/cash-opening', json={
            'date': first_day.isoformat(), 'amount': '500000', 'note': 'Пересчёт'} )
        assert response.status_code == 201
        assert client.post('/api/accountant/cash-opening', json={
            'date': first_day.isoformat(), 'amount': '500000', 'note': 'Повтор'}).status_code == 409
        today = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['ledger']
        assert today['cash_flow']['opening_balance'] == '1500000'
        assert today['cash_balance'] == '2500000'
        assert today['cash_opening']['note'] == 'Пересчёт'


def test_unpaid_expense_can_be_recorded_without_iiko_and_does_not_spend_cash(tmp_path):
    with demo_client(tmp_path) as client:
        response = client.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'salary_staff',
            'note': 'Смена сотрудника', 'amount': '400000', 'paid_amount': '0'})
        assert response.status_code == 201
        ledger = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['ledger']
        assert ledger['manual_debt_total'] == '400000'
        assert ledger['cash_balance'] is None
        assert ledger['salary_recorded_on_day'] == '0'
        assert not ledger['movements']
        assert client.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'salary_staff',
            'note': 'Ошибка', 'amount': '400000', 'paid_amount': '400001'}).status_code == 422


def test_day_uses_attendance_source_defaults_to_yesterday_and_does_not_touch_cashier(tmp_path):
    with demo_client(tmp_path) as client:
        default = client.get('/api/accountant/day').json()
        assert default['date'] == (today_tashkent() - timedelta(days=1)).isoformat()
        day = client.get('/api/accountant/day', params={'date': DAY.isoformat()})
        assert day.status_code == 200
        data = day.json()
        assert data['demo'] is False
        assert len(data['employees']) == 20
        assert data['payroll']['draft_total'] != '0'
        assert data['ledger']['cash_balance'] is None
        assert data['ledger']['salary_debt'] == '0'
        assert client.get('/api/cashier/expenses', params={'date': DAY.isoformat()}).json()['total'] == '350000'
        assert client.get('/api/accountant/day', params={'date': '2099-01-01'}).status_code == 422


def test_confirmed_payroll_becomes_debt_then_partial_payment_reduces_cash(tmp_path):
    with demo_client(tmp_path) as client:
        confirmed = client.post('/api/accountant/payroll/confirm', json={
            'date': DAY.isoformat(), 'approver': 'Финансы'})
        assert confirmed.status_code == 200
        assert client.post('/api/accountant/payroll/confirm', json={
            'date': DAY.isoformat(), 'approver': 'Финансы'}).json()['already_confirmed'] is True
        before = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert before['ledger']['salary_debt'] == before['payroll']['draft_total']
        assert before['ledger']['cash_balance'] is None

        received_day = (DAY + timedelta(days=1)).isoformat()
        client.app.state.cache.put(replace(demo_snapshot(DAY + timedelta(days=1)), demo=False,
                                           payments=(Payment('Демо', Decimal('550000')),)))
        accrual = before['ledger']['accruals'][0]
        payment = client.post('/api/accountant/salary-payments', json={
            'accrual_id': accrual['id'], 'date': received_day, 'amount': '100000'})
        assert payment.status_code == 201
        after = client.get('/api/accountant/day', params={'date': received_day}).json()
        assert after['ledger']['cash_balance'] == '100000'
        assert Decimal(after['ledger']['salary_debt']) == Decimal(before['ledger']['salary_debt']) - 100000
        assert client.get('/api/cashier/expenses', params={'date': received_day}).json()['total'] == '350000'


def test_unlinked_employee_gets_only_one_demo_exception_and_other_expenses_are_separate(tmp_path):
    with demo_client(tmp_path) as client:
        before = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        unlinked = next(item for item in before['employees'] if item['status'] == 'unlinked')
        exception = client.post('/api/accountant/exceptions', json={
            'date': DAY.isoformat(), 'employee_id': unlinked['employee_id'],
            'reason': 'Первый день', 'approver': 'Финансы'})
        assert exception.status_code == 201
        after = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert next(item for item in after['employees'] if item['employee_id'] == unlinked['employee_id'])['payable'] == '250000'
        assert client.post('/api/accountant/exceptions', json={
            'date': (DAY + timedelta(days=1)).isoformat(), 'employee_id': unlinked['employee_id'],
            'reason': 'Ещё день', 'approver': 'Финансы'}).status_code == 409

        client.app.state.cache.put(replace(demo_snapshot(DAY), demo=False,
                                           payments=(Payment('Демо', Decimal('850000')),)))
        assert client.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'ops_rent',
            'note': 'Аренда', 'amount': '100000'}).status_code == 201
        assert client.post('/api/accountant/procurement', json={
            'date': DAY.isoformat(), 'recipient': 'Шох', 'purpose': 'Закуп', 'amount': '150000'}).status_code == 201
        balance = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['ledger']
        assert balance['cash_balance'] == '250000'


def test_missing_rate_can_be_corrected_with_reason_and_public_host_cannot_see_roster(tmp_path):
    with demo_client(tmp_path) as client:
        employee = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['employees'][0]
        corrected = client.patch(f'/api/accountant/employees/{employee["employee_id"]}', json={
            'rate': '310000', 'group': 'Уборка', 'reason': 'Новая ставка'})
        assert corrected.status_code == 200
        updated = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert updated['employees'][0]['rate'] == '310000'
        assert client.get('/accountant', headers={'host': 'public.trycloudflare.com'}).status_code == 403
        assert client.get('/accountant/employees', headers={'host': 'public.trycloudflare.com'}).status_code == 403
        assert client.get('/api/accountant/day', headers={'host': 'public.trycloudflare.com'}).status_code == 403
        assert client.get('/api/accountant/employees/export', params={
            'date': DAY.isoformat(), 'scope': 'all'},
            headers={'host': 'public.trycloudflare.com'}).status_code == 403


def test_employee_registry_can_edit_name_role_and_salary(tmp_path):
    with demo_client(tmp_path) as client:
        employee = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['employees'][0]
        updated = client.patch(f'/api/accountant/employees/{employee["employee_id"]}', json={
            'name': 'Новое имя', 'role': 'официант', 'rate': '275000',
            'reason': 'Обновление реестра'})
        assert updated.status_code == 200
        person = next(item for item in client.get('/api/accountant/day',
                         params={'date': DAY.isoformat()}).json()['employees']
                       if item['employee_id'] == employee['employee_id'])
        assert person['name'] == 'Новое имя'
        assert person['role'] == 'официант'
        assert person['group'] == 'Обслуживание зала'
        assert person['rate'] == '275000'


def test_employee_registry_can_add_worker_and_change_group(tmp_path):
    with demo_client(tmp_path) as client:
        employee = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['employees'][0]
        changed = client.patch(f'/api/accountant/employees/{employee["employee_id"]}', json={
            'name': employee['name'], 'role': employee['role'], 'group': 'Бар',
            'rate': employee['rate'], 'reason': 'Новая группа'})
        assert changed.status_code == 200
        assert changed.json()['employee']['group'] == 'Бар'
        created = client.post('/api/accountant/employees', json={
            'name': 'Новый сотрудник', 'role': 'официант', 'group': 'Обслуживание зала',
            'rate': '200000'})
        assert created.status_code == 201
        assert created.json()['employee']['name'] == 'Новый сотрудник'
        employee_id = created.json()['employee']['id']
        assert client.delete(f'/api/accountant/employees/{employee_id}').status_code == 204
        assert client.delete(f'/api/accountant/employees/{employee_id}').status_code == 404


def test_expected_cashier_amount_is_read_only_and_requires_fresh_real_snapshot(tmp_path):
    with demo_client(tmp_path) as client:
        assert client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['expected_cashier'] is None
        snapshot = replace(demo_snapshot(DAY), demo=False)
        client.app.state.cache.put(snapshot)
        result = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert result['expected_cashier'] == '-350000'
        assert result['ledger']['cash_balance'] == '-350000'


def test_expense_catalog_and_cash_rollforward_use_actual_outflows(tmp_path):
    with demo_client(tmp_path) as client:
        catalog = client.get('/api/accountant/expenses/catalog')
        assert catalog.status_code == 200
        items = {item['code']: item for group in catalog.json()['groups'] for item in group['items']}
        assert items['salary_staff']['label'] == 'ЗП персонал'
        assert items['salary_technical']['label'] == 'Тех персонал'
        assert 'income_cash' not in items
        client.app.state.cache.put(replace(demo_snapshot(DAY), demo=False,
                                           payments=(Payment('Демо', Decimal('1350000')),)))
        assert client.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'salary_technical',
            'note': 'Доплата вне реестра', 'amount': '200000'}).status_code == 201
        assert client.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'unknown',
            'note': 'Неизвестно', 'amount': '1'}).status_code == 422
        summary = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()['ledger']
        assert summary['cash_flow'] == {
            'opening_balance': '0', 'received_from_cashier': '1000000',
            'other_receipts': '0',
            'salary_paid': '0', 'other_outflows': '200000',
            'closing_balance': '800000', 'missing_day': None,
            'first_day': DAY.isoformat()}
        assert summary['movements'][1]['item_code'] == 'salary_technical'
        assert summary['movements'][1]['description'] == 'Тех персонал · Доплата вне реестра'
        assert summary['salary_recorded_on_day'] == '200000'


def test_employee_exports_split_late_and_everyone_without_claiming_real_hikvision(tmp_path):
    with demo_client(tmp_path) as client:
        day = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        late_count = day['payroll']['late_count']
        late = client.get('/api/accountant/employees/export', params={
            'date': DAY.isoformat(), 'scope': 'late'})
        everyone = client.get('/api/accountant/employees/export', params={
            'date': DAY.isoformat(), 'scope': 'all'})
        assert late.status_code == everyone.status_code == 200
        late_sheet = load_workbook(BytesIO(late.content), data_only=True).active
        all_sheet = load_workbook(BytesIO(everyone.content), data_only=True).active
        assert late_sheet['B4'].value == late_count
        assert all_sheet['B4'].value == 20
        assert 'ДЕМО' not in late_sheet['A1'].value
        assert 'ДЕМО' not in all_sheet['A1'].value
        assert all_sheet['D7'].value in ('Вовремя', 'Опоздал', 'Не пришёл', 'Нет привязки', 'Данных нет')
        assert client.get('/api/accountant/employees/export', params={
            'date': DAY.isoformat(), 'scope': 'invalid'}).status_code == 422
        page = client.get('/accountant/employees')
        assert page.status_code == 200
        assert 'Все сотрудники' in page.text


def test_daily_cash_starts_from_cashier_handover_without_manual_confirmation(tmp_path):
    with demo_client(tmp_path) as client:
        snapshot = replace(demo_snapshot(DAY), demo=False,
                           payments=(Payment('Демо', Decimal('1000000')),),
                           cash_prepayment=Decimal(0))
        client.app.state.cache.put(snapshot)
        before = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert before['ledger']['cash_flow']['received_from_cashier'] == '650000'
        assert before['ledger']['cash_flow']['closing_balance'] == '650000'
        assert before['ledger']['cash_balance'] == '650000'

        expense = client.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'admin_other',
            'note': 'Ремонт', 'amount': '150000'})
        assert expense.status_code == 201
        client.post('/api/accountant/payroll/confirm', json={
            'date': DAY.isoformat(), 'approver': 'Финансы'})
        accrual = client.get('/api/accountant/day', params={
            'date': DAY.isoformat()}).json()['ledger']['accruals'][0]
        payment = client.post('/api/accountant/salary-payments', json={
            'accrual_id': accrual['id'], 'date': DAY.isoformat(), 'amount': '100000'})
        assert payment.status_code == 201
        after = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert after['ledger']['cash_flow'] == {
            'opening_balance': '0', 'received_from_cashier': '650000',
            'other_receipts': '0',
            'salary_paid': '100000', 'other_outflows': '150000',
            'closing_balance': '400000', 'missing_day': None,
            'first_day': DAY.isoformat()}
        assert after['ledger']['cash_balance'] == '400000'
        assert after['ledger']['salary_recorded_on_day'] == '100000'


def test_without_cashier_data_daily_balance_is_unknown_and_expense_is_rejected(tmp_path):
    with demo_client(tmp_path) as client:
        day = client.get('/api/accountant/day', params={'date': DAY.isoformat()}).json()
        assert day['ledger']['cash_balance'] is None
        assert day['ledger']['cash_flow']['received_from_cashier'] is None
        assert client.post('/api/accountant/expenses', json={
            'date': DAY.isoformat(), 'item_code': 'admin_other',
            'note': 'Нет данных кассира', 'amount': '1'}).status_code == 409
        assert client.post('/api/accountant/transfers', json={
            'cashier_date': DAY.isoformat(), 'received_date': DAY.isoformat(),
            'amount': '100'}).status_code == 410


def test_accountant_page_still_shows_staff_when_iiko_is_unavailable(tmp_path):
    with demo_client(tmp_path) as client:
        client.app.state.settings = Settings(login='test', password='test', store_id=1)

        async def unavailable(_day):
            raise DataError('iiko временно недоступен')

        client.app.state.iiko.load = unavailable
        response = client.get('/api/accountant/day', params={'date': DAY.isoformat()})
        assert response.status_code == 200
        data = response.json()
        assert len(data['employees']) == 20
        assert data['ledger']['cash_balance'] is None
        assert data['cashier_error'] == 'iiko временно недоступен'
