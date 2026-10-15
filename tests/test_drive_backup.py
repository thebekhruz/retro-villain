import asyncio
import json
from dataclasses import replace
from datetime import date, datetime
from io import BytesIO

import httpx
import openpyxl
import pytest

from legacy_app import create_app
from retro.config import DriveBackupConfig, Settings, parse_drive_backup_config
from retro.drive_backup import BackupStore, DriveBackup, Job, Skip, build, due_jobs, last_final_day
from retro.integrations.google_drive import SHEET, DriveError, GoogleDrive
from retro.modules.cashier.service import TZ, demo_snapshot

CONFIG = DriveBackupConfig('client', 'secret', 'refresh', folder_name='Резерв', start=date(2026, 10, 13))


def at(text):
    return datetime.fromisoformat(text).replace(tzinfo=TZ)


def test_day_is_final_only_after_morning_close():
    assert last_final_day(at('2026-10-16T06:29')) == date(2026, 10, 14)
    assert last_final_day(at('2026-10-16T06:30')) == date(2026, 10, 15)


def test_due_jobs_cover_days_and_closed_months():
    jobs = due_jobs(date(2026, 10, 30), at('2026-11-02T07:00'))
    assert [job.period for job in jobs] == [
        '2026-10-30', '2026-10-30', '2026-10-31', '2026-10-31',
        '2026-10', '2026-10', '2026-10', '2026-11-01', '2026-11-01']
    assert {job.kind for job in jobs if job.period == '2026-10'} == {'payroll', 'salary_day', 'reconciliation'}
    # Ноябрь ещё идёт — его месячных файлов нет.
    assert not [job for job in jobs if job.period == '2026-11']


def test_file_names_are_readable():
    assert Job('2026-10-14', 'cashier').name == '14.10.2026 · Касса'
    assert Job('2026-10', 'payroll').name == 'Октябрь 2026 · Зарплата · месяц'


def test_config_needs_all_three_secrets():
    assert parse_drive_backup_config({}) is None
    with pytest.raises(ValueError):
        parse_drive_backup_config({'GOOGLE_DRIVE_CLIENT_ID': 'x'})
    config = parse_drive_backup_config({'GOOGLE_DRIVE_CLIENT_ID': 'a', 'GOOGLE_DRIVE_CLIENT_SECRET': 'b',
                                        'GOOGLE_DRIVE_REFRESH_TOKEN': 'refresh-secret', 'GOOGLE_BACKUP_START': '2026-10-14'})
    assert config.start == date(2026, 10, 14) and 'refresh-secret' not in repr(config)


def test_drive_client_refreshes_token_once_and_converts_xlsx():
    seen = []

    def handler(request):
        seen.append(request)
        if request.url.host == 'oauth2.googleapis.com':
            assert b'grant_type=refresh_token' in request.content
            return httpx.Response(200, json={'access_token': 'AT', 'expires_in': 3600})
        assert request.headers['authorization'] == 'Bearer AT'
        if request.url.path == '/drive/v3/files':
            return httpx.Response(200, json={'id': 'folder-1'})
        return httpx.Response(200, json={'id': 'file-1'})

    async def scenario():
        drive = GoogleDrive(CONFIG, transport=httpx.MockTransport(handler))
        folder = await drive.create_folder('Резерв')
        file_id = await drive.upload_sheet('14.10.2026 · Касса', folder, b'PK-xlsx-bytes')
        await drive.close()
        return folder, file_id

    assert asyncio.run(scenario()) == ('folder-1', 'file-1')
    assert sum(request.url.host == 'oauth2.googleapis.com' for request in seen) == 1
    upload = seen[-1]
    assert upload.url.params['uploadType'] == 'multipart'
    body = upload.content
    metadata = json.loads(body.split(b'\r\n\r\n', 1)[1].split(b'\r\n', 1)[0])
    assert metadata == {'name': '14.10.2026 · Касса', 'mimeType': SHEET, 'parents': ['folder-1']}
    assert b'PK-xlsx-bytes' in body


def test_drive_errors_do_not_leak_tokens():
    def handler(request):
        if request.url.host == 'oauth2.googleapis.com':
            return httpx.Response(400, json={'error': 'invalid_grant'})
        raise AssertionError('no Drive call without a token')

    async def scenario():
        drive = GoogleDrive(CONFIG, transport=httpx.MockTransport(handler))
        try:
            await drive.create_folder('Резерв')
        finally:
            await drive.close()

    with pytest.raises(DriveError) as error:
        asyncio.run(scenario())
    assert 'invalid_grant' in str(error.value) and 'refresh' not in str(error.value)


class FakeDrive:
    def __init__(self):
        self.folders, self.files, self.fail = [], [], set()

    async def create_folder(self, name, parent=None):
        self.folders.append((name, parent))
        return f'folder-{len(self.folders)}'

    async def upload_sheet(self, name, parent, content):
        if name in self.fail:
            raise DriveError('Google Drive: загрузить файл не удалось (503 UNAVAILABLE).')
        self.files.append((name, parent, content))
        return f'file-{len(self.files)}'

    async def close(self):
        pass


def backup(tmp_path, clock, builder, drive=None):
    return DriveBackup(None, drive or FakeDrive(), BackupStore(tmp_path / 'backup.sqlite3'), CONFIG,
                       clock=clock, builder=builder)


def test_each_report_is_uploaded_once_into_month_folder(tmp_path):
    now = [at('2026-10-15T07:00')]

    async def builder(app, job):
        return f'{job.kind}:{job.period}'.encode()

    service = backup(tmp_path, lambda: now[0], builder)
    first = asyncio.run(service.run_once())
    assert [(job.period, job.kind) for job in first] == [
        ('2026-10-13', 'cashier'), ('2026-10-13', 'accountant'),
        ('2026-10-14', 'cashier'), ('2026-10-14', 'accountant')]
    assert service.drive.folders == [('Резерв', None), ('2026-10', 'folder-1')]
    assert {parent for _, parent, _ in service.drive.files} == {'folder-2'}
    assert asyncio.run(service.run_once()) == []          # повтор ничего не дублирует
    now[0] = at('2026-10-16T07:00')
    assert [job.period for job in asyncio.run(service.run_once())] == ['2026-10-15', '2026-10-15']
    assert len(service.drive.folders) == 2                # папки запомнены в базе


def test_failure_is_retried_later_without_blocking_other_reports(tmp_path):
    now = [at('2026-10-14T07:00')]
    drive = FakeDrive()
    drive.fail.add('13.10.2026 · Касса')

    async def builder(app, job):
        return b'x'

    service = backup(tmp_path, lambda: now[0], builder, drive)
    assert [job.kind for job in asyncio.run(service.run_once())] == ['accountant']
    assert asyncio.run(service.run_once()) == []          # ждём паузу перед повтором
    drive.fail.clear()
    now[0] = at('2026-10-14T07:31')
    assert [job.kind for job in asyncio.run(service.run_once())] == ['cashier']


def test_skipped_report_is_not_requested_again(tmp_path):
    calls = []

    async def builder(app, job):
        calls.append(job)
        raise Skip('демо-данные')

    service = backup(tmp_path, lambda: at('2026-10-14T07:00'), builder)
    asyncio.run(service.run_once())
    asyncio.run(service.run_once())
    assert len(calls) == 2 and service.drive.files == []


def test_claim_blocks_second_uploader(tmp_path):
    store = BackupStore(tmp_path / 'backup.sqlite3')
    job = Job('2026-10-14', 'cashier')
    assert store.claim(job, at('2026-10-15T07:00'))
    assert not store.claim(job, at('2026-10-15T07:10'))
    assert store.claim(job, at('2026-10-15T07:40'))       # заявка упавшего процесса истекла


def test_builders_produce_the_same_files_as_cabinet_downloads(tmp_path):
    class IikoStub:
        async def load(self, day):
            return replace(demo_snapshot(day), demo=False)

        async def load_prepayments(self, day):
            return ()

    settings = Settings(login='x', password='y', store_id=1, data_dir=tmp_path)
    app = create_app(settings)
    app.state.iiko = IikoStub()

    async def scenario():
        return {kind: await build(app, Job(period, kind)) for period, kind in (
            ('2026-10-14', 'cashier'), ('2026-10-14', 'accountant'), ('2026-10', 'payroll'),
            ('2026-10', 'salary_day'), ('2026-10', 'reconciliation'))}

    files = asyncio.run(scenario())
    cashier = openpyxl.load_workbook(BytesIO(files['cashier']))
    assert cashier['отчет']['B2'].value > 0
    for content in files.values():
        assert openpyxl.load_workbook(BytesIO(content)).sheetnames


def test_demo_cashier_day_is_skipped(tmp_path):
    app = create_app(Settings(data_dir=tmp_path))
    with pytest.raises(Skip):
        asyncio.run(build(app, Job('2026-10-14', 'cashier')))
