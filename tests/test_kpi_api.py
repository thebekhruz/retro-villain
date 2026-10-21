"""KPI официантов через API: менеджер смотрит рейтинг, директор настраивает, закрывает, выгружает Excel
и передаёт премии бухгалтеру. iiko подменён источником с готовыми чеками; SQLite и Postgres."""
from dataclasses import replace
from datetime import date, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from legacy_app import create_app
from retro.config import Settings
from retro.modules.cashier.service import today_tashkent
from retro.modules.kpi.rules import Line, Order
from test_accountant_design_parity import POSTGRES_URL
from test_manager_cabinet import postgres_app

USERS = {'director': ('secret', 'director'), 'karina': ('secret', 'manager'), 'boss': ('secret', 'admin'),
         'buh': ('secret', 'accountant')}
TEA, SOUP, COLA = 'tea-id', 'soup-id', 'cola-id'
TODAY = today_tashkent()
# Прошлая полная неделя: все дни закрыты, расчёт детерминирован.
START = TODAY - timedelta(days=TODAY.weekday() + 7)
END = START + timedelta(days=6)
TARGETS = [dict(key='tea', title='Фирменные чаи', group='напиток', target100=5.0, margin=36000, dishes=[[TEA, 'ЧАЙ Фруктовый']]),
           dict(key='soup', title='Эзогелин', group='блюдо', target100=4.0, margin=32000, dishes=[[SOUP, 'Суп Эзогелин']])]
# Порог шума 0: эти тесты про деление фонда; порог проверяется отдельно.
SCHEME = dict(share=0.25, noise_week=0, cap=1_000_000, min_checks=2, min_shifts=1, baseline_weeks=4, netting=True)


class FakeSource:
    """Чеки по дням: базовые 4 недели — у Алины чай в 1 чеке из 4; неделя периода — в 3 из 4."""

    def __init__(self):
        self.calls = []

    @staticmethod
    def order(waiter, day, index, items):
        lines = tuple(Line(dish, name, 1.0, 20000.0, 5000.0) for dish, name in items)
        return Order(id=f'{waiter}-{day}-{index}', day=day, register='Kassa-FiscalBox1', section='Ресторан',
                     waiters=frozenset({waiter}), lines=lines)

    async def orders(self, start, end):
        self.calls.append((start, end))
        out = []
        day = start
        while day <= end:
            in_period = START <= day <= END
            if day.weekday() < 4:   # четыре смены в неделю
                for index in range(4):
                    tea = index < (3 if in_period else 1)
                    out.append(self.order('Алина', day, index, [(COLA, 'Кола')] + ([(TEA, 'ЧАЙ Фруктовый')] if tea else [])))
                    out.append(self.order('Нурик', day, index, [(SOUP, 'Суп Эзогелин')] if index == 0 else [(COLA, 'Кола')]))
            day += timedelta(days=1)
        return out

    async def margins(self, start, end):
        return {TEA: (10.0, 420000.0, 60000.0), SOUP: (10.0, 370000.0, 50000.0)}


def settings(tmp_path, **extra):
    return Settings(data_dir=tmp_path, dashboard_panel_users=USERS, **extra)


def sqlite_app(tmp_path, **extra):
    # В `main` KPI по умолчанию выключен (T-522): тесты модуля включают его явно.
    return create_app(settings(tmp_path, **{'kpi_module': True, **extra}), expense_db_path=tmp_path / 'cashier.sqlite3',
                      accountant_db_path=tmp_path / 'accountant.sqlite3',
                      director_db_path=tmp_path / 'director.sqlite3', founder_db_path=tmp_path / 'founder.sqlite3')


@pytest.fixture(params=['sqlite', 'postgres'])
def client(request, tmp_path):
    if request.param == 'postgres':
        if not POSTGRES_URL:
            pytest.skip('RETRO_TEST_POSTGRES_URL не задан')
        with postgres_app(tmp_path, None) as app:
            app.state.accounts = app.state.accounts  # как в кабинете менеджера
            app.state.settings = replace(app.state.settings, kpi_module=True)  # T-522: в main выключен
            app.state.kpi_source = FakeSource()
            with TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)) as http:
                yield http
    else:
        app = sqlite_app(tmp_path)
        app.state.kpi_source = FakeSource()
        with TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)) as http:
            yield http


def director(client):
    client.auth = ('director', 'secret')
    return client


def manager(client):
    client.auth = ('karina', 'secret')
    return client


def configure(client, waiters=('Алина', 'Нурик')):
    """Директор ставит период на прошлую неделю с тестовыми позициями и отмечает официантов зала."""
    response = director(client).put('/api/director/kpi/settings', json=dict(
        scheme=SCHEME, targets=TARGETS, kind='week', start=START.isoformat(), end=END.isoformat()))
    assert response.status_code == 200, response.text
    for waiter in waiters:
        assert director(client).put('/api/director/kpi/links', json=dict(waiter=waiter, active=True)).status_code == 200
    return response.json()


def test_default_period_is_the_pilot_week_and_roles_are_separated(client):
    view = manager(client).get('/api/manager/kpi').json()
    assert view['period']['status'] == 'open' and view['period']['kind'] == 'week'
    assert view['period']['pilot'] is (view['period']['start'] == '2026-10-19')
    assert date.fromisoformat(view['period']['start']).weekday() == 0
    assert date.fromisoformat(view['period']['start']) >= date(2026, 10, 19)
    assert len(view['targets']) == 12 and view['scheme']['share'] == 0.4 and view['scheme']['noise_week'] == 1_500_000
    assert view['totals']['threshold'] == 1_500_000, 'порог шума за неделю по умолчанию'
    # Менеджеру настройки закрыты, директору чужой кабинет тоже.
    assert manager(client).put('/api/director/kpi/settings', json=dict(scheme={}, targets=TARGETS)).status_code == 403
    assert director(client).get('/api/manager/kpi').status_code == 403
    assert director(client).get('/api/director/kpi').status_code == 200


def test_rating_shares_the_team_fund_by_checks(client):
    configure(client)
    view = manager(client).get('/api/manager/kpi').json()
    assert view['period']['start'] == START.isoformat() and view['period']['days_passed'] == 7
    rows = {row['waiter']: row for row in view['rows']}
    alina, nurik = rows['Алина'], rows['Нурик']
    # База зала: чай у Алины 1 из 4 чеков, у Нурика 0 → 4 чая на 32 чека = 12,5 на 100; период: 12 на 32 = 37,5.
    # Прирост зала = +25 на 100 × 32 чека × 36 000 = 288 000; суп без изменений (4 на 32 в обе недели).
    assert view['totals']['team_uplift'] == 288000 and view['totals']['fund_total'] == 72000
    assert alina['checks'] == nurik['checks'] == 16 and alina['share'] == nurik['share'] == 0.5
    assert alina['premium'] == nurik['premium'] == 36000 and view['totals']['fund'] == 72000
    assert view['totals']['restaurant_net'] == 288000 - 72000
    assert alina['items'][0]['base'] == 12.5 and alina['items'][0]['now'] == 75 and alina['sold'] == 12
    assert alina['initials'] == 'А' and alina['photo'] is None
    progress = {item['key']: item for item in view['targets_progress']}
    # Ожидаемые чеки: 128 за 16 дней базы × 7 = 56; цель чая 5 на 100 → 2,8 порции, супа 4 → 2,24.
    # Цели по правилу «база + 70 %»: база 12,5 у обеих → 21,25 → 21,5 на 100; ожидаемые чеки 56 → 12,04 порции.
    assert progress['tea']['target'] == 21.5 and progress['tea']['target_source'] == 'base'
    assert progress['tea']['goal_qty'] == 12.0 and progress['tea']['sold'] == 12 and progress['tea']['remaining'] == 1
    assert progress['soup']['sold'] == 4 and progress['soup']['lag'] > 0 and progress['soup']['push'] is True
    assert view['totals']['goal_qty'] == 24.0 and view['totals']['sold'] == 16
    assert view['hall']['rates']['tea']['target'] == 21.5 and view['fetched_at']
    assert {t['key']: t['target_source'] for t in view['targets']} == {'tea': 'base', 'soup': 'base'}
    # Повтор идёт из кэша, «Обновить» читает iiko заново.
    calls = len(client.app.state.kpi_source.calls)
    manager(client).get('/api/manager/kpi')
    assert len(client.app.state.kpi_source.calls) == calls
    manager(client).get('/api/manager/kpi?refresh=true')
    assert len(client.app.state.kpi_source.calls) == calls + 2


def test_unmarked_logins_stay_out_of_kpi_but_show_in_login_cards(client):
    configure(client, waiters=('Алина',))
    view = manager(client).get('/api/manager/kpi').json()
    assert [row['waiter'] for row in view['rows']] == ['Алина']
    assert view['totals']['waiters'] == 1 and view['hall']['checks'] == 16, 'чеки неотмеченного не идут и в итоги зала'
    logins = {item['waiter']: item for item in view['logins']}
    assert logins['Алина']['active'] and logins['Алина']['checks'] == 16 and logins['Алина']['base_checks'] == 64
    assert not logins['Нурик']['active'] and not logins['Нурик']['linked'] and logins['Нурик']['checks'] == 16
    # Отметка «не считать» хранится так же, как «официант», и переживает закрытие периода (таблица общая).
    assert director(client).put('/api/director/kpi/links', json=dict(waiter='Нурик', active=False)).status_code == 200
    closed = director(client).post('/api/director/kpi/close', json={}).json()
    assert closed['next']
    page = director(client).get('/api/director/kpi/settings').json()
    assert page['links'] == {'Алина': {'employee_id': None, 'active': True}, 'Нурик': {'employee_id': None, 'active': False}}
    snapshot = director(client).get(f'/api/director/kpi?period={closed["period"]["id"]}').json()
    assert {item['waiter'] for item in snapshot['logins']} == {'Алина', 'Нурик'}


def test_registry_waiters_take_part_by_default_and_manual_marks_win(client):
    """T-469: официанты реестра участвуют без отметок; «не считать» исключает, «официант зала» добавляет."""
    director(client).put('/api/director/kpi/settings', json=dict(
        scheme=SCHEME, targets=TARGETS, kind='week', start=START.isoformat(), end=END.isoformat()))
    roster = client.app.state.accountant_roster
    alina = roster.add(name='Каримова Алина', role='официантка', rate='150000', group_name='Обслуживание зала')
    roster.add(name='Юсупов Нурбек', role='ранер', rate='120000', group_name='Обслуживание зала')
    view = manager(client).get('/api/manager/kpi').json()
    logins = {item['waiter']: item for item in view['logins']}
    assert logins['Алина']['source'] == 'registry' and logins['Алина']['active'] and logins['Алина']['employee_id'] == alina.id
    assert logins['Нурик']['source'] == 'unmatched' and not logins['Нурик']['active'], 'ранер в реестре официантом не считается'
    assert [row['waiter'] for row in view['rows']] == ['Алина'] and view['rows'][0]['employee_name'] == 'Каримова Алина'
    assert view['registry'] == dict(waiters=1, matched=1, manual=0, unmatched=1)
    # Ручные отметки сильнее реестра.
    director(client).put('/api/director/kpi/links', json=dict(waiter='Алина', active=False))
    director(client).put('/api/director/kpi/links', json=dict(waiter='Нурик', active=True))
    view = director(client).get('/api/director/kpi?refresh=true').json()
    logins = {item['waiter']: item for item in view['logins']}
    assert logins['Алина']['source'] == 'manual_off' and not logins['Алина']['active']
    assert logins['Нурик']['source'] == 'manual' and logins['Нурик']['active']
    assert [row['waiter'] for row in view['rows']] == ['Нурик']
    # Снять ручную отметку — снова решает реестр.
    assert director(client).delete('/api/director/kpi/links', params={'waiter': 'Алина'}).status_code == 200
    view = director(client).get('/api/director/kpi').json()
    assert {item['waiter']: item['source'] for item in view['logins']}['Алина'] == 'registry'
    assert manager(client).delete('/api/director/kpi/links', params={'waiter': 'Нурик'}).status_code == 403


def test_saved_legacy_scheme_gets_threshold_and_new_share(client):
    """Период, созданный до T-490 (доля 25 %, без порога), показывается и считается по новой схеме."""
    configure(client)
    store = client.app.state.kpi_store
    period = store.current(START)
    legacy = {k: v for k, v in period['settings']['scheme'].items() if k != 'noise_week'}
    store.update_settings(period['id'], dict(period['settings'], scheme=dict(legacy, share=0.25)))
    page = director(client).get('/api/director/kpi/settings').json()
    assert page['settings']['scheme']['share'] == 0.4 and page['settings']['scheme']['noise_week'] == 1_500_000
    view = manager(client).get('/api/manager/kpi').json()
    # Прирост зала 288 000 ниже порога недели 1,5 млн: фонд 0, всё ресторану, метка порога на кольце есть.
    assert view['totals']['threshold'] == 1_500_000 and view['totals']['fund'] == 0
    assert view['totals']['restaurant_net'] == view['totals']['team_uplift'] == 288000
    assert view['totals']['to_threshold'] == 1_500_000 - 288000 and view['totals']['above_threshold'] is False
    assert 0 < view['totals']['threshold_mark'] <= 1


def test_settings_validation_and_margin_fill(client):
    bad = director(client).put('/api/director/kpi/settings', json=dict(scheme=dict(share=0.9), targets=TARGETS))
    assert bad.status_code == 422 and 'половины' in bad.json()['detail']
    # Без маржи — берётся из iiko за 8 недель до периода: (420 000 − 60 000) / 10 = 36 000.
    saved = director(client).put('/api/director/kpi/settings', json=dict(
        scheme=SCHEME, targets=[dict(TARGETS[0], margin=None)], start=START.isoformat(), end=END.isoformat())).json()
    assert saved['settings']['targets'][0]['margin'] == 36000
    items = director(client).get('/api/director/kpi/items?q=ча').json()
    assert items == {'items': []}
    page = director(client).get('/api/director/kpi/settings').json()
    assert page['period']['start'] == START.isoformat() and page['employees'] == []


def test_close_creates_snapshot_next_period_excel_and_payouts(client):
    configure(client)
    roster = client.app.state.accountant_roster
    alina = roster.add(name='Каримова Алина', role='официант', rate='150000', group_name='Обслуживание зала')
    assert director(client).put('/api/director/kpi/links', json=dict(waiter='Алина', employee_id=alina.id)).status_code == 200
    view = manager(client).get('/api/manager/kpi').json()
    assert {row['waiter']: row['employee_id'] for row in view['rows']}['Алина'] == alina.id
    closed = director(client).post('/api/director/kpi/close', json=dict(withheld={'Алина': 10000})).json()
    assert closed['period']['status'] == 'closed' and closed['next']['start'] == (END + timedelta(days=1)).isoformat()
    periods = director(client).get('/api/director/kpi/periods').json()['periods']
    assert [p['status'] for p in periods] == ['open', 'closed'] and periods[1]['fund'] == closed['totals']['fund']
    # Закрытый период отдаётся из снимка, повторное закрытие невозможно.
    again = director(client).get(f'/api/director/kpi?period={closed["period"]["id"]}').json()
    assert again['period']['status'] == 'closed' and again['totals'] == closed['totals']
    assert director(client).post(f'/api/director/kpi/close?period={closed["period"]["id"]}', json={}).status_code == 409
    excel = director(client).get(f'/api/director/kpi/export?period={closed["period"]["id"]}')
    assert excel.status_code == 200 and excel.headers['content-type'].startswith('application/vnd.openxmlformats')
    book = load_workbook(BytesIO(excel.content))
    assert book.sheetnames == ['Премии', 'Позиции', 'Настройки']
    rows = list(book['Премии'].iter_rows(values_only=True))
    alina_row = next(row for row in rows if row[0] == 'Алина')
    assert alina_row[1] == 'Каримова Алина' and alina_row[8] == 10000 and alina_row[9] == alina_row[7] - 10000
    # Премии → доп. выплаты бухгалтера: привязанному заводится, непривязанный остаётся в списке.
    paid = director(client).post(f'/api/director/kpi/payouts?period={closed["period"]["id"]}', json={}).json()
    by_waiter = {item['waiter']: item for item in paid['results']}
    assert by_waiter['Нурик']['status'] == 'unlinked'
    assert by_waiter['Алина']['status'] in ('created', 'error'), by_waiter['Алина']
    if by_waiter['Алина']['status'] == 'created':
        payouts = client.app.state.accountant_finance.extra_payouts(TODAY, TODAY)
        assert any(p['note'].startswith('KPI официантов') for p in payouts)
        assert director(client).post(f'/api/director/kpi/payouts?period={closed["period"]["id"]}', json={}).status_code == 409


def test_kpi_module_flag_hides_api_and_tells_pages(tmp_path, monkeypatch):
    """T-522: KPI официантов обкатывается только на стенде. Без переменной — по
    умолчанию ветки (в `main` выключен); KPI_MODULE=0 — /api/*/kpi отвечают 404, а
    /api/config говорит страницам спрятать вкладку и карточку."""
    assert Settings().kpi_module is False
    monkeypatch.delenv('KPI_MODULE', raising=False)
    assert Settings.from_env().kpi_module is Settings.kpi_module
    monkeypatch.setenv('KPI_MODULE', '0')
    assert Settings.from_env().kpi_module is False
    monkeypatch.setenv('KPI_MODULE', '1')
    assert Settings.from_env().kpi_module is True

    app = create_app(settings(tmp_path, kpi_module=False), expense_db_path=tmp_path / 'cashier.sqlite3',
                     accountant_db_path=tmp_path / 'accountant.sqlite3',
                     director_db_path=tmp_path / 'director.sqlite3', founder_db_path=tmp_path / 'founder.sqlite3')
    app.state.kpi_source = FakeSource()
    with TestClient(app, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)) as http:
        assert director(http).get('/api/config').json()['kpi'] is False
        assert director(http).get('/api/director/kpi').status_code == 404
        assert director(http).post('/api/director/kpi/close', json={}).status_code == 404
        assert manager(http).get('/api/manager/kpi').status_code == 404
        assert manager(http).get('/manager').status_code == 200  # сам кабинет на месте
    plain = create_app(settings(tmp_path), expense_db_path=tmp_path / 'cashier.sqlite3',
                       accountant_db_path=tmp_path / 'accountant.sqlite3',
                       director_db_path=tmp_path / 'director.sqlite3', founder_db_path=tmp_path / 'founder.sqlite3')
    with TestClient(plain, base_url='http://127.0.0.1', client=('127.0.0.1', 50000)) as http:
        assert director(http).get('/api/config').json()['kpi'] is False  # умолчание main
        assert manager(http).get('/api/manager/kpi').status_code == 404
