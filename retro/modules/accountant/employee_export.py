"""Demo attendance exports for the accountant module."""

from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from retro.modules.cashier.service import TZ

from .payroll import PayrollRow

GREEN = '16483F'
INK = '21342E'
PINK = 'FCE6E9'
PALE = 'EAF0E9'
STATUS = {'on_time': 'Вовремя', 'late': 'Опоздал',
          'missing': 'Не пришёл', 'unlinked': 'Нет привязки'}


def export_employees(day: date, rows: list[PayrollRow], scope: str) -> bytes:
    """Export only supplied demo rows; the caller decides late versus all."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Опоздавшие' if scope == 'late' else 'Все сотрудники'
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'C7'
    for column, width in {'A': 7, 'B': 36, 'C': 27, 'D': 20,
                          'E': 16, 'F': 18, 'G': 20}.items():
        sheet.column_dimensions[column].width = width
    sheet.merge_cells('A1:G2')
    title = sheet['A1']
    title.value = 'RETRO MILLIY / ДЕМО / ' + ('ОПОЗДАВШИЕ' if scope == 'late' else 'ВСЕ СОТРУДНИКИ')
    title.fill = PatternFill('solid', fgColor=GREEN)
    title.font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    title.alignment = Alignment(vertical='center', indent=1)
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 25
    sheet['A3'] = 'Дата'
    sheet['B3'] = datetime.combine(day, time.min)
    sheet['B3'].number_format = 'dd.mm.yyyy'
    sheet['A4'] = 'Сотрудников'
    sheet['B4'] = len(rows)
    sheet['C4'] = 'Опоздали'
    sheet['D4'] = sum(row.status == 'late' for row in rows)
    for row_index in (3, 4):
        for col in range(1, 8):
            cell = sheet.cell(row_index, col)
            cell.fill = PatternFill('solid', fgColor=PALE)
            cell.font = Font(name='Calibri', size=11, bold=col in (1, 3), color=INK)
            cell.alignment = Alignment(vertical='center', indent=1)
        sheet.row_dimensions[row_index].height = 24
    sheet.merge_cells('A5:G5')
    sheet['A5'] = 'Демонстрационные проходы. Реальный Hikvision ресторана не подключён.'
    sheet['A5'].font = Font(name='Calibri', size=10, italic=True, color='6D7F73')
    sheet['A5'].alignment = Alignment(vertical='center', indent=1)
    sheet.row_dimensions[5].height = 27
    headers = ('№', 'Сотрудник', 'Группа', 'Статус', 'Первый вход', 'Ставка, сум', 'Начисление, сум')
    for col, label in enumerate(headers, 1):
        cell = sheet.cell(6, col, label)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', horizontal='center')
    sheet.row_dimensions[6].height = 29
    for index, person in enumerate(rows, 1):
        row = index + 6
        at = person.occurred_at.astimezone(TZ).replace(tzinfo=None) if person.occurred_at else None
        values = (index, person.name, person.group_name, STATUS[person.status],
                  at, float(person.rate) if person.rate is not None else None,
                  float(person.payable) if person.payable is not None else None)
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row, col, value)
            if col in (2, 3, 4):
                cell.data_type = 's'
            cell.fill = PatternFill('solid', fgColor=PINK if person.status == 'late'
                                    else ('FFFFFF' if index % 2 else 'F5F7F2'))
            cell.font = Font(name='Calibri', size=10,
                             bold=col == 4 and person.status == 'late',
                             color='A0344B' if person.status == 'late' else INK)
            cell.alignment = Alignment(vertical='center', indent=1)
            cell.border = Border(bottom=Side(style='hair', color='DDE5DB'))
        sheet.cell(row, 5).number_format = 'hh:mm'
        for col in (6, 7):
            sheet.cell(row, col).number_format = '#,##0.##'
        sheet.row_dimensions[row].height = 27
    if rows:
        sheet.auto_filter.ref = f'A6:G{6 + len(rows)}'
    else:
        sheet.merge_cells('A7:G7')
        sheet['A7'] = 'За выбранный день записей нет'
        sheet['A7'].alignment = Alignment(horizontal='center', vertical='center')
        sheet.row_dimensions[7].height = 38
    sheet.print_area = f'A1:G{max(7, 6 + len(rows))}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
