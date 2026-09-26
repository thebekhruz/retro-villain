"""Local employee roster imported from the approved payroll worksheet."""

# Аннотации не вычисляются при импорте: ниже в классе есть метод list(),
# и на Python 3.12 подпись «-> list[MonthlyEmployee]» бралась бы за него,
# а не за встроенный тип. На 3.14 аннотации ленивые и это не всплывает,
# поэтому локально всё работало, а на сервере приложение не поднималось.
from __future__ import annotations

import sqlite3
from contextlib import closing
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from retro.db import as_database, table_columns
from retro.runtime import secure_directory, secure_file
from retro.integrations.hikvision import HikvisionPerson
from retro.modules.cashier.service import today_tashkent


GROUPS = {
    'менеджер': 'Управление',
    'хостес': 'Встреча гостей',
    'официант': 'Обслуживание зала',
    'ранер': 'Обслуживание зала',
    'бармен': 'Бар',
    'няня': 'Присмотр за детьми',
    'техперсонал': 'Уборка',
    'охрана': 'Охрана',
}

UNASSIGNED_ROLE = 'Должность не указана'
UNASSIGNED_GROUP = 'Не распределено'


def group_for(role: str) -> str:
    normalized = ' '.join(role.casefold().split())
    if normalized.startswith('повар') or normalized == 'кондитер':
        return 'Кухня'
    if normalized not in GROUPS:
        raise ValueError(f'Неизвестная должность в реестре: {role}')
    return GROUPS[normalized]


def parse_rate(value) -> Decimal | None:
    if value is None or value == '':
        return None
    try:
        rate = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError('Дневная ставка должна быть числом.') from None
    if not rate.is_finite() or rate <= 0 or rate.as_tuple().exponent < -2:
        raise ValueError('Дневная ставка должна быть положительной суммой.')
    return rate


def parse_money(value, *, allow_zero=True) -> Decimal:
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError('Сумма должна быть числом.') from None
    if (not amount.is_finite() or amount < 0 or (amount == 0 and not allow_zero)
            or amount > Decimal('1000000000000') or amount.as_tuple().exponent < -2):
        raise ValueError('Сумма должна быть неотрицательной, не более 1 трлн сум и с точностью до тиына.')
    return amount


def normalized_hikvision_name(value: str) -> str:
    # Payroll exports and Hikvision may store the same full name in a
    # different order (surname first vs. given name first).  Sorting complete
    # whitespace-delimited parts keeps the match strict while ignoring order;
    # the two-sided uniqueness checks below still reject duplicate names.
    return ' '.join(sorted(value.casefold().split()))


@dataclass(frozen=True)
class Employee:
    id: int
    source_row: int
    name: str
    role: str
    group_name: str
    rate: Decimal | None
    hikvision_id: str | None
    # «Нет в Hikvision · отмечать вручную»: человек проходит мимо турникета
    # (охрана, уборка). Привязку к устройству такой сотрудник не получает, и
    # его день отмечает бухгалтер, как у любого непривязанного.
    manual_attendance: bool = False

    def json(self):
        return dict(id=self.id, name=self.name, role=self.role, group=self.group_name,
                    rate=str(self.rate) if self.rate is not None else None,
                    hikvision_registered=self.hikvision_id is not None,
                    manual_attendance=self.manual_attendance)


@dataclass(frozen=True)
class MonthlyEmployee:
    id: int
    external_key: str | None
    name: str
    role: str
    salary: Decimal
    schedule: str
    card: Decimal
    cash: Decimal
    advances: Decimal
    remaining: Decimal

    def json(self):
        return dict(id=self.id, external_key=self.external_key, name=self.name, role=self.role,
                    salary=str(self.salary), schedule=self.schedule, card=str(self.card),
                    cash=str(self.cash), advances=str(self.advances), remaining=str(self.remaining),
                    accounting_basis='manual_current_register',
                    warning='Ручной текущий реестр: поля выплат и остатка не сверены с движениями; месяц не задан.')


class RosterStore:
    def __init__(self, path: Path):
        self.db = as_database(path)
        # .path остаётся для скриптов обслуживания и тестов
        self.path = self.db.path
        self._initialize()

    def _open(self):
        return self.db.connect()

    def _initialize(self):
        with closing(self._open()) as connection, connection:
            connection.execute('PRAGMA journal_mode=WAL')
            connection.execute('''CREATE TABLE IF NOT EXISTS accountant_employees (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_row INTEGER NOT NULL UNIQUE,
                name TEXT NOT NULL,
                role TEXT NOT NULL,
                group_name TEXT NOT NULL,
                rate TEXT,
                hikvision_id TEXT UNIQUE
            )''')
            connection.execute('''CREATE TABLE IF NOT EXISTS accountant_roster_audit (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                employee_id INTEGER NOT NULL,
                changed_at TEXT NOT NULL,
                reason TEXT NOT NULL,
                old_rate TEXT,
                new_rate TEXT,
                old_group TEXT NOT NULL,
                new_group TEXT NOT NULL
            )''')
            connection.execute('''CREATE TABLE IF NOT EXISTS accountant_monthly_employees (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                external_key TEXT UNIQUE,
                name TEXT NOT NULL,
                role TEXT NOT NULL,
                salary TEXT NOT NULL,
                schedule TEXT NOT NULL DEFAULT '',
                card TEXT NOT NULL DEFAULT '0',
                cash TEXT NOT NULL DEFAULT '0',
                advances TEXT NOT NULL DEFAULT '0',
                remaining TEXT NOT NULL DEFAULT '0'
            )''')
            employee_columns = table_columns(connection, 'accountant_employees')
            if 'manual_attendance' not in employee_columns:
                connection.execute('ALTER TABLE accountant_employees '
                                   'ADD COLUMN manual_attendance INTEGER NOT NULL DEFAULT 0')
            columns = table_columns(connection, 'accountant_monthly_employees')
            if 'external_key' not in columns:
                connection.execute('ALTER TABLE accountant_monthly_employees ADD COLUMN external_key TEXT')
                connection.execute(
                    'CREATE UNIQUE INDEX IF NOT EXISTS accountant_monthly_external_key '
                    'ON accountant_monthly_employees(external_key) WHERE external_key IS NOT NULL')
            # Preserve values used for previous calendar days. The initial version
            # is a baseline, not a reconstruction of changes before this migration.
            connection.executescript('''
                CREATE TABLE IF NOT EXISTS accountant_employee_versions (
                    employee_id INTEGER NOT NULL, effective_day TEXT NOT NULL,
                    source_row INTEGER NOT NULL, name TEXT NOT NULL, role TEXT NOT NULL,
                    group_name TEXT NOT NULL, rate TEXT, hikvision_id TEXT,
                    deleted INTEGER NOT NULL DEFAULT 0,
                    PRIMARY KEY(employee_id, effective_day)
                );
                INSERT OR IGNORE INTO accountant_employee_versions
                    SELECT id,'0001-01-01',source_row,name,role,group_name,rate,hikvision_id,0
                    FROM accountant_employees WHERE id NOT IN
                        (SELECT employee_id FROM accountant_employee_versions);
                -- Историю ставок ведёт сам код (_stamp_version), а не триггеры:
                -- они были на диалекте SQLite и в Postgres не переносятся, а
                -- один денежный инвариант не должен жить в двух вариантах.
                -- У существующих баз триггеры сносим, иначе история писалась бы
                -- дважды.
                DROP TRIGGER IF EXISTS accountant_employee_insert_history;
                DROP TRIGGER IF EXISTS accountant_employee_update_history;
                DROP TRIGGER IF EXISTS accountant_employee_delete_history;
            ''')


    # ── История ставок ─────────────────────────────────────────────────────
    # `list(day)` читает версию с наибольшим effective_day <= день, поэтому
    # каждая правка реестра обязана оставить снимок на день правки. Раньше это
    # делали три триггера SQLite; теперь — этот метод, вызываемый из всех путей
    # записи в одной с ними транзакции.

    def _stamp_version(self, connection, employee_id: int, *, deleted: bool = False,
                       day: str | None = None) -> None:
        """Записать текущее состояние сотрудника в историю на указанный день."""
        effective = day or today_tashkent().isoformat()
        connection.execute(
            'INSERT INTO accountant_employee_versions '
            '(employee_id, effective_day, source_row, name, role, group_name, rate, '
            ' hikvision_id, deleted) '
            'SELECT id, ?, source_row, name, role, group_name, rate, hikvision_id, ? '
            'FROM accountant_employees WHERE id = ? '
            'ON CONFLICT(employee_id, effective_day) DO UPDATE SET '
            'source_row=excluded.source_row, name=excluded.name, role=excluded.role, '
            'group_name=excluded.group_name, rate=excluded.rate, '
            'hikvision_id=excluded.hikvision_id, deleted=excluded.deleted',
            (effective, 1 if deleted else 0, employee_id))

    def import_xlsx(self, source: Path, *, replace: bool = False) -> dict[str, int]:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            if 'ЗП' not in workbook.sheetnames:
                raise ValueError('В файле нет листа «ЗП».')
            sheet = workbook['ЗП']
            rows = []
            for cells in sheet.iter_rows(min_row=5, min_col=1, max_col=4):
                number, name, role, raw_rate = (cell.value for cell in cells)
                if not isinstance(number, int) or not isinstance(name, str) or not name.strip():
                    continue
                if not isinstance(role, str):
                    raise ValueError(f'Строка {cells[0].row}: у сотрудника нет должности.')
                rows.append((cells[0].row, name.strip(), role.strip(), group_for(role),
                             str(parse_rate(raw_rate)) if raw_rate is not None else None))
        finally:
            workbook.close()
        imported = existing = 0
        with closing(self._open()) as connection:
            with connection:
                if replace:
                    source_rows = {row[0] for row in rows}
                    placeholders = ','.join('?' for _ in source_rows)
                    # Снимок «удалён» снимаем до удаления, пока строки ещё есть.
                    for (gone_id,) in connection.execute(
                            'SELECT id FROM accountant_employees WHERE source_row NOT IN (%s)'
                            % placeholders, tuple(source_rows)).fetchall():
                        self._stamp_version(connection, gone_id, deleted=True)
                    connection.execute('DELETE FROM accountant_employees WHERE source_row NOT IN (%s)' %
                                       placeholders, tuple(source_rows))
                for row in rows:
                    cursor = connection.execute(
                        'SELECT id,name FROM accountant_employees WHERE source_row = ?',
                        (row[0],))
                    current = cursor.fetchone()
                    if current is None:
                        new_id = connection.execute(
                            'INSERT INTO accountant_employees '
                            '(source_row, name, role, group_name, rate) VALUES (?, ?, ?, ?, ?)',
                            row).lastrowid
                        # Импорт задаёт исходное состояние, поэтому версия с
                        # начала времён: прошлые дни должны видеть сотрудника.
                        self._stamp_version(connection, new_id, day='0001-01-01')
                        imported += 1
                    else:
                        if replace:
                            preserve_link = (normalized_hikvision_name(current[1]) ==
                                             normalized_hikvision_name(row[1]))
                            if not preserve_link:
                                raise ValueError('Строка импорта относится к другому сотруднику. '
                                                 'Нельзя заменять человека под прежним ID.')
                            connection.execute(
                                'UPDATE accountant_employees SET name=?, role=?, group_name=?, rate=?, '
                                'hikvision_id=CASE WHEN ? THEN hikvision_id ELSE NULL END '
                                'WHERE source_row=?',
                                (row[1], row[2], row[3], row[4], preserve_link, row[0]))
                            self._stamp_version(connection, current[0])
                            imported += 1
                        else:
                            existing += 1
        return {'imported': imported, 'existing': existing}

    def list(self, day=None) -> list[Employee]:
        with closing(self._open()) as connection:
            if day is None:
                rows = connection.execute('SELECT id, source_row, name, role, group_name, rate, hikvision_id, '
                                          'manual_attendance '
                                          'FROM accountant_employees ORDER BY source_row').fetchall()
            else:
                rows = connection.execute('''
                    SELECT employee_id,source_row,name,role,group_name,rate,
                        COALESCE(v.hikvision_id,
                            (SELECT e.hikvision_id FROM accountant_employees e WHERE e.id=v.employee_id),
                            (SELECT h.hikvision_id FROM accountant_employee_versions h
                             WHERE h.employee_id=v.employee_id AND h.hikvision_id IS NOT NULL
                             ORDER BY h.effective_day DESC LIMIT 1)),
                        COALESCE((SELECT e.manual_attendance FROM accountant_employees e
                                  WHERE e.id=v.employee_id), 0)
                    FROM accountant_employee_versions v WHERE deleted=0 AND effective_day=(
                        SELECT MAX(effective_day) FROM accountant_employee_versions h
                        WHERE h.employee_id=v.employee_id AND h.effective_day<=?)
                    ORDER BY source_row
                ''', (day.isoformat(),)).fetchall()
        # Ручная отметка сильнее привязки: турникет такого человека не видит,
        # и его вход не должен ни засчитываться, ни считаться прогулом.
        return [Employee(*row[:5], Decimal(row[5]) if row[5] is not None else None,
                         None if row[7] else row[6], bool(row[7])) for row in rows]

    def set_manual_attendance(self, employee_id: int, manual: bool) -> Employee:
        with closing(self._open()) as connection, connection:
            changed = connection.execute(
                'UPDATE accountant_employees SET manual_attendance = ? WHERE id = ?',
                (1 if manual else 0, employee_id)).rowcount
            if changed:
                self._stamp_version(connection, employee_id)
        if not changed:
            raise ValueError('Сотрудник не найден.')
        return next(person for person in self.list() if person.id == employee_id)

    def link_hikvision_people(self, people: tuple[HikvisionPerson, ...]) -> dict[str, int]:
        """Link only two-sided unique exact normalized names; never overwrite IDs."""
        unique_people = {person.employee_no: person for person in people if person.employee_no}
        employees = self.list()
        already_ids = {employee.hikvision_id for employee in employees if employee.hikvision_id}
        unlinked_by_name: dict[str, list[Employee]] = {}
        for employee in employees:
            if employee.hikvision_id is None and not employee.manual_attendance:
                key = normalized_hikvision_name(employee.name)
                if key:
                    unlinked_by_name.setdefault(key, []).append(employee)
        people_by_name: dict[str, list[HikvisionPerson]] = {}
        for person in unique_people.values():
            if person.employee_no in already_ids:
                continue
            key = normalized_hikvision_name(person.name or '')
            if key:
                people_by_name.setdefault(key, []).append(person)

        result = dict(people=len(unique_people), linked=0, already_linked=0,
                      ambiguous=0, unmatched=0)
        with closing(self._open()) as connection, connection:
            for person in unique_people.values():
                if person.employee_no in already_ids:
                    result['already_linked'] += 1
                    continue
                key = normalized_hikvision_name(person.name or '')
                candidates = unlinked_by_name.get(key, []) if key else []
                device_candidates = people_by_name.get(key, []) if key else []
                if key and (len(candidates) > 1 or len(device_candidates) > 1):
                    result['ambiguous'] += 1
                    continue
                if not key or not candidates:
                    result['unmatched'] += 1
                    continue
                try:
                    changed = connection.execute(
                        'UPDATE accountant_employees SET hikvision_id = ? '
                        'WHERE id = ? AND hikvision_id IS NULL',
                        (person.employee_no, candidates[0].id)).rowcount
                except sqlite3.IntegrityError:
                    changed = 0
                if changed:
                    self._stamp_version(connection, candidates[0].id)
                    result['linked'] += 1
                    already_ids.add(person.employee_no)
                else:
                    result['ambiguous'] += 1
        return result

    def import_hikvision_people(self, people: tuple[HikvisionPerson, ...]) -> dict[str, int]:
        """Create missing device people with explicit placeholders and link by employeeNo.

        This is intentionally separate from background polling: importing people changes the
        accountant roster and must only be triggered by an authenticated operator.
        """
        unique_people = {person.employee_no: person for person in people if person.employee_no}
        people_by_name: dict[str, list[HikvisionPerson]] = {}
        for person in unique_people.values():
            key = normalized_hikvision_name(person.name or '')
            if key:
                people_by_name.setdefault(key, []).append(person)

        result = dict(people=len(unique_people), created=0, linked=0,
                      already_linked=0, ambiguous=0)
        with closing(self._open()) as connection, connection:
            rows = connection.execute(
                'SELECT id,name,hikvision_id FROM accountant_employees ORDER BY source_row').fetchall()
            by_name: dict[str, list[tuple[int, str | None]]] = {}
            linked_ids = set()
            for employee_id, name, hikvision_id in rows:
                key = normalized_hikvision_name(name)
                if key:
                    by_name.setdefault(key, []).append((employee_id, hikvision_id))
                if hikvision_id:
                    linked_ids.add(hikvision_id)

            next_source_row = connection.execute(
                'SELECT COALESCE(MAX(source_row), 0) + 1 FROM accountant_employees').fetchone()[0]
            for person in unique_people.values():
                if person.employee_no in linked_ids:
                    result['already_linked'] += 1
                    continue
                key = normalized_hikvision_name(person.name or '')
                device_matches = people_by_name.get(key, []) if key else []
                roster_matches = by_name.get(key, []) if key else []
                if not key or len(device_matches) != 1 or len(roster_matches) > 1:
                    result['ambiguous'] += 1
                    continue
                if roster_matches:
                    employee_id, current_link = roster_matches[0]
                    if current_link:
                        result['ambiguous'] += 1
                        continue
                    connection.execute(
                        'UPDATE accountant_employees SET hikvision_id=? WHERE id=?',
                        (person.employee_no, employee_id))
                    self._stamp_version(connection, employee_id)
                    linked_ids.add(person.employee_no)
                    result['linked'] += 1
                    continue
                # last_insert_rowid() есть только в SQLite; слой отдаёт номер
                # через lastrowid на обоих диалектах.
                new_id = connection.execute(
                    'INSERT INTO accountant_employees '
                    '(source_row,name,role,group_name,rate,hikvision_id) VALUES (?,?,?,?,NULL,?)',
                    (next_source_row, person.name.strip(), UNASSIGNED_ROLE,
                     UNASSIGNED_GROUP, person.employee_no)).lastrowid
                self._stamp_version(connection, new_id, day='0001-01-01')
                by_name[key] = [(new_id, person.employee_no)]
                linked_ids.add(person.employee_no)
                next_source_row += 1
                result['created'] += 1
        return result

    def monthly_total(self) -> Decimal:
        """Return only salaries explicitly stored in the monthly payroll register."""
        return sum((person.salary for person in self.list_monthly()), Decimal(0))
    def list_monthly(self) -> list[MonthlyEmployee]:
        with closing(self._open()) as connection:
            rows = connection.execute(
                'SELECT id, external_key, name, role, salary, schedule, card, cash, advances, remaining '
                'FROM accountant_monthly_employees ORDER BY role, name, id').fetchall()
        return [MonthlyEmployee(row[0], row[1], row[2], row[3], Decimal(row[4]), row[5],
                                Decimal(row[6]), Decimal(row[7]), Decimal(row[8]), Decimal(row[9]))
                for row in rows]

    @staticmethod
    def _monthly_values(*, name, role, salary, schedule='', card='0', cash='0', advances='0',
                        remaining='0'):
        name = name.strip() if isinstance(name, str) else ''
        role = role.strip() if isinstance(role, str) else ''
        schedule = schedule.strip() if isinstance(schedule, str) else ''
        if not name or len(name) > 160 or not role or len(role) > 80:
            raise ValueError('Укажите корректные имя и должность.')
        if len(schedule) > 160:
            raise ValueError('График должен быть не длиннее 160 символов.')
        money = {field: parse_money(value) for field, value in {
            'salary': salary, 'card': card, 'cash': cash,
            'advances': advances, 'remaining': remaining,
        }.items()}
        return name, role, schedule, money

    def add_monthly(self, *, name: str, role: str, salary: str, schedule: str = '',
                    card: str = '0', cash: str = '0', advances: str = '0',
                    remaining: str = '0', external_key: str | None = None) -> MonthlyEmployee:
        name, role, schedule, money = self._monthly_values(
            name=name, role=role, salary=salary, schedule=schedule, card=card, cash=cash,
            advances=advances, remaining=remaining)
        key = external_key.strip() if isinstance(external_key, str) and external_key.strip() else None
        with closing(self._open()) as connection, connection:
            try:
                employee_id = connection.execute(
                    'INSERT INTO accountant_monthly_employees '
                    '(external_key,name,role,salary,schedule,card,cash,advances,remaining) '
                    'VALUES (?,?,?,?,?,?,?,?,?)',
                    (key, name, role, str(money['salary']), schedule, str(money['card']),
                     str(money['cash']), str(money['advances']), str(money['remaining']))).lastrowid
            except sqlite3.IntegrityError:
                raise ValueError('Сотрудник с таким внешним ключом уже существует.') from None
        return next(item for item in self.list_monthly() if item.id == employee_id)

    def update_monthly(self, employee_id: int, *, name: str, role: str, salary: str,
                       schedule: str = '', card: str = '0', cash: str = '0',
                       advances: str = '0', remaining: str = '0') -> MonthlyEmployee:
        name, role, schedule, money = self._monthly_values(
            name=name, role=role, salary=salary, schedule=schedule, card=card, cash=cash,
            advances=advances, remaining=remaining)
        with closing(self._open()) as connection, connection:
            changed = connection.execute(
                'UPDATE accountant_monthly_employees SET name=?,role=?,salary=?,schedule=?,card=?,cash=?, '
                'advances=?,remaining=? WHERE id=?',
                (name, role, str(money['salary']), schedule, str(money['card']), str(money['cash']),
                 str(money['advances']), str(money['remaining']), employee_id)).rowcount
        if not changed:
            raise ValueError('Сотрудник не найден.')
        return next(item for item in self.list_monthly() if item.id == employee_id)

    def update_monthly_basics(self, employee_id: int, *, name: str, role: str,
                              salary: str) -> MonthlyEmployee:
        """Имя, должность и оклад без касания выплат, которые ведёт бухгалтер."""
        name, role, _, money = self._monthly_values(name=name, role=role, salary=salary)
        with closing(self._open()) as connection, connection:
            changed = connection.execute(
                'UPDATE accountant_monthly_employees SET name=?,role=?,salary=? WHERE id=?',
                (name, role, str(money['salary']), employee_id)).rowcount
        if not changed:
            raise ValueError('Сотрудник не найден.')
        return next(item for item in self.list_monthly() if item.id == employee_id)

    def delete_monthly(self, employee_id: int):
        with closing(self._open()) as connection, connection:
            deleted = connection.execute(
                'DELETE FROM accountant_monthly_employees WHERE id = ?', (employee_id,)).rowcount
        if not deleted:
            raise ValueError('Сотрудник не найден.')

    def add(self, *, name: str, role: str, rate: str | None, group_name: str) -> Employee:
        name = name.strip()
        role = role.strip()
        if not name or len(name) > 160:
            raise ValueError('Укажите корректное имя сотрудника.')
        if not role or len(role) > 80:
            raise ValueError('Укажите корректную должность.')
        if group_name not in set(GROUPS.values()) | {'Кухня', UNASSIGNED_GROUP}:
            raise ValueError('Неизвестная группа.')
        parsed_rate = parse_rate(rate)
        with closing(self._open()) as connection, connection:
            source_row = connection.execute('SELECT COALESCE(MAX(source_row), 0) + 1 '
                                            'FROM accountant_employees').fetchone()[0]
            employee_id = connection.execute(
                'INSERT INTO accountant_employees (source_row, name, role, group_name, rate) '
                'VALUES (?, ?, ?, ?, ?)',
                (source_row, name, role, group_name,
                 str(parsed_rate) if parsed_rate is not None else None)).lastrowid
            # Новый сотрудник действует с начала времён: иначе прошлые дни его
            # не увидят, а начисления за них уже закрыты.
            self._stamp_version(connection, employee_id, day='0001-01-01')
        return next(person for person in self.list() if person.id == employee_id)

    def delete(self, employee_id: int):
        with closing(self._open()) as connection, connection:
            # Снимок снимаем до удаления: после него строки уже нет.
            self._stamp_version(connection, employee_id, deleted=True)
            deleted = connection.execute('DELETE FROM accountant_employees WHERE id = ?',
                                         (employee_id,)).rowcount
        if not deleted:
            raise ValueError('Сотрудник не найден.')

    def update(self, employee_id: int, *, name: str | None = None, role: str | None = None,
               rate: str | None,
               group_name: str | None = None, reason: str) -> Employee:
        if not reason.strip():
            raise ValueError('Укажите причину изменения.')
        with closing(self._open()) as connection:
            existing = connection.execute('SELECT name, role FROM accountant_employees WHERE id = ?',
                                          (employee_id,)).fetchone()
        if existing is None:
            raise ValueError('Сотрудник не найден.')
        name = (name or existing[0]).strip()
        role = (role or existing[1]).strip()
        if not name or len(name) > 160:
            raise ValueError('Укажите корректное имя сотрудника.')
        if not role or len(role) > 80:
            raise ValueError('Укажите корректную должность.')
        derived_group = group_name or group_for(role)
        if derived_group not in set(GROUPS.values()) | {'Кухня', UNASSIGNED_GROUP}:
            raise ValueError('Неизвестная группа.')
        parsed_rate = parse_rate(rate)
        with closing(self._open()) as connection:
            with connection:
                old = connection.execute('SELECT name, role, rate, group_name FROM accountant_employees WHERE id = ?',
                                         (employee_id,)).fetchone()
                if old is None:
                    raise ValueError('Сотрудник не найден.')
                connection.execute('UPDATE accountant_employees SET name = ?, role = ?, rate = ?, group_name = ? '
                                   'WHERE id = ?', (name, role,
                                   str(parsed_rate) if parsed_rate is not None else None,
                                   derived_group, employee_id))
                self._stamp_version(connection, employee_id)
                connection.execute('INSERT INTO accountant_roster_audit '
                                   '(employee_id, changed_at, reason, old_rate, new_rate, old_group, new_group) '
                                   'VALUES (?, ?, ?, ?, ?, ?, ?)',
                                   (employee_id, datetime.now().isoformat(), reason.strip(), old[2],
                                    str(parsed_rate) if parsed_rate is not None else None, old[3], derived_group))
        return next(person for person in self.list() if person.id == employee_id)
