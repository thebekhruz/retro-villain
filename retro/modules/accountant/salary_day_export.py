"""«Зарплата · день» в Excel для печати (ТЗ 09.10, Б-07).

Тот же ответ, что видит экран (`/salary-day/month`): строка — сотрудник с
должностью и ставкой, столбец — день выплаты с подписью смены («Выплата 09.10 ·
смена 08.10»), итоги по сотруднику и по дню. Доп. выплаты входят в итог дня
выплаты и сотрудника, как на экране, а сами расписаны вторым листом — с датой
смены, датой выплаты и назначением.

Вся ведомость или выборка — группа и поиск, как на экране; режим написан в
шапке обоих листов. Печать: альбомный лист по ширине страницы, шапка таблицы
повторяется на каждой странице, имена и ставки закреплены при прокрутке.
"""

import re
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

from .sheets_export import _head, _money, _note, _row, _sum, _title, _widths

MONTHS = ('январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 'июль', 'август',
          'сентябрь', 'октябрь', 'ноябрь', 'декабрь')
HEAD_ROW = 4
NO_GROUP = 'Без группы'

# Поиск — как на экране (employees-logic.js, matchesQuery): без регистра,
# «ё» = «е», кириллица находится латиницей (Alijon → Алижон).
CYR = {'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'е': 'e', 'ж': 'j', 'з': 'z', 'и': 'i', 'й': 'y',
       'к': 'k', 'л': 'l', 'м': 'm', 'н': 'n', 'о': 'o', 'п': 'p', 'р': 'r', 'с': 's', 'т': 't', 'у': 'u',
       'ф': 'f', 'х': 'x', 'ц': 's', 'ч': 'ch', 'ш': 'sh', 'щ': 'sh', 'ъ': '', 'ы': 'i', 'ь': '', 'э': 'e',
       'ю': 'yu', 'я': 'ya', 'ў': 'o', 'қ': 'q', 'ғ': 'g', 'ҳ': 'x'}


def _plain(value) -> str:
    return ' '.join(str(value or '').lower().replace('ё', 'е').split())


def _latin(value) -> str:
    text = ' '.join(str(value or '').lower().split()).replace('ё', 'yo')
    text = ''.join(CYR.get(ch, ch) for ch in text)
    text = re.sub(r"[ʻʼ’‘'`]", '', text)
    text = text.replace('kh', 'x').replace('zh', 'j').replace('ts', 's')
    text = re.sub(r'(^|[^sc])h', r'\1x', text)
    return text.replace('ye', 'e').replace('w', 'v')


def matches_query(query, *fields) -> bool:
    wanted = _plain(query)
    if not wanted:
        return True
    latin = _latin(wanted)
    return any(wanted in _plain(field) or (latin and latin in _latin(field)) for field in fields)


def select_people(people: list[dict], group: str | None, query: str | None,
                  role: str | None = None) -> list[dict]:
    # role — должность внутри группы (второй ряд чипов), сравнение как на экране.
    return [person for person in people
            if (not group or (person.get('group') or NO_GROUP) == group)
            and (not role or _plain(person.get('role')) == _plain(role))
            and matches_query(query, person.get('name'), person.get('role'))]


def _dm(day: str) -> str:
    return f'{day[8:10]}.{day[5:7]}'


def _dmy(day: str) -> str:
    return f'{day[8:10]}.{day[5:7]}.{day[:4]}'


def _previous(day: str) -> str:
    return (date.fromisoformat(day) - timedelta(days=1)).isoformat()


def _people(count: int) -> str:
    tail = count % 100
    word = ('сотрудник' if tail % 10 == 1 and tail != 11 else
            'сотрудника' if 2 <= tail % 10 <= 4 and not 12 <= tail <= 14 else 'сотрудников')
    return f'{count} {word}'


def _mode(group, query, shown: int, total: int, role=None) -> str:
    if not group and not query and not role:
        return f'Вся ведомость, {_people(total)}'
    parts = (([f'группа «{group}»'] if group else []) + ([f'должность «{role}»'] if role else [])
             + ([f'поиск «{query}»'] if query else []))
    # После «из» — родительный падеж: «из 4 сотрудников», «из 21 сотрудника».
    word = 'сотрудника' if total % 10 == 1 and total % 100 != 11 else 'сотрудников'
    return 'Выборка: ' + ', '.join(parts) + f' — {shown} из {total} {word}'


def _print(sheet, last_column: int, last_row: int) -> None:
    """Альбомный A4 по ширине страницы, шапка таблицы — на каждой странице."""
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f'1:{HEAD_ROW}'
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4
    sheet.page_margins.top = sheet.page_margins.bottom = 0.5
    sheet.print_area = f'A1:{get_column_letter(last_column)}{last_row}'


def salary_day_workbook(data: dict, *, group: str | None = None, query: str | None = None,
                        role: str | None = None) -> bytes:
    """`data` — ответ `/salary-day/month`; `group`, `role`, `query` — выборка экрана."""
    today = data['today']
    # Будущие дни месяца пусты — на бумаге они только сужают столбцы.
    days = [day for day in data['days'] if day <= today]
    people = select_people(data.get('people') or [], group, query, role)
    chosen = {person['id'] for person in people}
    extras = [item for item in data.get('extras') or [] if item['employee_id'] in chosen and item['paid_day'] in days]
    month = date.fromisoformat(data['month'] + '-01')
    period = (f'{_dmy(days[0])} – {_dmy(days[-1])}' if days else '—')
    mode = _mode(group, query, len(people), len(data.get('people') or []), role)
    workbook = Workbook()
    _sheet(workbook.active, data, days, people, extras, month, period, mode)
    _extras_sheet(workbook.create_sheet('Доп. выплаты'), extras, month, period, mode)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet(sheet, data, days, people, extras, month, period, mode):
    sheet.title = 'Ведомость'
    labels = (['№', 'Сотрудник', 'Должность', 'Ставка']
              + [f'Выплата {_dm(day)}\nсмена {_dm(_previous(day))}' for day in days]
              + ['Доп. выплаты', 'Итого'])
    width = len(labels)
    _title(sheet, f'RETRO MILLIY · Зарплата · день · {MONTHS[month.month - 1]} {month.year}', width)
    _note(sheet, 2, f'Период выплат {period}. {mode}. В столбце — деньги, выданные в этот день за смену '
          'предыдущего дня. Доп. выплаты входят в итоги и расписаны на листе «Доп. выплаты».', width)
    sheet.row_dimensions[2].height = 30
    _widths(sheet, [5, 30, 18, 12] + [11] * len(days) + [13, 14])
    _head(sheet, HEAD_ROW, labels)
    sheet.row_dimensions[HEAD_ROW].height = 32
    for column in range(1, width + 1):
        sheet.cell(HEAD_ROW, column).alignment = Alignment(
            vertical='center', horizontal='left' if column in (2, 3) else 'center', wrap_text=True)
    extra_by = {}
    for item in extras:
        key = (item['employee_id'], item['paid_day'])
        extra_by[key] = extra_by.get(key, Decimal(0)) + Decimal(item['amount'])
    row = HEAD_ROW + 1
    per_day = [Decimal(0)] * len(days)
    extra_per_day = [Decimal(0)] * len(days)
    for number, person in enumerate(people, 1):
        cells = [_money((person['cells'].get(day) or {}).get('amount')) or Decimal(0) for day in days]
        own_extra = [extra_by.get((person['id'], day), Decimal(0)) for day in days]
        for index in range(len(days)):
            per_day[index] += cells[index] + own_extra[index]
            extra_per_day[index] += own_extra[index]
        extra = _sum(own_extra)
        role = person.get('role') or ''
        if person.get('temporary'):
            role = (role + ' · временный').strip(' ·')
        if person.get('archived'):
            role = (role + ' · архив').strip(' ·')
        _row(sheet, row, [number, person['name'], role, _money(person.get('rate'))]
             + [value or None for value in cells] + [extra or None, _sum(cells) + extra], money_from=4)
        # Доп. выплата в этот день — ячейка подсвечена, сумма — в столбце «Доп. выплаты».
        for index, value in enumerate(own_extra):
            if value:
                sheet.cell(row, 5 + index).fill = PatternFill('solid', fgColor='FBF0D7')
        row += 1
    if not people:
        _note(sheet, row, 'Никого не нашли: выборка пустая.', width)
        row += 1
    total = _sum(per_day)
    _row(sheet, row, [None, 'Итого за день', None, None] + [value or None for value in per_day]
         + [_sum(extra_per_day) or None, total], bold=True, money_from=4)
    row += 1
    _row(sheet, row, [None, 'в т.ч. доп. выплаты', None, None] + [value or None for value in extra_per_day]
         + [_sum(extra_per_day) or None, _sum(extra_per_day) or None], money_from=4)
    for column in range(1, width + 1):
        sheet.cell(row - 1, column).fill = PatternFill('solid', fgColor='F5F7F1')
    sheet.freeze_panes = sheet.cell(HEAD_ROW + 1, 5).coordinate
    _print(sheet, width, row)


def _extras_sheet(sheet, extras, month, period, mode):
    labels = ['№', 'Дата выплаты', 'Дата смены', 'Сотрудник', 'Должность', 'Временный', 'Сумма',
              'Назначение', 'Записал', 'Когда']
    width = len(labels)
    _title(sheet, f'Доп. выплаты · {MONTHS[month.month - 1]} {month.year}', width)
    _note(sheet, 2, f'Период выплат {period}. {mode}. Расход — в «Финансах дня» за дату выплаты.', width)
    sheet.row_dimensions[2].height = 30
    _widths(sheet, [5, 13, 12, 28, 18, 11, 14, 40, 16, 16])
    _head(sheet, HEAD_ROW, labels)
    row = HEAD_ROW + 1
    for number, item in enumerate(extras, 1):
        stamp = item.get('created_at') or ''
        _row(sheet, row, [number, _dmy(item['paid_day']), _dmy(item['work_day']), item['name'], item['role'],
                          'да' if item.get('temporary') else '', _money(item['amount']), item['note'],
                          item.get('created_by') or '', f'{_dmy(stamp[:10])} {stamp[11:16]}' if stamp else ''],
             money_from=7)
        sheet.cell(row, 8).alignment = Alignment(wrap_text=True, vertical='center')
        row += 1
    if not extras:
        _note(sheet, row, 'Доп. выплат за период нет.', width)
        row += 1
    _row(sheet, row, [None, 'Итого', None, None, None, None, _sum(_money(item['amount']) for item in extras),
                      None, None, None], bold=True, money_from=7)
    sheet.freeze_panes = sheet.cell(HEAD_ROW + 1, 1).coordinate
    _print(sheet, width, row)

