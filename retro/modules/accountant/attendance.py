from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from retro.modules.cashier.service import TZ

LATE_AFTER = time(10, 0)
GREEN = '16483F'
PALE = 'EAF0E9'
PINK = 'FCE6E9'
INK = '21342E'


@dataclass(frozen=True)
class Entrance:
    name: str
    occurred_at: datetime


def is_late(occurred_at: datetime) -> bool:
    if occurred_at.tzinfo is None:
        raise ValueError('Время входа должно содержать часовой пояс.')
    return occurred_at.astimezone(TZ).time().replace(tzinfo=None) > LATE_AFTER


def export_entrances(day: date, entries: tuple[Entrance, ...] = (),
                     health: dict | None = None) -> bytes:
    """Styled XLSX of real first entries for the selected day."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Входы'
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'B7'
    for letter, width in {'A': 8, 'B': 37, 'C': 19, 'D': 18}.items():
        sheet.column_dimensions[letter].width = width
    sheet.merge_cells('A1:D2')
    title = sheet['A1']
    title.value = 'RETRO MILLIY  /  ВХОДЫ'
    title.fill = PatternFill('solid', fgColor=GREEN)
    title.font = Font(name='Calibri', size=20, bold=True, color='FFFFFF')
    title.alignment = Alignment(vertical='center', indent=1)
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 25

    sheet['A3'] = 'Дата'
    sheet['B3'] = datetime.combine(day, time.min)
    sheet['B3'].number_format = 'dd.mm.yyyy'
    sheet['C3'] = 'После 10:00'
    sheet['D3'] = 'Опоздание'
    sheet['A4'] = 'Входов'
    sheet['B4'] = len(entries)
    sheet['C4'] = 'Опоздали'
    sheet['D4'] = sum(is_late(entry.occurred_at) for entry in entries)
    for row in (3, 4):
        sheet.row_dimensions[row].height = 25
        for col in range(1, 5):
            cell = sheet.cell(row, col)
            cell.fill = PatternFill('solid', fgColor=PALE)
            cell.font = Font(name='Calibri', size=11, bold=col in (1, 3), color=INK)
            cell.alignment = Alignment(vertical='center', indent=1)

    sheet.merge_cells('A5:D5')
    status = (health or {}).get('status')
    sheet['A5'] = ('Hikvision не настроен; отсутствие входа не считается прогулом.'
                   if status == 'not_configured' else
                   'Источник Hikvision недоступен или день покрыт не полностью.'
                   if status in {'starting', 'network', 'timeout', 'unauthorized',
                                 'invalid_response', 'device_error', 'internal', 'stale'} else
                   'Первый подтверждённый вход каждого человека за день.')
    sheet['A5'].font = Font(name='Calibri', size=10, italic=True, color='6D7F73')
    sheet['A5'].alignment = Alignment(vertical='center', indent=1)
    sheet.row_dimensions[5].height = 27

    for col, value in enumerate(('№', 'Человек', 'Время входа', 'Статус'), 1):
        cell = sheet.cell(6, col, value)
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', indent=1)
    sheet.row_dimensions[6].height = 30

    if not entries:
        sheet.merge_cells('A7:D7')
        sheet['A7'] = 'Данные о входах пока не подключены'
        sheet['A7'].font = Font(name='Calibri', size=11, italic=True, color='839187')
        sheet['A7'].alignment = Alignment(vertical='center', horizontal='center')
        sheet['A7'].fill = PatternFill('solid', fgColor='F6F8F4')
        sheet.row_dimensions[7].height = 42
    else:
        for index, entry in enumerate(entries, 1):
            row = index + 6
            late = is_late(entry.occurred_at)
            at = entry.occurred_at.astimezone(TZ)
            values = (index, entry.name, at.replace(tzinfo=None), 'Опоздал' if late else 'Вовремя')
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row, col, value)
                if col in (2, 4):
                    cell.data_type = 's'  # Names from devices must stay text in Excel.
                cell.fill = PatternFill('solid', fgColor=PINK if late else ('FFFFFF' if index % 2 else 'F5F7F2'))
                cell.font = Font(name='Calibri', size=11, bold=col == 4 and late,
                                 color='A0344B' if late else INK)
                cell.alignment = Alignment(vertical='center', indent=1)
                cell.border = Border(bottom=Side(style='hair', color='DDE5DB'))
            sheet.cell(row, 3).number_format = 'hh:mm:ss'
            sheet.row_dimensions[row].height = 29
        sheet.auto_filter.ref = f'A6:D{6 + len(entries)}'

    last = max(7, 6 + len(entries))
    sheet.print_area = f'A1:{get_column_letter(4)}{last}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
