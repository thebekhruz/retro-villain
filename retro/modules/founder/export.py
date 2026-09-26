"""Отчёт бухгалтера за месяц в Excel для учредителя.

Деньги берутся из журнала бухгалтера, выручка и чеки — из iiko по кассам.
Если iiko недоступен, файл всё равно собирается: колонки iiko остаются
пустыми, а в шапке написано почему, — пустая ячейка не выдаёт себя за ноль.
"""

from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import overview

MONEY = '#,##0'
HEAD = PatternFill('solid', fgColor='143E35')


def text(cell, value):
    cell.value = value
    cell.data_type = 's'  # подписи из журнала не должны становиться формулами Excel


def closing_balance(state, day: date, handover):
    finance = state.accountant_finance
    manual = state.settings.manual_handover_only
    anchor = finance.cash_opening()
    carry_start = date.fromisoformat(anchor['day']) if anchor and manual else None
    summary = finance.daily_summary(
        day, handover, carry_history=not manual or (carry_start is not None and day >= carry_start),
        carry_start=carry_start)
    return summary['cash_balance']


def month_workbook(state, first: date, last: date, orders, orders_error):
    flows_rows = state.accountant_finance.cash_flows_between(first, last)
    flows = overview.daily_flows(flows_rows)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'По дням'
    title = f'Retro Milliy · отчёт бухгалтера · {first.strftime("%m.%Y")}'
    text(sheet['A1'], title)
    sheet['A1'].font = Font(bold=True, size=13)
    note = ('Выручка и чеки — iiko по кассам, без банкетного зала. Деньги — журнал бухгалтера.'
            if orders is not None else f'Колонки iiko пустые: {orders_error}')
    text(sheet['A2'], note)
    heads = ['Дата', 'Retro · выручка', 'Oxbridge · выручка', 'Чеки Retro', 'Чеки Oxbridge',
             'Получено от кассира', 'Прочие поступления', 'Зарплаты', 'Закуп', 'Прочие расходы',
             'Дивиденды', 'Остаток на конец дня']
    for column, head in enumerate(heads, start=1):
        cell = sheet.cell(row=4, column=column)
        text(cell, head)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    row_index = 5
    day = first
    while day <= last:
        values = flows.get(day.isoformat(), {})
        registers = (orders or {}).get(day.isoformat(), {})
        handover = values.get('handover')
        balance = closing_balance(state, day, handover)
        cells = [
            datetime.combine(day, datetime.min.time()),
            registers.get('retro', {}).get('revenue') if orders is not None else None,
            registers.get('school', {}).get('revenue') if orders is not None else None,
            registers.get('retro', {}).get('orders') if orders is not None else None,
            registers.get('school', {}).get('orders') if orders is not None else None,
            handover, values.get('receipt', Decimal(0)), values.get('salary', Decimal(0)),
            values.get('procurement', Decimal(0)), values.get('other', Decimal(0)),
            values.get('dividends', Decimal(0)), balance,
        ]
        for column, value in enumerate(cells, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            if column == 1:
                cell.number_format = 'dd.mm.yyyy'
            elif column not in (4, 5):
                cell.number_format = MONEY
        row_index += 1
        day += timedelta(days=1)
    for column, width in zip('ABCDEFGHIJKL', (12, 16, 16, 11, 11, 16, 14, 14, 14, 14, 14, 16)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'B5'

    categories = workbook.create_sheet('Расходы')
    text(categories['A1'], f'Куда ушли деньги · {first.strftime("%m.%Y")}')
    categories['A1'].font = Font(bold=True, size=13)
    spending = overview.expense_categories(flows_rows)
    for column, head in enumerate(('Категория', 'Сумма', 'Доля, %'), start=1):
        cell = categories.cell(row=3, column=column)
        text(cell, head)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEAD
    for index, item in enumerate(spending['categories'], start=4):
        text(categories.cell(row=index, column=1), item['label'])
        categories.cell(row=index, column=2, value=Decimal(item['amount'])).number_format = MONEY
        share = item['share_percent']
        categories.cell(row=index, column=3, value=Decimal(share) if share else None)
    total_row = 4 + len(spending['categories'])
    text(categories.cell(row=total_row, column=1), 'Итого')
    categories.cell(row=total_row, column=1).font = Font(bold=True)
    categories.cell(row=total_row, column=2, value=Decimal(spending['total'])).number_format = MONEY
    categories.column_dimensions['A'].width = 32
    categories.column_dimensions['B'].width = 16
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
