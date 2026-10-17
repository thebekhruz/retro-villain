"""Excel закрытого периода для бухгалтера: премии по людям, позиции, настройки."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HEAD = Font(bold=True)


def _sheet(workbook, title, header, rows, widths):
    sheet = workbook.create_sheet(title) if workbook.worksheets and workbook.worksheets[0].title != 'Sheet' \
        else workbook.active
    sheet.title = title
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for row in rows:
        sheet.append(row)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    return sheet


def kpi_workbook(view: dict, *, withheld=None) -> bytes:
    withheld = withheld or {}
    period = view['period']
    label = f"{period['start']} – {period['end']}"
    workbook = Workbook()
    people = []
    for row in view['rows']:
        hold = int(withheld.get(row['waiter'], 0) or 0)
        people.append([row['waiter'], row.get('employee_name') or '', row['checks'], row['shifts'],
                       'да' if row['eligible'] else 'нет', row.get('sold', 0), row.get('share', 0.0),
                       row['premium'], hold, max(row['premium'] - hold, 0)])
    _sheet(workbook, 'Премии',
           ['Логин iiko', 'Сотрудник', 'Чеков', 'Смен', 'Участвует', 'Продано целевых, порций', 'Доля фонда',
            'Премия, сум', 'Удержано, сум', 'К выплате, сум'], people,
           [18, 26, 8, 7, 10, 14, 10, 14, 14, 14])
    totals = view['totals']
    summary = workbook['Премии']
    summary.append([])
    summary.append(['Период', label])
    summary.append(['Прирост маржи зала по целевым позициям, сум', totals['team_uplift']])
    summary.append(['Фонд премии (доля прироста), сум', totals.get('fund_total', totals['fund'])])
    summary.append(['Выплачено с учётом потолка, сум', totals['fund']])
    summary.append(['Ресторану сверх премии, сум', totals['restaurant_net']])
    summary.append(['Данные iiko на', view.get('fetched_at') or ''])
    items = []
    rates = view['hall']['rates']
    for target in view['targets']:
        rate = rates.get(target['key'], {})
        items.append([target['title'], target['group'], target['margin'], rate.get('base'), rate.get('now'), target['target100'],
                      'вручную' if target.get('target_source') == 'manual' else 'база + 30 %',
                      '; '.join(name for _id, name in target['dishes'])])
    _sheet(workbook, 'Позиции', ['Позиция', 'Группа', 'Маржа на порцию, сум', 'База зала на 100 чеков',
                                 'Период на 100 чеков', 'Цель на 100 чеков', 'Откуда цель', 'Карточки iiko'], items,
           [34, 10, 16, 14, 14, 14, 14, 60])
    scheme = view['scheme']
    _sheet(workbook, 'Настройки', ['Параметр', 'Значение'], [
        ['Период', label], ['Доля прироста зала в фонд', scheme['share']], ['Делится', 'между участниками по числу чеков'],
        ['Потолок на человека в неделю, сум', scheme['cap']],
        ['Минимум чеков', scheme['min_checks']], ['Минимум смен', scheme['min_shifts']],
        ['База зала, недель', scheme['baseline_weeks']],
        ['Нетто по позициям', 'да' if scheme['netting'] else 'нет']], [34, 34])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
