from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill

from .service import PAYMENT_SOURCES
from .expenses import cash_to_finance

TEMPLATE = Path(__file__).parent / 'templates' / 'cashier.xlsx'
MONEY = '#,##0.00'


def text(cell, value):
    cell.value = value
    cell.data_type = 's'  # iiko labels must never become executable Excel formulas.


def export_report(snapshot, expenses=(), receipts=()):
    expense_total = sum((item.amount for item in expenses), Decimal(0))
    receipt_total = sum((item.amount for item in receipts), Decimal(0))
    handover = cash_to_finance(snapshot, expense_total, receipt_total)
    workbook = openpyxl.load_workbook(TEMPLATE)
    sheet = workbook['отчет']
    # Preserve layout/style; erase the example's transactions and out-of-scope formulas.
    for region in ['A4:B19', 'D2:D11', 'C13:D19', 'C22:D37', 'C40:D55', 'B20:B23', 'D20:D20', 'D38:D39']:
        for cells in sheet[region]:
            for cell in cells:
                if not isinstance(cell, MergedCell):
                    cell.value = None
    sheet['A1'] = datetime.combine(snapshot.day, datetime.min.time())
    sheet['A1'].number_format = 'dd.mm.yyyy'
    sheet['B2'] = snapshot.revenue
    sheet['B2'].number_format = MONEY
    for label, pos in [('Количество чеков', 'A25'), ('Средний чек', 'A26')]:
        sheet[pos] = label
    sheet['B25'] = snapshot.receipt_count
    sheet['B26'] = snapshot.average_receipt
    sheet['B26'].number_format = MONEY
    sheet['A20'], sheet['B20'] = 'ИТОГО НОВЫЕ ПРЕДОПЛАТЫ', snapshot.new_prepayment
    sheet['A21'], sheet['B21'] = 'ОБЩИЙ ПРИХОД', snapshot.revenue + snapshot.new_prepayment + receipt_total
    sheet['B20'].number_format = MONEY
    sheet['B21'].number_format = MONEY
    sheet['A27'], sheet['B27'] = 'Новые предоплаты наличными · iiko', snapshot.cash_prepayment
    sheet['A28'], sheet['B28'] = 'Расходы наличными (включая зарплату)', expense_total
    sheet['A29'], sheet['B29'] = 'К передаче в финансовый отдел', handover
    sheet['A30'], sheet['B30'] = 'Прочие поступления наличными', receipt_total
    for position in ('B27', 'B28', 'B29', 'B30'):
        sheet[position].number_format = MONEY
    sheet['A31'] = 'Расчёт: Демо + предоплаты наличными + прочие поступления − расходы.'
    sheet['A39'] = ('ДЕМОНСТРАЦИЯ — НЕ ОТЧЁТ iiko' if snapshot.demo else
                    'Сформировано из iiko • только Retro, без школы и зала Бехруз')
    amounts = {p.name: p.amount for p in snapshot.payments}
    sheet['C2'] = 'ВЫРУЧКА ПО ТИПАМ ОПЛАТЫ'
    sheet['C2'].font = Font(name='Calibri', size=11, bold=True)
    sheet['C2'].alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[2].height = 30
    for row, name in enumerate(PAYMENT_SOURCES, 3):
        text(sheet.cell(row, 3), name)
        sheet.cell(row, 3).font = Font(name='Calibri', size=11)
        sheet.cell(row, 3).alignment = Alignment(wrap_text=True, vertical='center')
        sheet.cell(row, 4, amounts[name]).number_format = MONEY
        sheet.row_dimensions[row].height = 30 if len(name) > 28 else 22
    sheet['C11'] = 'ИТОГО ВЫРУЧКА:'
    sheet['D11'] = f'=SUM(D3:D{2 + len(PAYMENT_SOURCES)})'
    sheet['D11'].number_format = MONEY
    sheet['C13'] = 'РАСХОДЫ КАССЫ'
    sheet['C13'].font = Font(name='Calibri', size=11, bold=True, color='173D38')
    sheet['C14'], sheet['D14'] = 'Название', 'Сумма, сум'
    for row, item in enumerate(expenses[:22], 15):
        text(sheet.cell(row, 3), item.description)
        sheet.cell(row, 4, item.amount).number_format = MONEY
    if len(expenses) > 22:
        sheet['C37'] = f'Ещё {len(expenses) - 22} — на листе «Расходы»'
    sheet['C38'], sheet['D38'] = 'ИТОГО РАСХОДЫ', expense_total
    sheet['C39'], sheet['D39'] = 'К ПЕРЕДАЧЕ', handover
    for position in ('D38', 'D39'):
        sheet[position].number_format = MONEY
    sheet.print_area = 'A1:D40'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_view.showGridLines = False

    # The second template tab was a manual scratchpad. Use it for exact iiko labels,
    # since they need not match the hard-coded payment categories on the first tab.
    detail = workbook['Касса']
    for merged in list(detail.merged_cells.ranges):
        detail.unmerge_cells(str(merged))
    for cells in detail:
        for cell in cells:
            cell.value = None
    detail['A1'] = 'Retro Milliy — оплаты'
    detail['A2'] = datetime.combine(snapshot.day, datetime.min.time())
    detail['A2'].number_format = 'dd.mm.yyyy'
    detail['A3'], detail['B3'] = 'Продажи после скидок', snapshot.revenue
    detail['A4'], detail['B4'] = 'Количество чеков', snapshot.receipt_count
    detail['A5'], detail['B5'] = 'Средний чек', snapshot.average_receipt
    detail['A7'], detail['B7'], detail['C7'] = 'Тип оплаты', 'Выручка, сум', 'Доля'
    for index, payment in enumerate(snapshot.payments, 8):
        text(detail.cell(index, 1), payment.name)
        detail.cell(index, 2, payment.amount)
        detail.cell(index, 3, payment.amount / snapshot.revenue if snapshot.revenue else None)
    end = 8 + len(snapshot.payments)
    detail.cell(end, 1, 'Итого по оплатам')
    detail.cell(end, 2, f'=SUM(B8:B{end-1})')
    detail.cell(end, 3, 1 if snapshot.revenue else None)
    detail.cell(end + 2, 1, 'Суммы по отдельным типам оплаты получены из iiko (PayTypes).')
    detail.cell(end + 3, 1, 'Дата: OpenDate.Typed. Удалённые заказы и позиции исключены.')
    detail.cell(end + 4, 1, 'Источник: https://retro3158.iikoweb.ru')
    detail.cell(end + 5, 1, 'Получено: ' + snapshot.fetched_at.strftime('%d.%m.%Y %H:%M') + ' (Ташкент)')
    detail.cell(end + 6, 1, 'Только касса Retro, без отделения «Бехруз (Свадьба)»; операции PAYMENT.')
    if snapshot.demo:
        detail.cell(end + 7, 1, 'ДЕМОНСТРАЦИЯ. Данные вымышлены, не использовать для учёта.')
    detail.column_dimensions['A'].width = 68
    detail.column_dimensions['B'].width = 24
    detail.column_dimensions['C'].width = 12
    for row in range(1, end + 8):
        detail.row_dimensions[row].height = 24
        for col in (1, 2, 3):
            c = detail.cell(row, col)
            c.font = Font(name='Calibri', size=11, color='173D38')
            c.border = Border()
            c.fill = PatternFill('solid', fgColor='FFFFFF')
            c.alignment = Alignment(vertical='center', wrap_text=True)
            if col == 2:
                c.number_format = '#,##0' if row == 4 else MONEY
            if col == 3:
                c.number_format = '0%'
    for row in (1, 7, end):
        for col in (1, 2, 3):
            c = detail.cell(row, col)
            c.fill = PatternFill('solid', fgColor='16483F')
            c.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    detail.sheet_view.showGridLines = False
    detail.freeze_panes = 'A8'
    detail.print_area = f'A1:C{end + 7}'
    detail.sheet_properties.pageSetUpPr.fitToPage = True
    detail.page_setup.fitToWidth = 1
    detail.page_setup.fitToHeight = 0

    manual = workbook.create_sheet('Расходы')
    manual.sheet_view.showGridLines = False
    manual['A1'] = 'RETRO MILLIY · РАСХОДЫ КАССИРА'
    manual['A2'] = datetime.combine(snapshot.day, datetime.min.time())
    manual['A2'].number_format = 'dd.mm.yyyy'
    manual['A4'], manual['B4'] = 'Демо · по данным iiko', next(
        (p.amount for p in snapshot.payments if p.name == 'Демо'), Decimal(0))
    manual['A5'], manual['B5'] = 'Расходы наличными · включая зарплату', expense_total
    manual['A6'], manual['B6'] = 'К передаче в финансовый отдел', handover
    manual['C4'], manual['D4'] = 'Предоплаты наличными · iiko', snapshot.cash_prepayment
    manual['C5'], manual['D5'] = 'Прочие поступления · вручную', receipt_total
    manual['C6'], manual['D6'] = 'Общий приход · все источники', snapshot.revenue + snapshot.new_prepayment + receipt_total
    manual['A7'], manual['B7'] = 'Название расхода', 'Сумма, сум'
    manual['C7'], manual['D7'] = 'Прочее поступление', 'Сумма, сум'
    for row, item in enumerate(expenses, 8):
        text(manual.cell(row, 1), item.description)
        manual.cell(row, 2, item.amount)
    for row, item in enumerate(receipts, 8):
        text(manual.cell(row, 3), item.description)
        manual.cell(row, 4, item.amount)
    last_row = max(8, 7 + len(expenses), 7 + len(receipts))
    if not expenses:
        manual['A8'] = 'Расходы не внесены'
    manual.column_dimensions['A'].width = 54
    manual.column_dimensions['B'].width = 22
    manual.column_dimensions['C'].width = 48
    manual.column_dimensions['D'].width = 22
    for row in range(1, last_row + 1):
        manual.row_dimensions[row].height = 25
        for col in (1, 2, 3, 4):
            cell = manual.cell(row, col)
            cell.font = Font(name='Calibri', size=11, color='173D38')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if col in (2, 4) and row >= 4:
                cell.number_format = MONEY
    for row in (1, 6, 7):
        for col in (1, 2, 3, 4):
            cell = manual.cell(row, col)
            cell.fill = PatternFill('solid', fgColor='16483F')
            cell.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    manual.freeze_panes = 'A8'
    manual.print_area = f'A1:D{last_row}'
    manual.sheet_properties.pageSetUpPr.fitToPage = True
    manual.page_setup.fitToWidth = 1
    manual.page_setup.fitToHeight = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
